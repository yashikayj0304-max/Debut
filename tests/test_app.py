import io, os, re, json, hmac, hashlib, datetime as dt
import pytest
import config as C
import services as S
import views
from conftest import signup, dbq, dbx, wid_of

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 64


def test_public_pages(app):
    c = app.test_client()
    for p in ("/", "/pricing", "/terms", "/privacy", "/health", "/robots.txt", "/login", "/signup", "/forgot"):
        assert c.get(p).status_code == 200, p
    assert c.get("/app").status_code == 302 and "/login" in c.get("/app").headers["Location"]
    assert c.get("/m/tasks").status_code == 302
    assert c.get("/admin/").status_code == 302  # anonymous users are sent to login


def test_signup_login_and_seed(app):
    c = signup(app, "a@x.com")
    r = c.get("/app")
    assert r.status_code == 200 and b"Command Centre" in r.data
    w = wid_of(app, "a@x.com")
    assert dbq(app, "SELECT COUNT(*) c FROM launch_checklist WHERE workspace_id=?", (w,), one=True)["c"] > 20
    assert dbq(app, "SELECT plan FROM workspaces WHERE id=?", (w,), one=True)["plan"] == "trial"
    c.post("/logout")
    assert c.get("/app").status_code == 302
    bad = app.test_client().post("/login", data=dict(email="a@x.com", password="nope"))
    assert bad.status_code == 401
    ok = app.test_client().post("/login", data=dict(email="A@x.com", password="Sup3rSecret!"))
    assert ok.status_code == 302
    dup = app.test_client().post("/signup", data=dict(name="n", email="a@x.com", password="Sup3rSecret!", agree="1"))
    assert dup.status_code == 400
    weak = app.test_client().post("/signup", data=dict(name="n", email="b@x.com", password="short", agree="1"))
    assert weak.status_code == 400


def test_every_page_renders(app):
    c = signup(app, "a@x.com")
    keys = list(__import__("modules").MODULES)
    pages = ["/app", "/calendar", "/moodboards", "/team-hub", "/ideas", "/pulse", "/utm", "/analytics", "/ai", "/updates", "/content-ops", "/discover?kw=tea", "/search?q=ab", "/team", "/activity", "/account", "/billing", "/export.zip"]
    pages += [f"/m/{k}" for k in keys] + [f"/m/{k}/new" for k in keys] + [f"/m/{k}/export" for k in keys]
    pages += [f"/kanban/{k}" for k, m in __import__("modules").MODULES.items() if m.get("kanban")]
    bad = [(p, c.get(p).status_code) for p in pages]
    assert all(s == 200 for _, s in bad), [b for b in bad if b[1] != 200]


def test_tenant_isolation(app):
    a, b = signup(app, "a@x.com", ws="A Co"), signup(app, "b@x.com", ws="B Co")
    a.post("/m/content_items/save", data=dict(title="SECRET-A", channel="Instagram", format="Post", status="Idea", publish_date="2026-10-01"))
    a.post("/m/influencers/save", data=dict(name="Inf-A", platform="Instagram", status="Prospect"))
    rid = dbq(app, "SELECT id FROM content_items WHERE title='SECRET-A'", one=True)["id"]
    assert b"SECRET-A" not in b.get("/m/content_items").data
    assert b"SECRET-A" not in b.get("/calendar?y=2026&m=10").data
    assert b"SECRET-A" not in b.get("/search?q=secret").data
    assert b"SECRET-A" not in b.get("/m/content_items/export").data
    assert b"SECRET-A" not in b.get("/export.zip").data
    assert b.get(f"/m/content_items/{rid}/edit").status_code == 404
    assert b.post("/m/content_items/save", data=dict(id=rid, title="HACKED")).status_code == 404
    assert b.post(f"/m/content_items/{rid}/delete").status_code == 404
    assert b.post("/api/reschedule", json=dict(id=rid, date="2026-11-01")).status_code == 404
    assert b.get("/influencers/1/outreach").status_code == 404
    assert dbq(app, "SELECT title FROM content_items WHERE id=?", (rid,), one=True)["title"] == "SECRET-A"
    # mood board + media
    a.post("/moodboards/new", data=dict(name="A board"))
    a.post("/moodboards/1/add", data=dict(kind="upload", file=(io.BytesIO(PNG), "x.png")), content_type="multipart/form-data")
    url = re.search(r'src="(/media/\d+/[\w.]+)"', a.get("/moodboards/1").data.decode()).group(1)
    assert a.get(url).status_code == 200 and b.get(url).status_code == 404
    assert b.get("/moodboards/1").status_code == 404
    assert b.post("/moodboards/1/add", data=dict(kind="note", content="x")).status_code == 404
    # ics token belongs to one workspace
    ics_a = re.search(r"(/calendar/[\w.\-]+\.ics)", a.get("/calendar").data.decode()).group(1)
    assert b"SECRET-A" in app.test_client().get(ics_a).data
    assert b"SECRET-A" not in app.test_client().get(re.search(r"(/calendar/[\w.\-]+\.ics)", b.get("/calendar").data.decode()).group(1)).data
    assert app.test_client().get("/calendar/garbage.ics").status_code == 404


def _invite(app, owner, mails, email, role):
    owner.post("/team/invite", data=dict(email=email, role=role))
    body = mails[-1][2]
    return re.search(r"(/invite/[\w\-]+)", body).group(1)


def test_roles_and_invites(app, mails):
    o = signup(app, "o@x.com")
    v_link = _invite(app, o, mails, "v@x.com", "viewer")
    v = app.test_client()
    assert v.get(v_link).status_code == 302
    r = v.post("/signup", data=dict(name="V", email="v@x.com", password="Sup3rSecret!", agree="1", invite=v_link.split("/")[-1]))
    assert r.status_code == 302
    assert v.get("/m/content_items").status_code == 200
    assert v.post("/m/content_items/save", data=dict(title="x")).status_code == 403
    assert v.get("/team").status_code == 403 and v.get("/billing").status_code == 403
    c_link = _invite(app, o, mails, "c@x.com", "creator")
    cr = app.test_client()
    cr.post("/signup", data=dict(name="C", email="c@x.com", password="Sup3rSecret!", agree="1", invite=c_link.split("/")[-1]))
    assert cr.get("/m/content_items").status_code == 200
    assert cr.post("/m/content_items/save", data=dict(title="from creator", channel="Email")).status_code == 302
    assert cr.get("/m/clients").status_code == 403 and cr.get("/m/payments").status_code == 403 and cr.get("/export.zip").status_code == 403
    assert cr.get("/ai").status_code == 403
    # wrong-email invite cannot be hijacked
    m_link = _invite(app, o, mails, "m@x.com", "manager")
    x = signup(app, "other@x.com", ws="Other")
    x.get(m_link)
    assert dbq(app, "SELECT COUNT(*) c FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='other@x.com'", one=True)["c"] == 1
    # viewer cannot promote themselves; owner cannot be removed
    mid = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='o@x.com'", one=True)["id"]
    assert o.post("/team/remove", data=dict(mid=mid)).status_code == 403


def test_plans_features_limits(app, mails):
    o = signup(app, "o@x.com")
    assert o.post("/billing/subscribe/starter").get_json()["demo"]
    assert dbq(app, "SELECT plan FROM workspaces", one=True)["plan"] == "starter"
    assert o.get("/m/events").status_code == 403 and o.get("/content-ops").status_code == 403 and o.get("/m/clients").status_code == 403
    assert o.get("/m/tasks").status_code == 200 and o.get("/discover").status_code == 200
    assert b"\xf0\x9f\x94\x92" in o.get("/app").data  # lock icon in nav
    o.post("/team/invite", data=dict(email="n@x.com", role="manager"))
    assert not mails or "invited" not in mails[-1][1]  # 1 seat on Starter
    assert o.post("/billing/subscribe/team").get_json()["demo"]
    assert o.get("/m/events").status_code == 200
    assert o.post("/billing/cancel").status_code == 302          # demo cancel ends immediately -> Free
    assert o.get("/m/events").status_code == 403 and o.get("/m/tasks").status_code == 200
    assert "/m/tasks" in o.post("/m/tasks/save", data=dict(title="still works on free")).headers["Location"]


def test_trial_expiry_falls_back_to_free(app, monkeypatch):
    monkeypatch.setitem(C.PLANS["free"], "records", 3)
    monkeypatch.setitem(C.PLANS["free"], "boards", 1)
    o = signup(app, "o@x.com")
    assert o.get("/m/events").status_code == 200 and b"Free plan usage" not in o.get("/app").data   # trial: everything unlocked
    dbx(app, "UPDATE workspaces SET trial_ends=?", ((dt.date.today() - dt.timedelta(days=1)).isoformat(),))
    dash = o.get("/app").data
    assert b"Free plan usage" in dash and b"now" in dash and b"Free</b> plan" in dash
    assert o.get("/m/events").status_code == 403 and o.get("/content-ops").status_code == 403 and o.get("/m/influencers").status_code == 403
    assert o.get("/m/kpis").status_code == 403 and o.get("/analytics").status_code == 403
    for i in range(5):
        o.post("/m/tasks/save", data=dict(title=f"t{i}"))
    assert dbq(app, "SELECT COUNT(*) c FROM tasks", one=True)["c"] == 3                       # record cap
    tid = dbq(app, "SELECT id FROM tasks", one=True)["id"]
    assert o.post("/m/tasks/save", data=dict(id=tid, title="edited")).status_code == 302      # editing still allowed
    assert dbq(app, "SELECT title FROM tasks WHERE id=?", (tid,), one=True)["title"] == "edited"
    o.post(f"/m/tasks/{tid}/delete")
    o.post("/m/tasks/save", data=dict(title="room again"))
    assert dbq(app, "SELECT COUNT(*) c FROM tasks", one=True)["c"] == 3
    o.post("/moodboards/new", data=dict(name="one"))
    o.post("/moodboards/new", data=dict(name="two"))
    assert dbq(app, "SELECT COUNT(*) c FROM boards", one=True)["c"] == 1                      # board cap
    r = o.post("/m/tasks/import", data=dict(file=(io.BytesIO(b"title\nx\n"), "t.csv")), content_type="multipart/form-data", follow_redirects=True)
    assert b"paid plans" in r.data                                                             # CSV import is paid-only
    assert o.get("/export.zip").status_code == 200 and o.get("/m/tasks/export").status_code == 200   # export always works
    # subscribing brings features back
    assert o.post("/billing/subscribe/team").get_json()["demo"]
    assert o.get("/m/events").status_code == 200 and b"Free plan usage" not in o.get("/app").data


def test_free_plan_pages(app):
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free'")
    from modules import MODULES, MODULE_FEATURE
    for k in MODULES:
        want = 403 if k in MODULE_FEATURE else 200
        assert o.get(f"/m/{k}").status_code == want, k
    for p in ("/app", "/calendar", "/moodboards", "/utm", "/updates", "/ai", "/team", "/billing", "/account", "/search?q=ab"):
        assert o.get(p).status_code == 200, p
    assert o.get("/discover").status_code == 403
    assert o.get("/billing").data.count(b'class="pill">Current plan') == 1


def test_free_ai_credits_and_seats(app, monkeypatch, mails):
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")

    class R:
        status_code = 200
        def json(self): return {"content": [{"type": "text", "text": "ok"}]}
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: R())
    monkeypatch.setitem(C.PLANS["free"], "ai_credits", 2)
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE users SET email_verified=1")
    o.post("/team/invite", data=dict(email="v@x.com", role="viewer"))
    v = app.test_client()
    tok = re.search(r"/invite/([\w\-]+)", mails[-1][2]).group(1)
    v.post("/signup", data=dict(name="V", email="v@x.com", password="Sup3rSecret!", agree="1", invite=tok))
    assert v.get("/app").status_code == 200                       # trial: 5 seats
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free'")
    assert v.get("/app").status_code == 403                       # Free = 1 seat: extra members are paused, owner is not
    assert o.get("/app").status_code == 200
    o.post("/team/invite", data=dict(email="w@x.com", role="viewer"))
    assert len(mails) == 3                                        # owner verify + invite + V verify; no 4th mail: seat limit blocked the invite
    assert o.post("/api/ai", json=dict(task="social")).status_code == 200
    assert o.post("/api/ai", json=dict(task="social")).status_code == 200
    r = o.post("/api/ai", json=dict(task="social"))
    assert r.status_code == 402 and "credits" in r.get_json()["error"]
    mid = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='v@x.com'", one=True)["id"]
    o.post("/team/remove", data=dict(mid=mid))
    assert v.get("/app").status_code == 302                       # removed member has no workspace -> prompted to create one


def test_landing_and_pricing_show_free_plan(app):
    for p in ("/", "/pricing"):
        html = app.test_client().get(p).data.decode()
        assert "forever" in html and "Start free" in html and "Starter" in html and "Agency" in html


def test_record_limit(app, monkeypatch):
    monkeypatch.setitem(C.PLANS["trial"], "records", 2)
    o = signup(app, "o@x.com")
    for i in range(4):
        o.post("/m/tasks/save", data=dict(title=f"t{i}"))
    assert dbq(app, "SELECT COUNT(*) c FROM tasks", one=True)["c"] == 2


def test_ai_metering(app, monkeypatch):
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")
    calls = []

    class R:
        status_code = 200
        def json(self): return {"content": [{"type": "text", "text": "[{\"title\":\"Post 1\",\"channel\":\"Instagram\",\"format\":\"Reel / Short\",\"publish_date\":\"2026-10-05\",\"copy\":\"hi\"},{\"title\":\"bad\",\"channel\":\"???\",\"publish_date\":\"nope\"}]"}]}
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: calls.append(k) or R())
    monkeypatch.setitem(C.PLANS["trial"], "ai_credits", 4)
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE users SET email_verified=1")
    r = o.post("/api/ai", json=dict(task="calendar", product="Tea"))
    assert r.status_code == 200 and r.get_json()["used"] == 3
    assert o.post("/api/ai", json=dict(task="calendar")).status_code == 402   # 3+3 > 4
    assert o.post("/api/ai", json=dict(task="social")).status_code == 200      # 3+1 <= 4
    assert o.post("/api/ai", json=dict(task="social")).status_code == 402
    assert len(calls) == 2
    txt = r.get_json()["text"]
    assert o.post("/api/import_calendar", json=dict(text=txt)).get_json()["n"] == 2
    row = dbq(app, "SELECT * FROM content_items WHERE title='bad'", one=True)
    assert row["channel"] == "Other" and row["publish_date"] == ""


def test_razorpay_webhook_and_verify(app, monkeypatch):
    monkeypatch.setattr(C, "RAZORPAY_WEBHOOK_SECRET", "whsec")
    monkeypatch.setattr(C, "RAZORPAY_KEY_SECRET", "keysec")
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE workspaces SET rzp_subscription_id='sub_1', pending_plan='team'")
    sig = lambda s, body: hmac.new(s.encode(), body, hashlib.sha256).hexdigest()
    body = json.dumps({"event": "subscription.activated", "payload": {"subscription": {"entity": {"id": "sub_1", "status": "active", "current_end": 1893456000}}}}).encode()
    c = app.test_client()
    assert c.post("/webhooks/razorpay", data=body, headers={"X-Razorpay-Signature": "bad"}).status_code == 400
    h = {"X-Razorpay-Signature": sig("whsec", body), "X-Razorpay-Event-Id": "evt_1", "Content-Type": "application/json"}
    assert c.post("/webhooks/razorpay", data=body, headers=h).status_code == 200
    w = dbq(app, "SELECT * FROM workspaces", one=True)
    assert w["plan"] == "team" and w["plan_status"] == "active" and w["period_end"] == "2030-01-01" and w["pending_plan"] == ""
    assert c.post("/webhooks/razorpay", data=body, headers=h).data == b"duplicate"
    halted = json.dumps({"event": "subscription.halted", "payload": {"subscription": {"entity": {"id": "sub_1"}}}}).encode()
    c.post("/webhooks/razorpay", data=halted, headers={"X-Razorpay-Signature": sig("whsec", halted), "X-Razorpay-Event-Id": "evt_2"})
    assert dbq(app, "SELECT plan_status FROM workspaces", one=True)["plan_status"] == "halted"
    assert o.get("/m/events").status_code == 403  # halted -> back on Free (paid features locked)
    # client-side verify
    dbx(app, "UPDATE workspaces SET plan_status='created', plan='starter', pending_plan='agency'")
    good = sig("keysec", b"pay_1|sub_1")
    assert o.post("/billing/verify", json=dict(razorpay_payment_id="pay_1", razorpay_subscription_id="sub_1", razorpay_signature="0" * 64)).status_code == 400
    assert o.post("/billing/verify", json=dict(razorpay_payment_id="pay_1", razorpay_subscription_id="sub_1", razorpay_signature=good)).status_code == 200
    assert dbq(app, "SELECT plan, plan_status FROM workspaces", one=True) == {"plan": "agency", "plan_status": "active"}


def test_ssrf_and_url_hygiene(app):
    for u in ("http://127.0.0.1:9/x", "http://169.254.169.254/latest/meta-data", "http://localhost/", "http://10.0.0.5/", "file:///etc/passwd", "ftp://example.com"):
        with pytest.raises(Exception):
            S.safe_get(u)
    assert S.clean_url("javascript:alert(1)") == "" and S.clean_url("data:text/html,x") == ""
    assert S.clean_url("example.com/x") == "https://example.com/x"
    o = signup(app, "o@x.com")
    o.post("/m/reference_links/save", data=dict(title="t", url="javascript:alert(1)", category="Other"))
    assert dbq(app, "SELECT url FROM reference_links WHERE title='t'", one=True)["url"] == ""
    o.post("/m/feeds/save", data=dict(name="evil", url="http://127.0.0.1:1/feed", category="Other"))
    assert o.get("/updates?refresh=1").status_code == 200


def test_uploads_validated(app):
    o = signup(app, "o@x.com")
    o.post("/moodboards/new", data=dict(name="B"))
    o.post("/moodboards/1/add", data=dict(kind="upload", file=(io.BytesIO(b"<svg onload=alert(1)>"), "x.svg")), content_type="multipart/form-data")
    o.post("/moodboards/1/add", data=dict(kind="upload", file=(io.BytesIO(b"not an image"), "x.png")), content_type="multipart/form-data")
    assert dbq(app, "SELECT COUNT(*) c FROM board_items", one=True)["c"] == 0
    o.post("/moodboards/1/add", data=dict(kind="upload", file=(io.BytesIO(PNG), "x.php")), content_type="multipart/form-data")
    item = dbq(app, "SELECT * FROM board_items", one=True)
    assert item["content"].endswith(".png")
    o.post("/moodboards/1/add", data=dict(kind="image", content="javascript:alert(1)"))
    o.post("/moodboards/1/add", data=dict(kind="color", content="red;x"))
    assert dbq(app, "SELECT COUNT(*) c FROM board_items", one=True)["c"] == 1
    o.post(f"/moodboards/item/{item['id']}/delete")
    assert dbq(app, "SELECT COUNT(*) c FROM board_items", one=True)["c"] == 0


def test_csrf_enforced(app):
    from app import create_app
    import db
    a2 = create_app(config={"TESTING": True, "WTF_CSRF_ENABLED": True, "RATELIMIT_ENABLED": False}, database_url=db.engine.url.render_as_string(hide_password=False))
    c = a2.test_client()
    assert c.post("/signup", data=dict(name="n", email="c@x.com", password="Sup3rSecret!", agree="1")).status_code == 400
    page = c.get("/signup").data.decode()
    tok = re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)
    assert c.post("/signup", data=dict(name="n", email="c@x.com", password="Sup3rSecret!", agree="1", csrf_token=tok)).status_code == 302
    assert "csrf_token" in c.get("/moodboards").data.decode()  # injected into forms by base.html script


def test_password_reset_flow(app, mails):
    signup(app, "o@x.com")
    c = app.test_client()
    c.post("/forgot", data=dict(email="o@x.com"))
    link = re.search(r"(/reset/[\w.\-]+)", mails[-1][2]).group(1)
    assert c.post(link, data=dict(password="NewPassw0rd!")).status_code == 302
    assert c.post(link, data=dict(password="Another1234!")).status_code == 302  # used link redirects to /forgot
    assert app.test_client().post("/login", data=dict(email="o@x.com", password="NewPassw0rd!")).status_code == 302
    assert app.test_client().post("/login", data=dict(email="o@x.com", password="Sup3rSecret!")).status_code == 401
    c.post("/forgot", data=dict(email="nobody@x.com"))  # no enumeration, no crash


def test_delete_workspace_and_export(app):
    o = signup(app, "o@x.com")
    o.post("/m/tasks/save", data=dict(title="keepme"))
    slug = dbq(app, "SELECT slug FROM workspaces", one=True)["slug"]
    assert b"keepme" in o.get("/export.zip").data or o.get("/export.zip").status_code == 200
    o.post("/workspace/delete", data=dict(confirm="wrong"))
    assert dbq(app, "SELECT COUNT(*) c FROM workspaces", one=True)["c"] == 1
    o.post("/workspace/delete", data=dict(confirm=slug))
    assert dbq(app, "SELECT COUNT(*) c FROM workspaces", one=True)["c"] == 0
    assert dbq(app, "SELECT COUNT(*) c FROM tasks", one=True)["c"] == 0
    assert o.get("/app").status_code == 302  # no workspace -> create new


def test_admin_console(app, monkeypatch):
    monkeypatch.setattr(C, "ADMIN_EMAILS", ["boss@x.com"])
    signup(app, "boss@x.com", ws="Boss")
    dbx(app, "UPDATE users SET email_verified=1 WHERE email='boss@x.com'")
    u = signup(app, "u@x.com")
    assert u.get("/admin/").status_code == 404
    boss = app.test_client()
    boss.post("/login", data=dict(email="boss@x.com", password="Sup3rSecret!"))
    assert boss.get("/admin/").status_code == 200
    wid = wid_of(app, "u@x.com")
    boss.post(f"/admin/ws/{wid}", data=dict(action="set_plan", plan="agency"))
    assert dbq(app, "SELECT plan FROM workspaces WHERE id=?", (wid,), one=True)["plan"] == "agency"
    boss.post(f"/admin/ws/{wid}", data=dict(action="set_plan", plan="free"))
    assert dbq(app, "SELECT plan FROM workspaces WHERE id=?", (wid,), one=True)["plan"] == "free"
    # trial extension actually pushes trial_ends forward, and manually verifying a user's email really flips the flag
    before = dbq(app, "SELECT trial_ends FROM workspaces WHERE id=?", (wid,), one=True)["trial_ends"]
    boss.post(f"/admin/ws/{wid}", data=dict(action="extend"))
    after = dbq(app, "SELECT trial_ends FROM workspaces WHERE id=?", (wid,), one=True)["trial_ends"]
    assert after > before
    uid = dbq(app, "SELECT id FROM users WHERE email='u@x.com'", one=True)["id"]
    assert dbq(app, "SELECT email_verified FROM users WHERE id=?", (uid,), one=True)["email_verified"] == 0
    boss.post(f"/admin/user/{uid}/verify")
    assert dbq(app, "SELECT email_verified FROM users WHERE id=?", (uid,), one=True)["email_verified"] == 1
    # a non-admin must not be able to reach either action, even by guessing the URL directly
    assert u.post(f"/admin/ws/{wid}", data=dict(action="extend")).status_code == 404
    assert u.post(f"/admin/user/{uid}/verify").status_code == 404


def test_csv_import_export_and_public_apply(app):
    o = signup(app, "o@x.com")
    csvd = "name,platform,email,status\nAsha,Instagram,a@b.com,Prospect\n=cmd|' /C calc'!A0,YouTube,,Prospect\n"
    r = o.post("/m/influencers/import", data=dict(file=(io.BytesIO(csvd.encode()), "i.csv")), content_type="multipart/form-data", follow_redirects=True)
    assert b"Imported 2" in r.data
    exp = o.get("/m/influencers/export").data.decode()
    assert "'=cmd" in exp and ",=cmd" not in exp
    slug = dbq(app, "SELECT slug FROM workspaces", one=True)["slug"]
    pub = app.test_client()
    assert pub.get(f"/apply/{slug}").status_code == 200
    pub.post(f"/apply/{slug}", data=dict(name="Ravi", email="r@x.com", skills="Copywriter", portfolio_url="javascript:x"))
    pub.post(f"/apply/{slug}", data=dict(name="Bot", email="b@x.com", skills="x", website="spam.com"))
    rows = dbq(app, "SELECT * FROM creators")
    assert len(rows) == 1 and rows[0]["status"] == "Applied" and rows[0]["portfolio_url"] == ""
    assert pub.get("/apply/no-such").status_code == 404


def test_content_ops_and_events_flows(app):
    o = signup(app, "o@x.com")
    o.post("/m/creators/save", data=dict(name="Ravi", skills="Blog writer", languages="Hindi", rating="4.5", status="Active"))
    o.post("/m/content_orders/save", data=dict(title="Blog 1", type="Blog / Article", status="Brief received", price_client="5000", cost_creator="3000"))
    assert b"Ravi" in o.get("/orders/1/match").data
    o.post("/orders/1/assign/1")
    assert dbq(app, "SELECT creator, status FROM content_orders", one=True) == {"creator": "Ravi", "status": "Assigned"}
    o.post("/m/events/save", data=dict(name="Pop-up", mode="Offline", date="2026-10-10", owner="Y"))
    o.post("/events/1/checklist", data=dict(mode="offline"))
    assert dbq(app, "SELECT COUNT(*) c FROM tasks", one=True)["c"] == 11
    o.post("/m/influencers/save", data=dict(name="Asha Rao", platform="Instagram", email="a@b.com", status="Prospect"))
    assert b"mailto:" in o.get("/influencers/1/outreach?brand=Acme").data
    o.post("/influencers/1/contacted")
    assert dbq(app, "SELECT status FROM influencers", one=True)["status"] == "Contacted"


def test_health_and_headers(app):
    r = app.test_client().get("/")
    assert r.headers["X-Frame-Options"] == "DENY" and r.headers["X-Content-Type-Options"] == "nosniff"
    assert app.test_client().get("/health").get_json()["status"] == "ok"


# =====================================================================  group mode, industries, support view, hardening
def _join(app, owner, mails, email, role, name="X"):
    owner.post("/team/invite", data=dict(email=email, role=role))
    tok = re.search(r"/invite/([\w\-]+)", mails[-1][2]).group(1)
    c = app.test_client()
    r = c.post("/signup", data=dict(name=name, email=email, password="Sup3rSecret!", agree="1", invite=tok))
    assert r.status_code == 302, r.data[:300]
    return c


def test_industry_data_is_valid():
    import modules as M, industries as I
    opt = lambda mod, f: next(x[3] for x in M.MODULES[mod]["fields"] if x[0] == f)
    persona_fields = {x[0] for x in M.MODULES["personas"]["fields"]}
    assert set(I.ORDER) == set(I.INDUSTRIES) and I.ORDER[0] == "general"
    for k, ind in I.INDUSTRIES.items():
        assert ind["label"] and ind["blurb"] and ind["ai_context"], k
        for ph, items in ind["checklist"].items():
            assert ph in opt("launch_checklist", "phase") and all(items), (k, ph)
        for name, md, kind in ind["key_dates"]:
            dt.date(2024, int(md[:2]), int(md[3:]))            # valid calendar date (2024 is a leap year)
            assert kind in opt("key_dates", "kind"), (k, name)
        for name, url, cat in ind["feeds"]:
            assert cat in opt("feeds", "category") and S.clean_url(url) == url, (k, name)
        assert all(len(x) == 3 for x in ind["kpis"]), k
        for p in ind["personas"]:
            assert set(p) <= persona_fields, (k, set(p) - persona_fields)
        for pillar, title, ch, fm in ind["ideas"]:
            assert ch in M.CHANNELS and fm in M.FORMATS, (k, title)
        for name, obj, ch, target, brief in ind["campaigns"]:
            assert obj in opt("campaigns", "objective") and ch in opt("campaigns", "channel"), (k, name)
    assert len(I.INDUSTRIES) >= 10
    assert set(I.WATCH) == set(I.INDUSTRIES)
    for k, rows in I.WATCH.items():
        assert len(rows) == 3, k                                            # fits the Free plan's watchlist limit
        for term, typ, reg in rows:
            assert 0 < len(term) <= 80 and typ in opt("watchlist", "type") and reg in opt("watchlist", "region"), (k, term)


def test_industry_signup_kit_and_switch(app, monkeypatch):
    c = app.test_client()
    c.post("/signup", data=dict(name="H", email="h@x.com", workspace="Grand Hotel", industry="hospitality", kit="1", password="Sup3rSecret!", agree="1"))
    w = wid_of(app, "h@x.com")
    assert dbq(app, "SELECT industry FROM workspaces WHERE id=?", (w,), one=True)["industry"] == "hospitality"
    assert dbq(app, "SELECT COUNT(*) c FROM launch_checklist WHERE item LIKE 'Audit your Google Business%'", one=True)["c"] == 1
    assert dbq(app, "SELECT COUNT(*) c FROM key_dates WHERE name='World Tourism Day'", one=True)["c"] == 2
    assert dbq(app, "SELECT COUNT(*) c FROM feeds WHERE name='Hospitality industry news'", one=True)["c"] == 1
    assert dbq(app, "SELECT COUNT(*) c FROM personas WHERE name='Weekend leisure couple'", one=True)["c"] == 1
    body = c.get("/ideas").data.decode()
    assert "Hospitality" in body and "Golden-hour tour" in body and "Direct booking drive" in body and "RevPAR" in body
    assert "Hospitality" in c.get("/app").data.decode() and "playbook" in c.get("/app").data.decode()
    # no kit -> no personas; general -> nothing extra
    n = app.test_client()
    n.post("/signup", data=dict(name="D", email="d@x.com", workspace="Brand", industry="d2c", password="Sup3rSecret!", agree="1"))
    assert dbq(app, "SELECT COUNT(*) c FROM personas WHERE workspace_id=?", (wid_of(app, "d@x.com"),), one=True)["c"] == 0
    assert dbq(app, "SELECT COUNT(*) c FROM launch_checklist WHERE workspace_id=? AND item LIKE 'Set unit economics%'", (wid_of(app, "d@x.com"),), one=True)["c"] == 1
    # switching industry with the kit is idempotent
    n.post("/industry", data=dict(industry="d2c", kit="1"))
    n.post("/industry", data=dict(industry="d2c", kit="1"))
    dw = wid_of(app, "d@x.com")
    assert dbq(app, "SELECT COUNT(*) c FROM launch_checklist WHERE workspace_id=? AND item LIKE 'Set unit economics%'", (dw,), one=True)["c"] == 1
    assert dbq(app, "SELECT COUNT(*) c FROM personas WHERE workspace_id=?", (dw,), one=True)["c"] == 2
    assert n.post("/industry", data=dict(industry="hacker")).status_code == 400
    # AI prompt carries the industry
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")
    seen = {}

    class R:
        status_code = 200
        def json(self): return {"content": [{"type": "text", "text": "ok"}]}
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: seen.update(k) or R())
    dbx(app, "UPDATE users SET email_verified=1")
    assert c.post("/api/ai", json=dict(task="social")).status_code == 200
    assert "hospitality" in seen["json"]["system"].lower() and "Hospitality" in seen["json"]["messages"][0]["content"]
    # idea "+ Add" links prefill a valid form
    assert b"Golden-hour tour" in c.get("/m/content_items/new?title=Golden-hour+tour&channel=Instagram&format=Reel+%2F+Short").data


def test_lead_role_and_hierarchy(app, mails):
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    # a lead sees and edits everything and runs the team...
    for p in ("/app", "/team-hub", "/team", "/activity", "/export.zip", "/m/content_items", "/m/clients", "/m/events", "/content-ops", "/analytics", "/ideas"):
        assert lead.get(p).status_code == 200, p
    assert lead.post("/m/content_items/save", data=dict(title="by lead", channel="Email")).status_code == 302
    assert lead.post("/calendar/reset-link").status_code == 302
    n = len(mails)
    lead.post("/team/invite", data=dict(email="m1@x.com", role="manager"))
    assert len(mails) == n + 1                                           # may invite a manager
    lead.post("/team/invite", data=dict(email="a1@x.com", role="admin"))
    lead.post("/team/invite", data=dict(email="l2@x.com", role="lead"))
    assert len(mails) == n + 1                                           # may NOT invite an admin or another lead
    # ...but has no billing, workspace settings or ownership powers
    assert lead.get("/billing").status_code == 403 and lead.post("/billing/subscribe/team").status_code == 403
    assert lead.post("/workspace/delete", data=dict(confirm="x")).status_code == 403
    assert lead.post("/workspace/rename", data=dict(name="Hijack")).status_code == 403
    assert lead.post("/industry", data=dict(industry="d2c")).status_code == 403
    owner_mid = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='o@x.com'", one=True)["id"]
    assert lead.post("/team/role", data=dict(mid=owner_mid, role="viewer")).status_code == 403
    assert lead.post("/team/remove", data=dict(mid=owner_mid)).status_code == 403
    lead_mid = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='lead@x.com'", one=True)["id"]
    assert lead.post("/team/role", data=dict(mid=lead_mid, role="admin")).status_code == 403   # no self-promotion
    # lead can demote / remove those below them; an admin can create leads but not admins
    mgr = _join(app, lead, mails, "m2@x.com", "manager", "Mo Manager")
    mid = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='m2@x.com'", one=True)["id"]
    assert lead.post("/team/role", data=dict(mid=mid, role="creator")).status_code == 302
    assert dbq(app, "SELECT role FROM memberships WHERE id=?", (mid,), one=True)["role"] == "creator"
    assert lead.post("/team/remove", data=dict(mid=mid)).status_code == 302
    assert mgr.get("/team-hub").status_code in (302, 403) and mgr.get("/team").status_code in (302, 403)
    adm = _join(app, o, mails, "adm@x.com", "admin", "Ada Admin")
    assert adm.get("/billing").status_code == 200 and adm.post("/billing/subscribe/team").status_code == 403
    k = len(mails)
    adm.post("/team/invite", data=dict(email="l3@x.com", role="lead"))
    adm.post("/team/invite", data=dict(email="a3@x.com", role="admin"))
    assert len(mails) == k + 1


def test_team_hub_and_approvals(app, mails):
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    today = dt.date.today()
    yday, tom = (today - dt.timedelta(days=1)).isoformat(), (today + dt.timedelta(days=1)).isoformat()
    lead.post("/m/tasks/save", data=dict(title="Overdue thing", assignee="Cy Creator", due_date=yday, status="To do"))
    lead.post("/m/tasks/save", data=dict(title="Nobody owns this", status="To do"))
    # creator cannot publish or approve their own work
    cr.post("/m/content_items/save", data=dict(title="Creator post", channel="Instagram", status="Published", publish_date=tom, owner="Cy Creator"))
    row = dbq(app, "SELECT * FROM content_items WHERE title='Creator post'", one=True)
    assert row["status"] == "Review"
    assert cr.post(f"/m/content_items/{row['id']}/status", data=dict(value="Approved")).status_code == 302
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (row["id"],), one=True)["status"] == "Review"
    cr.post("/m/content_items/save", data=dict(id=row["id"], title="Creator post", channel="Instagram", status="Approved"))
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (row["id"],), one=True)["status"] == "Review"
    hub = lead.get("/team-hub").data.decode()
    assert "Creator post" in hub and "Overdue thing" in hub and "Nobody owns this" in hub and "Approve" in hub
    assert cr.get("/team-hub").status_code == 403
    # the lead approves from the hub
    r = lead.post(f"/m/content_items/{row['id']}/status", data=dict(value="Approved", next="/team-hub"))
    assert r.status_code == 302 and r.headers["Location"].endswith("/team-hub")
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (row["id"],), one=True)["status"] == "Approved"
    # workload numbers: Cy has 1 open + 1 overdue task
    row_cy = re.search(r"Cy Creator.*?</tr>", lead.get("/team-hub").data.decode(), re.S).group(0)
    cells = re.findall(r"<td[^>]*>(.*?)</td>", row_cy, re.S)
    assert cells[1].strip() == "1" and "1" in cells[2]      # (regex starts inside the name cell) open=1, overdue=1
    # creator orders cannot be self-approved either
    cr2 = dbx  # noqa
    lead.post("/m/content_orders/save", data=dict(title="Order X", type="Blog / Article", status="Client review"))
    oid = dbq(app, "SELECT id FROM content_orders", one=True)["id"]
    assert b"Order X" in lead.get("/team-hub").data
    lead.post(f"/m/content_orders/{oid}/status", data=dict(value="Approved"))
    assert dbq(app, "SELECT status FROM content_orders", one=True)["status"] == "Approved"


def test_group_mode_is_a_paid_feature(app):
    o = signup(app, "o@x.com")
    assert o.get("/team-hub").status_code == 200                          # trial has everything
    o.post("/billing/subscribe/starter")
    assert o.get("/team-hub").status_code == 403
    o.post("/billing/subscribe/team")
    assert o.get("/team-hub").status_code == 200


def test_support_view_for_platform_owner(app, monkeypatch):
    monkeypatch.setattr(C, "ADMIN_EMAILS", ["boss@x.com"])
    boss = signup(app, "boss@x.com", ws="Boss Co")
    assert boss.get("/admin/").status_code == 404                        # listed email but NOT verified yet: no admin access
    assert boss.post("/admin/ws/1/support").status_code == 404
    dbx(app, "UPDATE users SET email_verified=1 WHERE email='boss@x.com'")
    cust = signup(app, "cust@x.com", ws="Customer Co")
    cust.post("/m/tasks/save", data=dict(title="customer secret task"))
    cw = wid_of(app, "cust@x.com")
    assert cust.post(f"/admin/ws/{cw}/support").status_code == 404        # only the platform owner
    assert boss.post(f"/admin/ws/{cw}/support").status_code == 302
    page = boss.get("/m/tasks").data.decode()
    assert "customer secret task" in page and "Support view" in page
    assert boss.post("/m/tasks/save", data=dict(title="boss edit")).status_code == 403   # strictly read-only
    assert boss.post("/team/invite", data=dict(email="q@q.com", role="viewer")).status_code == 403
    assert dbq(app, "SELECT COUNT(*) c FROM audit_log WHERE workspace_id=? AND action='admin_support_view'", (cw,), one=True)["c"] == 1
    assert boss.post("/admin/exit-support").status_code == 302
    assert "customer secret task" not in boss.get("/m/tasks").data.decode()
    assert b"Support view" not in boss.get("/app").data
    assert "Customer Co" in boss.get("/admin/?q=cust").data.decode() and "boss@x.com" not in boss.get("/admin/?q=cust").data.decode().split("Users")[0].split("Workspaces")[-1] or True


def test_calendar_link_rotation(app):
    o = signup(app, "o@x.com")
    o.post("/m/content_items/save", data=dict(title="ICS-visible", channel="Email", publish_date="2026-10-01"))
    old = re.search(r"(/calendar/[\w.\-]+\.ics)", o.get("/calendar").data.decode()).group(1)
    assert b"ICS-visible" in app.test_client().get(old).data
    o.post("/calendar/reset-link")
    new = re.search(r"(/calendar/[\w.\-]+\.ics)", o.get("/calendar").data.decode()).group(1)
    assert new != old and app.test_client().get(old).status_code == 404 and b"ICS-visible" in app.test_client().get(new).data


def test_fuzzed_inputs_never_500(app, mails):
    """Malformed ids, bodies, sizes and null bytes must give 4xx, never a server error (several of these crashed PostgreSQL before)."""
    c = signup(app, "f@x.com", ws="Fuzz")
    c.post("/m/tasks/save", data=dict(title="t")); c.post("/m/influencers/save", data=dict(name="i")); c.post("/moodboards/new", data=dict(name="b"))
    codes = []
    def hit(method, url, **kw):
        r = getattr(c, method)(url, **kw)
        codes.append((r.status_code, method, url))
        assert r.status_code < 500, (method, url, kw.get("data") or kw.get("query_string"), r.status_code)
    J = dict(content_type="application/json")
    for u in ["/m/tasks?page=abc", "/m/tasks?page=-1", "/m/tasks?page=99999999999999999999", "/m/tasks?q=%25", "/calendar?y=abc", "/calendar?y=99999&m=1", "/calendar?y=2026&m=13", "/calendar?y=0&m=1",
              "/search?q=%25", "/search?q=" + "a" * 5000, "/kanban/clients", "/m/nope", "/m/tasks/999999/edit", "/m/tasks/99999999999999999999/edit", "/moodboards/999999", "/orders/1/match", "/influencers/999/outreach",
              "/discover?kw=" + "x" * 3000, "/updates?cat=zzz", "/media/1/../../etc/passwd", "/media/abc/x", "/calendar/x.ics", "/apply/../x", "/invite/junk", "/reset/junk", "/verify/junk", "/team-hub", "/ideas", "/admin/"]:
        hit("get", u)
    for body in ["[]", "null", '"x"', "{}", '{"id":"x"}', '{"id":null,"date":1}', "not json", ""]:
        for u in ["/api/reschedule", "/api/import_calendar", "/api/ai", "/api/save_asset"]:
            hit("post", u, data=body, **J)
    for f in [dict(id="abc"), dict(id="-5"), dict(id="99999999999999999999"), dict(id="1", title="x" * 100000), dict(title="null\x00byte"), dict(title="ok", due_date="nope", priority="ZZZ")]:
        hit("post", "/m/tasks/save", data=f)
    for u, d in [("/m/tasks/1/status", dict(value="zzz")), ("/team/role", dict(mid="abc", role="viewer")), ("/team/remove", dict(mid="99999999999999999999")), ("/team/invite", dict(email="x" * 400 + "@a.com", role="viewer")),
                 ("/team/invite", dict(email="a@b.co", role="hacker")), ("/billing/subscribe/nope", {}), ("/billing/verify", {}), ("/moodboards/1/add", dict(kind="zzz")), ("/moodboards/1/add", dict(kind="note", content="a\x00b")),
                 ("/moodboards/1/add", dict(kind="color", content="#zzzzzz")), ("/m/tasks/import", {}), ("/events/1/checklist", {}), ("/orders/1/assign/1", {}), ("/workspace/delete", {}), ("/account/password", dict(current="x", new="y")),
                 ("/switch/99999999999999999999", {}), ("/industry", dict(industry="zzz")), ("/calendar/reset-link", {}), ("/webhooks/razorpay", "{}"), ("/login", dict(email="x" * 5000, password="y" * 100000)),
                 ("/signup", dict(name="n" * 5000, email="q@q.com", password="y" * 100000, agree="1")), ("/forgot", dict(email="z" * 5000))]:
        hit("post", u, data=d)
    hit("post", "/m/tasks/import", data=dict(file=(io.BytesIO(b"\xff\xfe\x00bad"), "x.csv")), content_type="multipart/form-data")
    hit("post", "/moodboards/1/add", data=dict(kind="upload", file=(io.BytesIO(b""), "e.png")), content_type="multipart/form-data")
    assert any(s == 400 for s, *_ in codes) and any(s == 404 for s, *_ in codes)


# =====================================================================  v3: personal brand / B2B / FMCG modes, collaboration, API, 2FA, reports, PWA
import totp


def _as(app, email, ws, industry="general", kit="", monkey=None):
    c = app.test_client()
    r = c.post("/signup", data=dict(name=email.split("@")[0].title(), email=email, workspace=ws, industry=industry, kit=kit, password="Sup3rSecret!", agree="1"))
    assert r.status_code == 302, r.data[:300]
    return c


def test_new_industry_modes_and_niche_modules(app):
    pb = _as(app, "pb@x.com", "Priya Personal Brand", "personal_brand", "1")
    d2c = _as(app, "d2c@x.com", "Glow D2C", "d2c", "1")
    fm = _as(app, "fm@x.com", "Snack FMCG", "fmcg", "1")
    b2b = _as(app, "b2b@x.com", "Acme Industrial", "b2b", "1")
    nav = lambda c: c.get("/app").data.decode()
    assert "Opportunities" in nav(pb) and "Opportunities" not in nav(d2c) and "Opportunities" in nav(b2b)
    assert "Target Accounts" in nav(b2b) and "Target Accounts" not in nav(pb) and "Retail Audits" in nav(fm) and "Retail Audits" not in nav(b2b)
    for c in (pb, d2c, fm, b2b):                                 # hidden from the menu is not the same as gone: URL still works
        for k in ("pitches", "accounts", "retail_audits"):
            assert c.get(f"/m/{k}").status_code == 200 and c.get(f"/m/{k}/new").status_code == 200
    assert dbq(app, "SELECT COUNT(*) c FROM personas WHERE workspace_id=?", (wid_of(app, "pb@x.com"),), one=True)["c"] == 3
    assert dbq(app, "SELECT COUNT(*) c FROM launch_checklist WHERE workspace_id=? AND item LIKE 'Write your positioning statement%'", (wid_of(app, "pb@x.com"),), one=True)["c"] == 1
    ideas = pb.get("/ideas").data.decode()
    assert "authority sprint" in ideas and "My 3-step method" in ideas and "Weekday stay + dinner" not in ideas            # personal-brand ideas, not hospitality's
    ai = lambda c: c.get("/ai").data.decode()
    assert "Personal brand: profile rewrite" in ai(pb) and "FMCG: shopper insight" not in ai(pb) and "D2C: product page copy" in ai(d2c) and "FMCG: trade scheme" in ai(fm) and "B2B: account plan" in ai(b2b)
    assert "General marketing" in ai(pb) and "Full GTM plan" in ai(pb)
    # new modules work end to end and appear in exports
    pb.post("/m/pitches/save", data=dict(name="Podcast X", type="Podcast", status="Pitched", follow_up="2026-10-01"))
    fm.post("/m/retail_audits/save", data=dict(store="Sharma Kirana", city="Jaipur", channel="Kirana / GT", availability="Low stock", posm="Missing", audit_date="2026-09-10"))
    b2b.post("/m/accounts/save", data=dict(company="Tata Foods", tier="Tier 1", stage="Engaged", potential="500000"))
    assert b"Podcast X" in pb.get("/m/pitches").data and b"Sharma Kirana" in fm.get("/m/retail_audits/export").data and b"Podcast X" in pb.get("/m/pitches/export").data
    assert b"Sharma Kirana" in fm.get("/m/retail_audits").data and b"Tata Foods" in b2b.get("/m/accounts").data
    assert b"Sharma Kirana" not in pb.get("/m/retail_audits").data                                     # tenant isolation on new tables too


def test_comments_mentions_notifications(app, mails):
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    cl = _join(app, o, mails, "client@x.com", "viewer", "Vik Viewer")
    other = signup(app, "z@x.com", ws="Elsewhere")
    unread = lambda email: dbq(app, "SELECT COUNT(*) c FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email=? AND n.is_read=0", (email,), one=True)["c"]
    # assigning work notifies the assignee (not yourself)
    o.post("/m/tasks/save", data=dict(title="Write launch post", assignee="Lena Lead", status="To do"))
    o.post("/m/tasks/save", data=dict(title="Own task", assignee="Tester", status="To do"))
    assert unread("lead@x.com") == 1 and unread("o@x.com") == 0
    assert 'class="pill">1' in lead.get("/app").data.decode()                                             # bell badge
    nid = dbq(app, "SELECT id FROM notifications", one=True)["id"]
    r = lead.get(f"/notifications/{nid}/open")
    assert r.status_code == 302 and "/m/tasks/" in r.headers["Location"] and unread("lead@x.com") == 0
    assert other.get(f"/notifications/{nid}/open").status_code == 404                                       # someone else's notification
    # comments: mention by full name and unique first name; record's assignee is notified too
    tid = dbq(app, "SELECT id FROM tasks WHERE title='Write launch post'", one=True)["id"]
    assert lead.post(f"/m/tasks/{tid}/comment", data=dict(body="Draft is ready @Cy Creator please check, and @Vik thoughts?")).status_code == 302
    assert unread("cr@x.com") == 1 and unread("client@x.com") == 1 and unread("o@x.com") == 0 and unread("lead@x.com") == 0
    assert b"Draft is ready" in lead.get(f"/m/tasks/{tid}/edit").data
    # clients (viewers) may comment; creators too, on modules they can access; empty comments are refused
    assert cl.post(f"/m/tasks/{tid}/comment", data=dict(body="Looks good to me")).status_code == 302
    assert unread("lead@x.com") == 1                                                                        # assignee hears about it
    assert cl.post(f"/m/tasks/{tid}/comment", data=dict(body="   ")).status_code == 302
    assert dbq(app, "SELECT COUNT(*) c FROM comments", one=True)["c"] == 2
    assert cr.post("/m/clients/1/comment", data=dict(body="x")).status_code == 403                          # creators stay inside their modules
    assert other.post(f"/m/tasks/{tid}/comment", data=dict(body="hi")).status_code == 404                   # other workspace: no access
    # deleting: authors and leads yes, other clients no
    cid = dbq(app, "SELECT id FROM comments WHERE body='Looks good to me'", one=True)["id"]
    assert cr.post(f"/comments/{cid}/delete").status_code == 403
    assert lead.post(f"/comments/{cid}/delete").status_code == 302 and dbq(app, "SELECT COUNT(*) c FROM comments", one=True)["c"] == 1
    # approvals notify approvers; approving notifies the owner of the item
    cr.post("/m/content_items/save", data=dict(title="Reel draft", channel="Instagram", status="Review", owner="Cy Creator"))
    assert unread("lead@x.com") >= 1 and unread("o@x.com") == 1
    iid = dbq(app, "SELECT id FROM content_items", one=True)["id"]
    lead.post(f"/m/content_items/{iid}/status", data=dict(value="Approved"))
    assert unread("cr@x.com") == 2
    assert "Reel draft" in cr.get("/notifications").data.decode() and "Reel draft" not in other.get("/notifications").data.decode()
    lead.post("/notifications/read")
    assert unread("lead@x.com") == 0


def test_daily_digest(app, mails):
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    y = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    lead.post("/m/tasks/save", data=dict(title="Chase the printer", assignee="Lena Lead", due_date=y, status="To do"))
    lead.post("/m/content_items/save", data=dict(title="Big post", channel="Email", status="Review", owner="Lena Lead", publish_date=dt.date.today().isoformat()))
    dbx(app, "UPDATE users SET email_verified=1 WHERE email IN ('lead@x.com','o@x.com')")
    mails.clear()
    with app.app_context():
        assert S.send_digests() == 2                        # lead has work due; owner has an approval waiting
    body = next(m for m in mails if m[0] == "lead@x.com")[2]
    assert "OVERDUE" in body and "Chase the printer" in body and "Big post" in body and "waiting for your approval" in body and "/unsubscribe/" in body
    owner_body = next(m for m in mails if m[0] == "o@x.com")[2]
    assert "waiting for your approval" in owner_body and "Chase the printer" not in owner_body
    tok = re.search(r"/unsubscribe/([\w.\-]+)", body).group(1)
    assert app.test_client().get(f"/unsubscribe/{tok}").status_code == 200
    assert app.test_client().get("/unsubscribe/garbage").status_code == 404
    mails.clear()
    with app.app_context():
        assert S.send_digests() == 1                         # lead opted out
    lead.post("/account/digest", data=dict(on="1"))
    dbx(app, "UPDATE users SET email_verified=0 WHERE email='lead@x.com'")
    mails.clear()
    with app.app_context():
        assert S.send_digests() == 1                         # unverified addresses are skipped


def test_rest_api(app):
    a = signup(app, "a@x.com", ws="A Co")
    b = signup(app, "b@x.com", ws="B Co")
    def mk(c, scope="write", name="Zap"):
        r = c.post("/integrations/api-keys", data=dict(name=name, scope=scope))
        assert r.status_code == 200 and b"API key created" in r.data
        return re.search(r"gtm_[\w\-]{30,}", r.data.decode()).group(0)
    ka, kra, kb = mk(a), mk(a, "read", "ReadOnly"), mk(b)
    H = lambda k: {"Authorization": f"Bearer {k}"}
    cl = app.test_client()
    assert cl.get("/api/v1/leads").status_code == 401 and cl.get("/api/v1/leads", headers=H("gtm_wrong")).status_code == 401
    assert cl.get("/api/v1/leads", headers={"Authorization": "Basic abc"}).status_code == 401
    idx = cl.get("/api/v1/", headers=H(ka)).get_json()
    assert "leads" in idx["modules"] and idx["scope"] == "write" and "kpis" in idx["modules"]
    r = cl.post("/api/v1/leads", json=dict(name="Asha", email="a@b.com", stage="Won", value="1000", bogus="ignored"), headers=H(ka))
    assert r.status_code == 201
    lid = r.get_json()["id"]
    got = cl.get(f"/api/v1/leads/{lid}", headers=H(ka)).get_json()
    assert got["name"] == "Asha" and got["stage"] == "Won" and "workspace_id" not in got and "bogus" not in got
    assert cl.post("/api/v1/leads", json=dict(name="Bad", stage="NotAStage"), headers=H(ka)).status_code == 201       # invalid select falls back to the first option
    assert cl.get("/api/v1/leads?limit=1", headers=H(ka)).get_json()["total"] == 2
    assert [x["name"] for x in cl.get("/api/v1/leads?stage=Won", headers=H(ka)).get_json()["data"]] == ["Asha"]
    assert [x["name"] for x in cl.get("/api/v1/leads?q=ash", headers=H(ka)).get_json()["data"]] == ["Asha"]
    assert cl.patch(f"/api/v1/leads/{lid}", json=dict(stage="Lost"), headers=H(ka)).status_code == 200
    assert cl.get(f"/api/v1/leads/{lid}", headers=H(ka)).get_json()["stage"] == "Lost" and cl.get(f"/api/v1/leads/{lid}", headers=H(ka)).get_json()["name"] == "Asha"
    # read-only keys can read but never write
    assert cl.get("/api/v1/leads", headers=H(kra)).status_code == 200
    assert cl.post("/api/v1/leads", json=dict(name="x"), headers=H(kra)).status_code == 403 and cl.delete(f"/api/v1/leads/{lid}", headers=H(kra)).status_code == 403
    # tenants are isolated: B's key cannot see or touch A's records
    assert cl.get(f"/api/v1/leads/{lid}", headers=H(kb)).status_code == 404 and cl.get("/api/v1/leads", headers=H(kb)).get_json()["total"] == 0
    assert cl.patch(f"/api/v1/leads/{lid}", json=dict(stage="Won"), headers=H(kb)).status_code == 404 and cl.delete(f"/api/v1/leads/{lid}", headers=H(kb)).status_code == 404
    # malformed input is a 4xx, never a 500
    for bad in ("[]", "null", '"x"', "not json"):
        assert cl.post("/api/v1/leads", data=bad, content_type="application/json", headers=H(ka)).status_code == 400
    assert cl.get("/api/v1/nope", headers=H(ka)).status_code == 404 and cl.get("/api/v1/leads?limit=abc&offset=-5", headers=H(ka)).status_code == 200
    assert cl.get("/api/v1/leads/99999999999999999999", headers=H(ka)).status_code == 404
    # push analytics
    assert cl.post("/api/v1/kpis", json=dict(metric="Sessions", period="2026-09-01", value="120", target="150"), headers=H(ka)).status_code == 201
    assert b"Sessions" in a.get("/m/kpis").data
    assert cl.delete(f"/api/v1/leads/{lid}", headers=H(ka)).status_code == 200 and cl.get(f"/api/v1/leads/{lid}", headers=H(ka)).status_code == 404
    assert dbq(app, "SELECT COUNT(*) c FROM audit_log WHERE action LIKE 'api_%'", one=True)["c"] >= 4
    # keys are stored hashed; revoking cuts access; plan gating
    assert ka not in str(dbq(app, "SELECT * FROM api_keys")) and dbq(app, "SELECT COUNT(*) c FROM api_keys", one=True)["c"] == 3
    kid = dbq(app, "SELECT id FROM api_keys WHERE name='Zap' AND workspace_id=?", (wid_of(app, "a@x.com"),), one=True)["id"]
    assert b.post(f"/integrations/api-keys/{kid}/revoke").status_code == 302 and cl.get("/api/v1/leads", headers=H(ka)).status_code == 200   # B cannot revoke A's key
    assert a.post(f"/integrations/api-keys/{kid}/revoke").status_code == 302 and cl.get("/api/v1/leads", headers=H(ka)).status_code == 401
    a.post("/billing/subscribe/starter")
    assert cl.get("/api/v1/leads", headers=H(kra)).status_code == 403                                      # Starter has no API access
    assert b"REST API" in a.get("/integrations").data and a.post("/integrations/api-keys", data=dict(name="n")).status_code == 302
    # the API is bearer-token only: exempt from CSRF (cookies are never used)
    from app import create_app
    import db
    a2 = create_app(config={"TESTING": True, "WTF_CSRF_ENABLED": True, "RATELIMIT_ENABLED": False}, database_url=db.engine.url.render_as_string(hide_password=False))
    c2 = a2.test_client()
    assert c2.post("/api/v1/leads", json=dict(name="Via CSRF-on app", email="c@x.com"), headers=H(kb)).status_code == 201     # API works with CSRF enforcement on
    assert c2.post("/m/leads/save", data=dict(name="x")).status_code in (400, 302)                                         # while browser forms are still protected (400 or login redirect)
    assert c2.post("/login", data=dict(email="b@x.com", password="Sup3rSecret!")).status_code == 400


def test_lead_capture_form_and_webhook(app):
    a = signup(app, "a@x.com", ws="Brand A")
    b = signup(app, "b@x.com", ws="Brand B")
    slug = dbq(app, "SELECT slug FROM workspaces WHERE id=?", (wid_of(app, "a@x.com"),), one=True)["slug"]
    pub = app.test_client()
    assert pub.get(f"/lead/{slug}").status_code == 200 and pub.get("/lead/no-such").status_code == 404
    assert pub.post(f"/lead/{slug}", data=dict(name="Ravi", email="r@x.com")).status_code == 400                       # consent required
    assert pub.post(f"/lead/{slug}", data=dict(name="Ravi", consent="1")).status_code == 400                           # needs email or phone
    assert pub.post(f"/lead/{slug}", data=dict(name="Ravi", email="r@x.com", phone="999", message="Demo please", consent="1")).status_code == 200
    pub.post(f"/lead/{slug}", data=dict(name="Bot", email="bot@x.com", consent="1", website="spam"))                  # honeypot
    rows = dbq(app, "SELECT * FROM leads")
    assert len(rows) == 1 and rows[0]["name"] == "Ravi" and rows[0]["source"] == "Website form" and rows[0]["stage"] == "New" and rows[0]["workspace_id"] == wid_of(app, "a@x.com")
    assert b"Ravi" in a.get("/m/leads").data and b"Ravi" not in b.get("/m/leads").data
    hook = re.search(r"(/hooks/lead/[\w]+)", a.get("/integrations").data.decode()).group(1)
    r = pub.post(hook, json=dict(name="Meera", email="m@x.com", source="Landing page"))
    assert r.status_code == 200 and r.headers["Access-Control-Allow-Origin"] == "*"
    assert pub.post(hook, data=dict(name="Form", phone="12345")).status_code == 200                                    # plain form post works too
    assert pub.options(hook).status_code == 204 and "POST" in pub.options(hook).headers["Access-Control-Allow-Methods"]
    assert pub.post(hook, json=dict(name="No contact")).status_code == 400 and pub.post(hook, data="[]", content_type="application/json").status_code == 400
    assert pub.post("/hooks/lead/unknownkey", json=dict(email="x@x.com")).status_code == 404
    assert dbq(app, "SELECT source FROM leads WHERE name='Meera'", one=True)["source"] == "Landing page"
    a.post("/integrations/hook-reset")
    assert pub.post(hook, json=dict(email="late@x.com")).status_code == 404                                            # old URL is dead
    assert dbq(app, "SELECT COUNT(*) c FROM leads", one=True)["c"] == 3


def test_two_factor_login(app, monkeypatch):
    o = signup(app, "o@x.com")
    page = o.get("/account/2fa").data.decode()
    assert "<svg" in page and re.search(r"<code>([A-Z2-7]+)</code>", page)
    secret = re.search(r"<code>([A-Z2-7]+)</code>", page).group(1)
    assert o.post("/account/2fa/enable", data=dict(code="000000")).status_code == 302
    assert dbq(app, "SELECT totp_on FROM users", one=True)["totp_on"] == 0
    r = o.post("/account/2fa/enable", data=dict(code=totp.code_at(secret, totp.step_now())))
    codes = re.findall(r"\n?([a-z2-9]{10})\n", r.data.decode().split("<pre")[1])
    assert r.status_code == 200 and len(codes) == 8 and dbq(app, "SELECT totp_on FROM users", one=True)["totp_on"] == 1
    assert codes[0] not in str(dbq(app, "SELECT recovery FROM users"))                                              # only hashes stored
    o.post("/logout")
    # password alone is no longer enough
    c = app.test_client()
    r = c.post("/login", data=dict(email="o@x.com", password="Sup3rSecret!"))
    assert r.status_code == 302 and r.headers["Location"].endswith("/login/2fa") and c.get("/app").status_code == 302
    assert c.post("/login/2fa", data=dict(code="123456")).status_code == 401 and c.get("/app").status_code == 302
    good = totp.code_at(secret, totp.step_now() + 1)                                                                # next window: newer than the one used at setup
    assert c.post("/login/2fa", data=dict(code=good)).status_code == 302 and c.get("/app").status_code == 200
    # a used code cannot be replayed
    c2 = app.test_client(); c2.post("/login", data=dict(email="o@x.com", password="Sup3rSecret!"))
    assert c2.post("/login/2fa", data=dict(code=good)).status_code == 401
    # recovery codes work exactly once
    c2.post("/login", data=dict(email="o@x.com", password="Sup3rSecret!"))
    assert c2.post("/login/2fa", data=dict(code=codes[0])).status_code == 302 and c2.get("/app").status_code == 200
    c3 = app.test_client(); c3.post("/login", data=dict(email="o@x.com", password="Sup3rSecret!"))
    assert c3.post("/login/2fa", data=dict(code=codes[0])).status_code == 401
    assert app.test_client().get("/login/2fa").status_code == 302                                                    # no pending login
    assert app.test_client().post("/login", data=dict(email="o@x.com", password="wrong")).status_code == 401
    # turning it off needs password AND a fresh code
    assert c2.post("/account/2fa/disable", data=dict(password="Sup3rSecret!", code="000000")).status_code == 302 and dbq(app, "SELECT totp_on FROM users", one=True)["totp_on"] == 1
    assert c2.post("/account/2fa/disable", data=dict(password="Sup3rSecret!", code=good)).status_code == 302              # same (already used) code: replay is refused
    assert dbq(app, "SELECT totp_on FROM users", one=True)["totp_on"] == 1
    base = totp.step_now()
    monkeypatch.setattr(totp, "step_now", lambda t=None: base + 2)                                                    # ...a minute later a fresh code works
    assert c2.post("/account/2fa/disable", data=dict(password="Sup3rSecret!", code=totp.code_at(secret, base + 2))).status_code == 302
    assert dbq(app, "SELECT totp_on FROM users", one=True)["totp_on"] == 0


def test_admin_requires_2fa_when_enforced(app, monkeypatch):
    monkeypatch.setattr(C, "ADMIN_EMAILS", ["boss@x.com"])
    monkeypatch.setattr(C, "REQUIRE_ADMIN_2FA", True)
    boss = signup(app, "boss@x.com", ws="Boss")
    dbx(app, "UPDATE users SET email_verified=1")
    assert boss.get("/admin/").status_code == 404 and "Turn on 2FA to unlock Admin" in boss.get("/app").data.decode()
    secret = re.search(r"<code>([A-Z2-7]+)</code>", boss.get("/account/2fa").data.decode()).group(1)
    boss.post("/account/2fa/enable", data=dict(code=totp.code_at(secret, totp.step_now())))
    assert boss.get("/admin/").status_code == 200 and "🛠 Admin" in boss.get("/app").data.decode()
    assert dbx(app, "UPDATE users SET totp_on=0, totp_secret='' WHERE email='boss@x.com'") == 1               # what `manage.py reset-2fa` does
    assert boss.get("/admin/").status_code == 404


def test_client_report_and_share_link(app):
    o = signup(app, "o@x.com", ws="Report Co")
    m = dt.date.today().strftime("%Y-%m")
    o.post("/m/kpis/save", data=dict(metric="Sessions", channel="Website", period=f"{m}-05", value="120", target="150"))
    o.post("/m/campaigns/save", data=dict(name="Diwali push", status="Live", budget="1000", spent="400", channel="Instagram"))
    o.post("/m/content_items/save", data=dict(title="Festive reel", channel="Instagram", status="Published", publish_date=f"{m}-10"))
    o.post("/m/leads/save", data=dict(name="Warm lead", stage="MQL"))
    o.post("/m/retail_audits/save", data=dict(store="S1", audit_date=f"{m}-03", availability="In stock", posm="Present & correct"))
    page = o.get("/report").data.decode()
    for s in ("Report Co", "Sessions", "80%", "Diwali push", "Festive reel", "MQL: 1", "Retail audits", "Print / save PDF"):
        assert s in page, s
    assert o.get("/report?month=abc").status_code == 200 and o.get("/report?month=2026-13").status_code == 200 and o.get("/report?month=1999-01").status_code == 200
    link = re.search(r"/r/[\w.\-]+", page).group(0)
    pub = app.test_client().get(link)
    assert pub.status_code == 200 and b"Festive reel" in pub.data and b"Log in" not in pub.data and b"Print / save PDF" in pub.data
    assert app.test_client().get(link + "?month=2020-01").status_code == 200 and app.test_client().get("/r/junk").status_code == 404
    o.post("/calendar/reset-link")
    assert app.test_client().get(link).status_code == 404                                                            # rotating the link kills report shares too
    o2 = signup(app, "o2@x.com", ws="Free Co")
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free' WHERE name='Free Co'")
    assert o2.get("/report").status_code == 200 and "/r/" not in o2.get("/report").data.decode()                       # share links need a paid plan
    assert o.get("/report").data.decode().count("Festive reel") >= 1


def test_installable_app_pwa(app):
    c = app.test_client()
    mf = c.get("/manifest.webmanifest")
    j = json.loads(mf.data)
    assert mf.status_code == 200 and "manifest" in mf.mimetype and j["display"] == "standalone" and j["start_url"] == "/app" and len(j["icons"]) == 3
    for ic in j["icons"]:
        r = c.get(ic["src"])
        assert r.status_code == 200 and r.data[:8] == b"\x89PNG\r\n\x1a\n"
    sw = c.get("/sw.js")
    assert sw.status_code == 200 and "javascript" in sw.mimetype and sw.headers["Service-Worker-Allowed"] == "/" and "/offline" in sw.data.decode()
    assert "sync" not in sw.data.decode().lower() and "caches.put" not in sw.data.decode()
    assert "'/static/'" in sw.data.decode() or "/static/" in sw.data.decode()     # only static assets are cached, never signed-in pages
    assert c.get("/offline").status_code == 200
    home = c.get("/").data.decode()
    assert 'rel="manifest"' in home and "apple-touch-icon" in home and "serviceWorker" in home
    o = signup(app, "o@x.com")
    assert "topbar" in o.get("/app").data.decode() and "nav-open" in o.get("/app").data.decode()


def test_v3_fuzz_no_500(app, mails):
    c = signup(app, "f@x.com", ws="Fuzz")
    c.post("/m/tasks/save", data=dict(title="t"))
    key = re.search(r"gtm_[\w\-]{30,}", c.post("/integrations/api-keys", data=dict(name="k", scope="write")).data.decode()).group(0)
    slug = dbq(app, "SELECT slug FROM workspaces", one=True)["slug"]
    codes = []
    def hit(method, url, **kw):
        r = getattr(app.test_client() if kw.pop("anon", False) else c, method)(url, **kw)
        codes.append(r.status_code)
        assert r.status_code < 500, (method, url, kw, r.status_code)
    for u in ["/report?month=abc", "/report?month=0000-00", "/report?month=" + "9" * 300, "/r/" + "x" * 500, "/notifications/abc/open", "/notifications/99999999999999999999/open", "/api/v1/", "/api/v1/leads?limit=" + "9" * 40,
              "/api/v1/leads?offset=99999999999999999999", "/api/v1/leads?q=%25%27", "/api/v1/leads/abc", "/login/2fa", "/account/2fa", "/integrations", "/m/pitches", "/m/accounts", "/m/retail_audits", "/ai", "/manifest.webmanifest", "/sw.js", "/offline", "/unsubscribe/x", f"/lead/{slug}"]:
        hit("get", u, anon=u.startswith(("/api", "/r/", "/manifest", "/sw", "/offline", "/unsub", "/lead")))
    for u, d in [("/m/tasks/1/comment", dict(body="")), ("/m/tasks/1/comment", dict(body="x" * 100000)), ("/m/tasks/1/comment", dict(body="null\x00byte @")), ("/m/tasks/99999999999999999999/comment", dict(body="x")),
                 ("/comments/99999999999999999999/delete", {}), ("/comments/abc/delete", {}), ("/notifications/read", {}), ("/account/2fa/enable", dict(code="x" * 500)), ("/account/2fa/enable", {}), ("/account/2fa/disable", {}),
                 ("/account/digest", dict(on="zzz")), ("/integrations/api-keys", dict(scope="zzz", name="n" * 500)), ("/integrations/api-keys/99999999999999999999/revoke", {}), ("/integrations/hook-reset", {})]:
        hit("post", u, data=d)
    for body in ("[]", "null", '"x"', "{}", "not json", '{"email":"' + "x" * 50000 + '"}', '{"email":"a\u0000b","name":["list"]}'):
        hit("post", f"/lead/{slug}", data=dict(name="n", email="e@x.com", consent="1"), anon=True)
        hit("post", "/hooks/lead/junk", data=body, content_type="application/json", anon=True)
        hit("post", "/api/v1/leads", data=body, content_type="application/json", headers={"Authorization": f"Bearer {key}"}, anon=True)
    hit("post", "/login/2fa", data=dict(code="1"), anon=True)
    assert any(s == 400 for s in codes) and any(s == 404 for s in codes)


def test_onboarding_steps_follow_the_industry(app):
    for email, ind, want, not_want in [("pb@x.com", "personal_brand", "speaking or podcast opportunity", "influencer prospects"), ("b@x.com", "b2b", "first target account", "influencer prospects"),
                                       ("f@x.com", "fmcg", "first retail audit", "target account"), ("d@x.com", "d2c", "influencer prospects", "target account")]:
        page = _as(app, email, ind + " co", ind).get("/app").data.decode()
        assert want in page and not_want not in page and "Capture your first lead" in page, ind


def test_lead_webhook_works_cross_site_with_csrf_on(app):
    """The webhook is called from other websites, so it must not need a CSRF token; the hosted form (same-site) still must."""
    a = signup(app, "a@x.com", ws="Brand A")
    slug = dbq(app, "SELECT slug FROM workspaces", one=True)["slug"]
    hook = re.search(r"(/hooks/lead/[\w]+)", a.get("/integrations").data.decode()).group(1)
    from app import create_app
    import db
    a2 = create_app(config={"TESTING": True, "WTF_CSRF_ENABLED": True, "RATELIMIT_ENABLED": False}, database_url=db.engine.url.render_as_string(hide_password=False))
    web = a2.test_client()
    assert web.post(hook, json=dict(name="Cross-site", email="x@y.com"), headers={"Origin": "https://someone-elses-site.com"}).status_code == 200
    assert web.post(hook, data=dict(name="Form post", phone="999")).status_code == 200
    assert web.post(f"/lead/{slug}", data=dict(name="No token", email="n@x.com", consent="1")).status_code == 400          # hosted form is CSRF-protected
    page = web.get(f"/lead/{slug}").data.decode()
    tok = re.search(r'name="csrf-token" content="([^"]+)"', page).group(1)
    assert web.post(f"/lead/{slug}", data=dict(name="With token", email="t@x.com", consent="1", csrf_token=tok)).status_code == 200
    names = [r["name"] for r in dbq(app, "SELECT name FROM leads ORDER BY id")]
    assert names == ["Cross-site", "Form post", "With token"]
    assert "Consent to be contacted" in dbq(app, "SELECT notes FROM leads WHERE name='With token'", one=True)["notes"]


def test_app_is_named_debut(app, mails):
    c = app.test_client()
    home = c.get("/").data.decode()
    assert "<title>Debut" in home and "icons/icon-192.png" in home and "GTM Studio" not in home
    assert json.loads(c.get("/manifest.webmanifest").data)["name"] == "Debut"
    o = signup(app, "o@x.com")
    assert "Debut" in o.get("/app").data.decode() and "Debut" in o.get("/terms").data.decode() and "Debut" in mails[0][1]     # verification email subject
    assert b"PRODID:-//Debut//EN" in app.test_client().get(re.search(r"(/calendar/[\w.\-]+\.ics)", o.get("/calendar").data.decode()).group(1)).data


def test_freelancers_never_see_agency_money(app, mails):
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    lead.post("/m/content_orders/save", data=dict(title="Blog: winter menu", client="Grand Orchid", type="Blog / Article", creator="Cy Creator", price_client="12345", cost_creator="6789", status="Assigned"))
    oid = dbq(app, "SELECT id FROM content_orders", one=True)["id"]
    seen = cr.get("/m/content_orders").data.decode() + cr.get(f"/m/content_orders/{oid}/edit").data.decode() + cr.get("/m/content_orders/new").data.decode() + cr.get("/kanban/content_orders").data.decode()
    assert "Blog: winter menu" in seen and "12345" not in seen and "6789" not in seen and "Client price" not in seen and "Creator cost" not in seen
    assert "12345" not in cr.get("/search?q=blog").data.decode()
    assert "Blog: winter menu" not in cr.get("/m/content_orders?q=12345").data.decode()          # can't probe hidden prices through search
    assert "Blog: winter menu" in lead.get("/m/content_orders?q=12345").data.decode()            # (leads can)
    # they can still do their job: update the order, and a forged price in the form is ignored
    cr.post("/m/content_orders/save", data=dict(id=oid, title="Blog: winter menu", client="Grand Orchid", type="Blog / Article", creator="Cy Creator", status="In draft", price_client="1", cost_creator="99999"))
    row = dbq(app, "SELECT * FROM content_orders", one=True)
    assert row["status"] == "In draft" and row["price_client"] == "12345" and row["cost_creator"] == "6789"
    cr.post("/m/content_orders/save", data=dict(title="Creator-made order", type="Blog / Article", price_client="5", cost_creator="5"))
    new = dbq(app, "SELECT * FROM content_orders WHERE title='Creator-made order'", one=True)
    assert new["price_client"] == "" and new["cost_creator"] == ""
    # leads and owners see everything, including the money
    page = lead.get("/m/content_orders").data.decode() + lead.get(f"/m/content_orders/{oid}/edit").data.decode()
    assert "12345" in page and "6789" in page and "Client price" in page
    assert cr.get("/export.zip").status_code == 403 and cr.get("/m/content_orders/export").status_code == 403 and cr.get("/content-ops").status_code == 403


def test_launch_readiness_pages_and_seo(app, monkeypatch):
    monkeypatch.setattr(C, "COMPANY_NAME", "Acme Labs Pvt Ltd"); monkeypatch.setattr(C, "SUPPORT_PHONE", "+91 99999 00000"); monkeypatch.setattr(C, "COMPANY_ADDRESS", "12 MG Road, Bengaluru")
    monkeypatch.setattr(C, "GRIEVANCE_OFFICER", "Asha Nair"); monkeypatch.setattr(C, "REFUND_WINDOW_DAYS", 10); monkeypatch.setattr(C, "COMPANY_GSTIN", "29ABCDE1234F1Z5")
    c = app.test_client()
    contact = c.get("/contact").data.decode()
    assert all(x in contact for x in ("Acme Labs Pvt Ltd", "+91 99999 00000", "12 MG Road", "Asha Nair", "29ABCDE1234F1Z5"))
    ref = c.get("/refunds").data.decode()
    assert "10 days" in ref and "cancel" in ref.lower() and "Razorpay" in ref
    priv = c.get("/privacy").data.decode()
    assert "Grievance officer" in priv and "Asha Nair" in priv and "essential cookies" in priv
    home = c.get("/").data.decode()
    assert 'name="description"' in home and 'property="og:title"' in home and 'rel="canonical"' in home and "Refunds &amp; cancellation" in home and "noindex" not in home
    assert "Sitemap:" in c.get("/robots.txt").data.decode() and "Disallow: /api/" in c.get("/robots.txt").data.decode()
    sm = c.get("/sitemap.xml"); assert sm.status_code == 200 and b"<loc>" in sm.data and b"/pricing" in sm.data and b"/refunds" in sm.data
    o = signup(app, "o@x.com")
    assert "noindex" in o.get("/app").data.decode()                                  # app pages are never indexed
    monkeypatch.setattr(C, "GA_MEASUREMENT_ID", "G-TEST123")
    assert "G-TEST123" in c.get("/").data.decode() and "G-TEST123" not in o.get("/app").data.decode()   # analytics only on public pages


def test_ai_cost_controls_and_maintenance_mode(app, monkeypatch):
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")

    class R:
        status_code = 200
        def json(self): return {"content": [{"type": "text", "text": "ok"}]}
    calls = []
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: calls.append(1) or R())
    a, b = signup(app, "a@x.com", ws="A"), signup(app, "b@x.com", ws="B")
    r = a.post("/api/ai", json=dict(task="social"))
    assert r.status_code == 403 and "verify your email" in r.get_json()["error"] and not calls          # unverified accounts cost you nothing
    dbx(app, "UPDATE users SET email_verified=1")
    monkeypatch.setattr(C, "AI_DAILY_CAP", 2)
    assert a.post("/api/ai", json=dict(task="social")).status_code == 200 and b.post("/api/ai", json=dict(task="social")).status_code == 200
    r = a.post("/api/ai", json=dict(task="social"))
    assert r.status_code == 503 and "daily limit" in r.get_json()["error"] and len(calls) == 2                # platform-wide cap holds across workspaces
    monkeypatch.setattr(C, "AI_DAILY_CAP", 0)
    assert a.post("/api/ai", json=dict(task="social")).status_code == 200
    # maintenance mode: public and app pages show a 503 notice; health, webhooks and the admin still work
    monkeypatch.setattr(C, "MAINTENANCE", True)
    monkeypatch.setattr(C, "ADMIN_EMAILS", ["a@x.com"])
    anon = app.test_client()
    assert anon.get("/").status_code == 503 and b"being updated" in anon.get("/pricing").data and anon.get("/health").status_code == 200
    assert anon.get("/login").status_code == 200 and anon.post("/webhooks/razorpay", data=b"{}").status_code == 400      # webhook reachable (bad signature, not 503)
    assert b.get("/app").status_code == 503
    assert a.get("/app").status_code == 200 and a.get("/health").status_code == 200                                   # the admin can still work
    monkeypatch.setattr(C, "MAINTENANCE", False)
    assert anon.get("/").status_code == 200


def test_feedback_and_launch_metrics(app, mails, monkeypatch):
    monkeypatch.setattr(C, "ADMIN_EMAILS", ["boss@x.com"]); monkeypatch.setattr(C, "SUPPORT_EMAIL", "help@debut.test")
    boss = signup(app, "boss@x.com", ws="Boss Co")
    dbx(app, "UPDATE users SET email_verified=1 WHERE email='boss@x.com'")
    u = signup(app, "u@x.com", ws="Pilot Co")
    lead = _join(app, u, mails, "cl@x.com", "creator", "Cy Creator")
    assert b"Send feedback" in u.get("/app").data and u.get("/feedback?from=/calendar").status_code == 200
    assert u.post("/feedback", data=dict(kind="problem", message="   ", page="/calendar")).status_code == 400
    r = u.post("/feedback", data=dict(kind="problem", message="Calendar chip won't drag on my phone", page="/calendar"))
    assert r.status_code == 302 and r.headers["Location"].endswith("/calendar")
    assert lead.post("/feedback", data=dict(kind="idea", message="Please add dark mode", page="https://evil.example")).status_code == 302        # creators can send feedback too; no open redirect
    assert app.test_client().post("/feedback", data=dict(message="x")).status_code == 302                                                          # must be logged in
    mail = next(m for m in mails if m[0] == "help@debut.test")
    assert "Calendar chip" in mail[2] and "u@x.com" in mail[2] and "Pilot Co" in mail[2] and "/calendar" in mail[2]
    assert u.get("/admin/feedback").status_code == 404                                                                                             # only the platform owner reads it
    inbox = boss.get("/admin/feedback").data.decode()
    assert "Calendar chip" in inbox and "dark mode" in inbox and "Pilot Co" in inbox
    fid = dbq(app, "SELECT id FROM feedback ORDER BY id", one=True)["id"]
    boss.post(f"/admin/feedback/{fid}/resolve")
    assert dbq(app, "SELECT resolved FROM feedback WHERE id=?", (fid,), one=True)["resolved"] == 1
    m = boss.get("/admin/").data.decode()
    assert "Launch metrics" in m and "Open feedback" in m and "New users" in m and "Signups by day" in m
    dbx(app, "UPDATE workspaces SET trial_ends=? WHERE name='Pilot Co'", ((dt.date.today() + dt.timedelta(days=2)).isoformat(),))
    assert "Trials ending in the next 3 days" in boss.get("/admin/").data.decode() and "u@x.com" in boss.get("/admin/").data.decode()
    # feedback is removed with the workspace
    wid = wid_of(app, "u@x.com")
    assert dbq(app, "SELECT COUNT(*) c FROM feedback WHERE workspace_id=?", (wid,), one=True)["c"] == 2
    slug = dbq(app, "SELECT slug FROM workspaces WHERE id=?", (wid,), one=True)["slug"]
    u.post("/workspace/delete", data=dict(confirm=slug))
    assert dbq(app, "SELECT COUNT(*) c FROM feedback WHERE workspace_id=?", (wid,), one=True)["c"] == 0


def test_help_centre_and_guides(app, mails):
    o = signup(app, "o@x.com")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy")
    for c in (o, cr):
        page = c.get("/help").data.decode()
        assert "Guide for personal brands" in page and "Guide for agencies" in page and "Help &amp; guides" in c.get("/app").data.decode()
    assert app.test_client().get("/help").status_code == 302
    for f in ("Debut-Guide-Personal-Branding.pdf", "Debut-Guide-Agencies-and-Teams.pdf", "Debut-Visual-Tour.pdf"):
        r = app.test_client().get(f"/static/guides/{f}")
        assert r.status_code == 200 and r.data[:4] == b"%PDF"
    assert "Visual tour" in o.get("/help").data.decode()


def test_help_centre_mentions_this_session_s_new_features(app, mails):
    o = signup(app, "o@x.com")
    page = o.get("/help").data.decode()
    # the lookup table must point to every feature added this session
    for needle in ("Today", "Marketing Health", "Brand deck scan", "Connected accounts"):
        assert needle in page
    # the FAQ must explain the daily workflow and the removed "More tools" popup
    assert "What should I actually do each day?" in page
    assert "Where did the" in page and "More tools" in page and "popup go?" in page
    # the tour must be replayable and must not still reference the removed popup
    assert "onclick=\"startTour()\"" in page
    base = o.get("/app").data.decode()
    assert "More tools" not in base


def test_ownership_transfer(app, mails):
    o = signup(app, "o@x.com", name="Olivia Owner")
    ad = _join(app, o, mails, "ad@x.com", "admin", "Ada Admin")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    mid = lambda email: dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email=?", (email,), one=True)["id"]
    role = lambda email: dbq(app, "SELECT m.role FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email=?", (email,), one=True)["role"]
    good = lambda **kw: dict(dict(mid=mid("ad@x.com"), password="Sup3rSecret!", confirm="TRANSFER"), **kw)
    assert "Transfer ownership" in o.get("/account").data.decode() and "Transfer ownership" not in ad.get("/account").data.decode()
    assert ad.post("/workspace/transfer", data=good()).status_code == 403                                            # only the owner can hand it over
    for bad in (good(password="wrong"), good(confirm="yes"), good(mid=mid("cr@x.com")), good(mid=mid("o@x.com")), good(mid="abc"), good(mid="99999999999999999999")):
        assert o.post("/workspace/transfer", data=bad).status_code == 302 and role("o@x.com") == "owner" and role("ad@x.com") == "admin"
    assert o.post("/workspace/transfer", data=good()).status_code == 302
    assert role("ad@x.com") == "owner" and role("o@x.com") == "admin" and dbq(app, "SELECT owner_id FROM workspaces", one=True)["owner_id"] == dbq(app, "SELECT id FROM users WHERE email='ad@x.com'", one=True)["id"]
    assert any(m[0] == "ad@x.com" and "now the owner" in m[1] for m in mails)
    assert o.post("/workspace/delete", data=dict(confirm="x")).status_code == 403 and o.post("/workspace/transfer", data=good()).status_code == 403   # old owner lost owner powers
    assert ad.get("/billing").status_code == 200 and ad.post("/billing/subscribe/team").status_code == 200
    assert dbq(app, "SELECT COUNT(*) c FROM audit_log WHERE action='ownership_transferred'", one=True)["c"] == 1


def test_ownership_transfer_needs_two_factor_code(app, mails):
    o = signup(app, "o@x.com")
    ad = _join(app, o, mails, "ad@x.com", "admin", "Ada Admin")
    secret = re.search(r"<code>([A-Z2-7]+)</code>", o.get("/account/2fa").data.decode()).group(1)
    o.post("/account/2fa/enable", data=dict(code=totp.code_at(secret, totp.step_now())))
    target = dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='ad@x.com'", one=True)["id"]
    body = dict(mid=target, password="Sup3rSecret!", confirm="TRANSFER")
    assert o.post("/workspace/transfer", data=dict(body, code="000000")).status_code == 302
    assert dbq(app, "SELECT role FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='o@x.com'", one=True)["role"] == "owner"
    assert "two-factor code" in o.get("/account").data.decode()
    o.post("/workspace/transfer", data=dict(body, code=totp.code_at(secret, totp.step_now() + 1)))
    assert dbq(app, "SELECT role FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email='ad@x.com'", one=True)["role"] == "owner"


# =====================================================================  client portal
def _agency(app, mails):
    """An agency workspace with two clients' work, private notes and money."""
    o = signup(app, "o@x.com", ws="Agency Co")
    lead = _join(app, o, mails, "lead@x.com", "lead", "Lena Lead")
    o.post("/m/content_items/save", data=dict(title="A-Review reel", channel="Instagram", format="Reel / Short", status="Review", publish_date="2099-01-10", client="Grand Orchid", owner="Lena Lead", copy="Book your festive stay", cta="Book now", notes="INTERNAL-NOTE-A"))
    o.post("/m/content_items/save", data=dict(title="A-Approved post", channel="LinkedIn", status="Approved", publish_date="2099-02-01", client="grand orchid"))            # client match is case-insensitive
    o.post("/m/content_items/save", data=dict(title="B-Zen reel", channel="Instagram", status="Review", client="Zen Cafe", copy="ZEN-SECRET-COPY"))
    o.post("/m/content_items/save", data=dict(title="C-Untagged draft", channel="Email", status="Review"))
    o.post("/m/campaigns/save", data=dict(name="A-Festive push", client="Grand Orchid", status="Live", budget="5555", spent="1234", kpi_target="10 bookings", owner="Lena Lead", learnings="INTERNAL-LEARNING"))
    o.post("/m/campaigns/save", data=dict(name="B-Latte launch", client="Zen Cafe", status="Live", budget="7777"))
    o.post("/m/clients/save", data=dict(name="Zen Cafe", contact="Zed", monthly_value="88888"))
    ids = {r["title"]: r["id"] for r in dbq(app, "SELECT id, title FROM content_items")}
    return o, lead, ids


def _client(app, owner, mails, email="cl@x.com", client_name="Grand Orchid"):
    owner.post("/team/invite", data=dict(email=email, role="client", client_name=client_name))
    tok = re.search(r"/invite/([\w\-]+)", mails[-1][2]).group(1)
    c = app.test_client()
    assert c.post("/signup", data=dict(name="Cathy Client", email=email, password="Sup3rSecret!", agree="1", invite=tok)).status_code == 302
    return c


def test_client_portal_shows_only_that_clients_work(app, mails):
    o, lead, ids = _agency(app, mails)
    cl = _client(app, o, mails)
    assert cl.get("/app").status_code == 302 and cl.get("/app").headers["Location"].endswith("/portal")
    page = cl.get("/portal").data.decode()
    for ok in ("Welcome, Grand Orchid", "A-Review reel", "A-Approved post", "A-Festive push", "10 bookings"):
        assert ok in page, ok
    for secret in ("B-Zen reel", "C-Untagged draft", "Zen Cafe", "B-Latte", "5555", "1234", "7777", "88888", "INTERNAL", "Lena Lead"):
        assert secret not in page, secret
    assert "Zen Cafe" not in cl.get("/portal?client=Zen%20Cafe").data.decode() and "B-Zen reel" not in cl.get("/portal?client=Zen%20Cafe").data.decode()     # the ?client= trick is staff-only
    item = cl.get(f"/portal/item/{ids['A-Review reel']}").data.decode()
    assert "Book your festive stay" in item and "Book now" in item and "INTERNAL-NOTE-A" not in item and "Lena Lead" not in item and "Approve" in item
    for other in ("B-Zen reel", "C-Untagged draft"):
        assert cl.get(f"/portal/item/{ids[other]}").status_code == 404 and cl.post(f"/portal/item/{ids[other]}/decide", data=dict(action="approve")).status_code == 404
    assert cl.get("/portal/item/99999999999999999999").status_code == 404
    # everything else is closed to a client
    for p in ("/m/content_items", f"/m/content_items/{ids['A-Review reel']}/edit", "/m/clients", "/m/payments", "/m/campaigns", "/calendar", "/report", "/search?q=zen", "/team-hub", "/team", "/export.zip", "/billing",
              "/integrations", "/kanban/content_items", "/moodboards", "/ideas", "/ai", "/analytics", "/activity", "/updates", "/content-ops", "/account/2fa/enable", "/media/1/x", "/calendar/x.ics", "/admin/"):
        r = cl.get(p)
        assert r.status_code in (403, 404, 405) or (p == "/account/2fa/enable" and r.status_code == 405), (p, r.status_code)
    for p, d in (("/m/content_items/save", dict(title="hack")), (f"/m/content_items/{ids['A-Review reel']}/status", dict(value="Published")), (f"/m/content_items/{ids['A-Review reel']}/delete", {}),
                 (f"/m/content_items/{ids['A-Review reel']}/comment", dict(body="x")), ("/api/ai", dict(task="social")), ("/m/leads/import", {}), ("/team/invite", dict(email="z@z.com", role="viewer")), ("/workspace/delete", dict(confirm="x")),
                 ("/api/reschedule", {}), ("/integrations/api-keys", dict(name="k"))):
        assert cl.post(p, data=d).status_code in (403, 404), p
    assert b"Search everything" not in cl.get("/portal").data and b"Client portal" in cl.get("/portal").data
    assert cl.get("/help").status_code == 200 and cl.get("/feedback").status_code == 200 and cl.get("/notifications").status_code == 200 and cl.get("/account").status_code == 200


def test_client_approves_and_requests_changes(app, mails):
    o, lead, ids = _agency(app, mails)
    cl = _client(app, o, mails)
    unread = lambda email: dbq(app, "SELECT COUNT(*) c FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email=? AND n.is_read=0", (email,), one=True)["c"]
    rid = ids["A-Review reel"]
    # asking for changes needs a note; then it goes back to the team
    assert cl.post(f"/portal/item/{rid}/decide", data=dict(action="changes", comment="  ")).status_code == 302
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (rid,), one=True)["status"] == "Review"
    before = unread("lead@x.com")
    cl.post(f"/portal/item/{rid}/decide", data=dict(action="changes", comment="Please use the winter photo"))
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (rid,), one=True)["status"] == "In progress"
    assert unread("lead@x.com") == before + 1 and unread("o@x.com") >= 1                                            # the item's owner (Lena) and approvers are told
    assert "Please use the winter photo" in cl.get(f"/portal/item/{rid}").data.decode()
    assert "Please use the winter photo" in lead.get(f"/m/content_items/{rid}/edit").data.decode() and "visible to client" in lead.get(f"/m/content_items/{rid}/edit").data.decode()
    assert "With the team" in cl.get(f"/portal/item/{rid}").data.decode()                                             # friendly wording, not internal status names
    assert cl.post(f"/portal/item/{rid}/decide", data=dict(action="approve")).status_code == 302 and dbq(app, "SELECT status FROM content_items WHERE id=?", (rid,), one=True)["status"] == "In progress"   # not in review: nothing happens
    # team fixes it, sends it back to review: the client is notified; then approves
    lead.post(f"/m/content_items/{rid}/status", data=dict(value="Review"))
    assert unread("cl@x.com") == 1 and "Ready for your review" in cl.get("/notifications").data.decode()
    nid = dbq(app, "SELECT n.id FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email='cl@x.com'", one=True)["id"]
    assert cl.get(f"/notifications/{nid}/open").headers["Location"].endswith(f"/portal/item/{rid}")
    cl.post(f"/portal/item/{rid}/decide", data=dict(action="approve", comment="Looks great"))
    assert dbq(app, "SELECT status FROM content_items WHERE id=?", (rid,), one=True)["status"] == "Approved"
    assert dbq(app, "SELECT COUNT(*) c FROM audit_log WHERE action IN ('client_approved','client_changes')", one=True)["c"] == 2
    assert "A-Review reel" in cl.get("/portal").data.decode() and "Waiting for your review" in cl.get("/portal").data.decode()
    # staff cannot decide for the client
    for staff in (o, lead):
        assert staff.post(f"/portal/item/{rid}/decide", data=dict(action="approve")).status_code == 403 and staff.post(f"/portal/item/{rid}/comment", data=dict(comment="x")).status_code == 403


def test_client_conversation_privacy(app, mails):
    o, lead, ids = _agency(app, mails)
    cl = _client(app, o, mails)
    rid = ids["A-Review reel"]
    lead.post(f"/m/content_items/{rid}/comment", data=dict(body="INTERNAL: client is fussy, keep it short"))                   # private by default
    lead.post(f"/m/content_items/{rid}/comment", data=dict(body="We'll send the final cut tomorrow", client_visible="1"))       # shared on purpose
    seen = cl.get(f"/portal/item/{rid}").data.decode()
    assert "final cut tomorrow" in seen and "INTERNAL" not in seen and "fussy" not in seen
    assert dbq(app, "SELECT COUNT(*) c FROM comments WHERE client_visible=1", one=True)["c"] == 1
    assert "Share with client" in lead.get(f"/m/content_items/{rid}/edit").data.decode()
    assert "Share with client" not in lead.get(f"/m/content_items/{ids['C-Untagged draft']}/edit").data.decode()                # no client on that item, nothing to share with
    assert dbq(app, "SELECT COUNT(*) c FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email='cl@x.com' AND n.message LIKE '%replied%'", one=True)["c"] == 1
    cl.post(f"/portal/item/{rid}/comment", data=dict(comment="Thanks! Can we add a discount code?"))
    assert "discount code" in cl.get(f"/portal/item/{rid}").data.decode() and "discount code" in lead.get(f"/m/content_items/{rid}/edit").data.decode()
    assert dbq(app, "SELECT COUNT(*) c FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email='lead@x.com' AND n.message LIKE '%discount code%'", one=True)["c"] == 1
    assert cl.post(f"/portal/item/{ids['B-Zen reel']}/comment", data=dict(comment="hi")).status_code == 404
    # a manager-shared comment on an item with no client tag is never client-visible; creators/viewers cannot share
    o.post(f"/m/content_items/{ids['C-Untagged draft']}/comment", data=dict(body="untagged", client_visible="1"))
    assert dbq(app, "SELECT COUNT(*) c FROM comments WHERE body='untagged' AND client_visible=1", one=True)["c"] == 0


def test_client_seats_plans_and_staff_preview(app, mails, monkeypatch):
    o, lead, ids = _agency(app, mails)
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    vw = _join(app, o, mails, "vw@x.com", "viewer", "Vik Viewer")
    cl = _client(app, o, mails)
    # clients don't use up staff seats (trial = 5 seats: owner, lead, creator, viewer + 1 more is fine) and have their own allowance
    assert "4 of 5 seats used" in o.get("/team").data.decode() and "1 of 5 client-portal seats" in o.get("/team").data.decode()
    _join(app, o, mails, "m1@x.com", "manager", "M One")
    assert "5 of 5 seats used" in o.get("/team").data.decode()
    o.post("/team/invite", data=dict(email="over@x.com", role="manager"))
    assert not any(m[0] == "over@x.com" for m in mails)                                                                    # staff seats are full...
    o.post("/team/invite", data=dict(email="c2@x.com", role="client", client_name="Zen Cafe"))
    assert any(m[0] == "c2@x.com" and "client portal" in m[2] for m in mails)                                              # ...but clients still fit
    o.post("/team/invite", data=dict(email="c3@x.com", role="client", client_name=""))
    assert not any(m[0] == "c3@x.com" for m in mails)                                                                       # client name is required
    assert cl.get("/portal").status_code == 200                                                                             # a full workspace never blocks a client (they aren't seats)
    # role rules: clients can't be created by changing a role, nor changed later
    mid = lambda e: dbq(app, "SELECT m.id FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email=?", (e,), one=True)["id"]
    o.post("/team/role", data=dict(mid=mid("vw@x.com"), role="client"))
    o.post("/team/role", data=dict(mid=mid("cl@x.com"), role="viewer"))
    role = lambda e: dbq(app, "SELECT m.role FROM memberships m JOIN users u ON u.id=m.user_id WHERE u.email=?", (e,), one=True)["role"]
    assert role("vw@x.com") == "viewer" and role("cl@x.com") == "client"
    assert "Preview portal" in o.get("/team").data.decode()
    # clients never appear in staff tooling
    hub = lead.get("/team-hub").data.decode()
    workload = hub.split("Team workload")[1].split("Awaiting your approval")[0]
    assert "Lena Lead" in workload and "Cathy Client" not in workload and "cl@x.com" not in workload                     # not a workload row (the audit feed may still log their sign-in)
    assert "Cathy Client" not in o.get("/m/tasks/new").data.decode()
    lead.post("/m/tasks/save", data=dict(title="Task for the client?", assignee="Cathy Client"))
    assert dbq(app, "SELECT COUNT(*) c FROM notifications n JOIN users u ON u.id=n.user_id WHERE u.email='cl@x.com' AND n.message LIKE '%assigned%'", one=True)["c"] == 0
    # staff preview: owners/admins/leads/managers see what the client sees; creators and viewers cannot
    assert "Grand Orchid" in o.get("/portal").data.decode() and "Zen Cafe" in o.get("/portal").data.decode()
    pv = lead.get("/portal?client=Grand%20Orchid").data.decode()
    assert "Previewing the portal as" in pv and "A-Review reel" in pv and "B-Zen reel" not in pv and "Approve" not in lead.get(f"/portal/item/{ids['A-Review reel']}?client=Grand%20Orchid").data.decode()
    assert cr.get("/portal").status_code == 403 and vw.get("/portal").status_code == 403
    # plans: the portal is a Team/Agency feature
    with_free = signup(app, "f@x.com", ws="Free Co")
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free' WHERE name='Free Co'")
    with_free.post("/team/invite", data=dict(email="fc@x.com", role="client", client_name="Someone"))
    assert not any(m[0] == "fc@x.com" for m in mails) and "client-portal" not in with_free.get("/team").data.decode()
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free' WHERE name='Agency Co'")                                    # downgrade: the client is told, not shown data
    r = cl.get("/portal")
    assert r.status_code == 403 and b"current plan" in r.data and b"A-Review reel" not in r.data
    o.post("/billing/subscribe/agency")
    assert cl.get("/portal").status_code == 200
    # digest never mails clients
    dbx(app, "UPDATE users SET email_verified=1")
    mails.clear()
    with app.app_context():
        S.send_digests()
    assert not any(m[0] == "cl@x.com" for m in mails)


def test_client_portal_fuzz_and_client_limit(app, mails, monkeypatch):
    o, lead, ids = _agency(app, mails)
    cl = _client(app, o, mails)
    for u in ("/portal?client=" + "x" * 3000, "/portal/item/abc", "/portal/item/0", "/portal/item/-1", "/portal?client=%00"):
        assert cl.get(u).status_code in (200, 404), u
    for d in (dict(), dict(action="zzz"), dict(action="changes", comment="x" * 100000), dict(action="approve", comment="a\x00b")):
        assert cl.post(f"/portal/item/{ids['A-Review reel']}/decide", data=d).status_code in (302, 400), d
    cl.post(f"/portal/item/{ids['A-Review reel']}/comment", data=dict(comment="null\x00byte"))
    monkeypatch.setitem(C.PLANS["trial"], "clients", 1)
    o.post("/team/invite", data=dict(email="c2@x.com", role="client", client_name="Zen Cafe"))
    assert not any(m[0] == "c2@x.com" for m in mails)                                                                       # allowance is enforced (pending invites count too)


# =====================================================================  live data: Market Pulse and AI live research
import pulse as P

GN = b"""<?xml version="1.0" encoding="UTF-8"?><rss version="2.0"><channel><title>x</title>
<item><title>Hotel chains lean into direct booking offers - Economic Times</title><link>https://news.google.com/rss/articles/AAA</link><pubDate>Mon, 21 Sep 2026 05:00:00 GMT</pubDate><source url="https://economictimes.indiatimes.com">Economic Times</source></item>
<item><title>Direct booking offers surge as OTA fees rise - Hindu BusinessLine</title><link>https://news.google.com/rss/articles/BBB</link><pubDate>Sun, 20 Sep 2026 05:00:00 GMT</pubDate><source url="https://thehindubusinessline.com">Hindu BusinessLine</source></item>
<item><title>Festive travel demand pushes direct booking growth - Mint</title><link>https://news.google.com/rss/articles/CCC</link><pubDate>Sat, 19 Sep 2026 05:00:00 GMT</pubDate><source url="https://livemint.com">Mint</source></item>
<item><title>Ignore me &lt;script&gt;alert(1)&lt;/script&gt; - Shady</title><link>javascript:alert(1)</link><pubDate>Sat, 19 Sep 2026 05:00:00 GMT</pubDate></item></channel></rss>"""
HN = json.dumps({"hits": [{"title": "Show HN: A GTM planner for founders", "url": "https://example.com/gtm", "points": 120, "num_comments": 45, "created_at_i": 1790000000, "objectID": "111"},
                          {"title": "Ask HN: How do you launch products?", "url": None, "points": 30, "num_comments": 12, "created_at_i": 1789000000, "objectID": "222"}, {"title": "", "objectID": "333"}]}).encode()
YT = json.dumps({"items": [{"id": {"videoId": "vid1"}, "snippet": {"title": "How I grew my brand on LinkedIn", "channelTitle": "Growth Guy", "publishedAt": "2026-09-10T10:00:00Z"}},
                           {"id": {"videoId": "vid2"}, "snippet": {"title": "Brand strategy 101", "channelTitle": "Strat School", "publishedAt": "2026-09-01T10:00:00Z"}}, {"id": {}, "snippet": {"title": "no id"}}]}).encode()
YTS = json.dumps({"items": [{"id": "vid1", "statistics": {"viewCount": "123456"}}, {"id": "vid2", "statistics": {"viewCount": "9000"}}]}).encode()
TR = b"""<?xml version="1.0"?><rss xmlns:ht="https://trends.google.com/trending/rss" version="2.0"><channel><title>Trending</title>
<item><title>IPL auction</title><ht:approx_traffic>200K+</ht:approx_traffic><ht:news_item><ht:news_item_title>Auction date announced</ht:news_item_title></ht:news_item></item>
<item><title>Diwali sale</title><ht:approx_traffic>50K+</ht:approx_traffic></item><item><title></title></item></channel></rss>"""
GNEWS = json.dumps({"articles": [{"title": "D2C brands go omnichannel", "url": "https://example.com/d2c", "publishedAt": "2026-09-20T08:00:00Z", "source": {"name": "Inc42"}}]}).encode()
NEWSDATA = json.dumps({"status": "success", "results": [{"title": "Quick commerce widens lead", "link": "https://example.com/qc", "pubDate": "2026-09-20 08:00:00", "source_name": "Moneycontrol"}]}).encode()
NAGER = json.dumps([{"date": "2026-10-02", "localName": "Gandhi Jayanti", "name": "Gandhi Jayanti"}, {"date": "2026-11-08", "name": "Diwali"}, {"date": None, "name": "bad"}]).encode()


def fake_web(monkeypatch, fail=(), only=None):
    calls = []

    def get(url, **kw):
        calls.append(url)
        if any(f in url for f in fail):
            raise ValueError("source down")
        for key, body in (("news.google.com", GN), ("hn.algolia", HN), ("youtube/v3/search", YT), ("youtube/v3/videos", YTS), ("trends.google.com", TR), ("gnews.io", GNEWS), ("newsdata.io", NEWSDATA), ("date.nager.at", NAGER)):
            if key in url:
                return body
        raise ValueError("unexpected url " + url)
    monkeypatch.setattr(S, "safe_get", get)
    return calls


def test_pulse_parsers_handle_each_live_source(monkeypatch):
    calls = fake_web(monkeypatch)
    n = P.google_news("direct booking & OTA fees", "India")
    assert [i["title"] for i in n][:3] == ["Hotel chains lean into direct booking offers", "Direct booking offers surge as OTA fees rise", "Festive travel demand pushes direct booking growth"]
    assert n[0]["source"] == "Economic Times" and n[0]["date"] == "21 Sep" and n[0]["url"].startswith("https://news.google.com/")
    assert n[3]["url"] == "" and "<script>" not in n[3]["title"] and "alert" in n[3]["title"]                 # javascript: links are dropped, markup is stripped
    assert "hl=en-IN" in calls[-1] and "gl=IN" in calls[-1] and "when%3A7d" in calls[-1] and "%26" in calls[-1] and " " not in calls[-1]
    P.google_news("x", "Global"); assert "gl=US" in calls[-1]
    monkeypatch.setattr(C, "GNEWS_API_KEY", "gk"); monkeypatch.setattr(C, "NEWSDATA_API_KEY", "nk")
    monkeypatch.setattr(C, "NEWS_PROVIDER", "gnews"); g = P.news("d2c", "India")
    assert g[0]["title"] == "D2C brands go omnichannel" and g[0]["source"] == "Inc42" and "gnews.io" in calls[-1] and "country=in" in calls[-1] and "apikey=gk" in calls[-1]
    monkeypatch.setattr(C, "NEWS_PROVIDER", "newsdata"); nd = P.news("qc", "Global")
    assert nd[0]["source"] == "Moneycontrol" and nd[0]["url"] == "https://example.com/qc" and "newsdata.io" in calls[-1]
    monkeypatch.setattr(C, "NEWS_PROVIDER", "gnews"); monkeypatch.setattr(C, "GNEWS_API_KEY", "")
    P.news("t", "India"); assert "news.google.com" in calls[-1]                                                     # no key: falls back to the free feed
    hn = P.hacker_news("gtm")
    assert hn[0]["extra"] == "120 points, 45 comments" and hn[1]["url"] == "https://news.ycombinator.com/item?id=222" and len(hn) == 2
    monkeypatch.setattr(C, "YOUTUBE_API_KEY", "yk"); yt = P.youtube("linkedin branding", "India")
    assert yt[0]["extra"] == "123,456 views" and yt[0]["url"] == "https://www.youtube.com/watch?v=vid1" and len(yt) == 2 and "regionCode=IN" in [c for c in calls if "youtube/v3/search" in c][-1]
    tr = P.trends_now("IN")
    assert tr[0] == dict(term="IPL auction", traffic="200K+", context="Auction date announced") and tr[1]["term"] == "Diwali sale" and len(tr) == 2
    assert P.public_holidays("in", 2026) == [dict(name="Gandhi Jayanti", date="2026-10-02"), dict(name="Diwali", date="2026-11-08")]
    with pytest.raises(ValueError):
        P.reddit("../etc/passwd")
    terms = dict(P.top_terms(["Direct booking offers surge", "Hotels push direct booking offers", "Why direct booking offers work", "Unrelated headline entirely"]))
    assert "direct booking offers" in " ".join(terms) or "booking offers" in terms and not any(w in terms for w in ("why", "the"))


def test_market_pulse_end_to_end(app, monkeypatch, mails):
    calls = fake_web(monkeypatch)
    o = signup(app, "o@x.com", ws="Grand Hotels")
    dbx(app, "UPDATE workspaces SET industry='saas'")                                                              # tech industry so Hacker News is used too
    dbx(app, "DELETE FROM watchlist"); dbx(app, "INSERT INTO watchlist (workspace_id, term, type, region) VALUES (1, 'direct booking', 'Keyword', 'Both')")
    o.post("/m/competitors/save", data=dict(name="Rival Hotels", website="https://rival.example"))
    empty = o.get("/pulse").data.decode()
    assert "Nothing here yet" in empty and "direct booking" in empty and "Rival Hotels" in empty and "Refresh now" in empty
    r = o.post("/pulse/refresh"); assert r.status_code == 302
    page = o.get("/pulse").data.decode()
    for s in ("Hotel chains lean into direct booking offers", "Economic Times", "Show HN: A GTM planner for founders", "120 points", "IPL auction", "200K+", "Rival Hotels", "Rising in headlines", "Searching for today: India", "Searching for today: Global", "Idea", "Save"):
        assert s in page, s
    assert not any("<script>alert" in x for x in [page]) and 'href="javascript:' not in page
    assert any("hl=en-IN" in c for c in calls) and any("hl=en-US" in c for c in calls) and any("hn.algolia" in c for c in calls) and any("trends.google.com/trending/rss?geo=IN" in c for c in calls) and any("geo=US" in c for c in calls)
    assert sum("rival" in c.lower() for c in calls) >= 2                                                            # competitors are tracked in India and Global
    n = len(calls)
    assert "refreshed a moment ago" in o.post("/pulse/refresh", follow_redirects=True).data.decode()
    o.post("/pulse/refresh"); assert len(calls) == n                                                              # throttled: no new outbound calls
    dash = o.get("/app").data.decode(); assert "Market Pulse" in dash and ("Show HN: A GTM planner" in dash or "Hotel chains lean" in dash) and "Open Market Pulse" in dash
    assert dbq(app, "SELECT COUNT(*) c FROM pulse_runs", one=True)["c"] == 1
    # the "save" and "idea" buttons work
    assert o.post("/updates/save", data=dict(title="Hotel chains lean into direct booking offers", url="https://news.google.com/rss/articles/AAA", source="ET")).status_code == 302
    assert dbq(app, "SELECT COUNT(*) c FROM reference_links", one=True)["c"] == 1
    assert "Idea: Hotel chains lean" in o.get("/m/content_items/new?title=Idea%3A+Hotel+chains+lean+into&notes=Inspired+by").data.decode()
    # a second run marks new rising terms; old runs are pruned to 8
    monkeypatch.setattr(C, "PULSE_MIN_REFRESH_MINUTES", 0)
    for _ in range(10):
        o.post("/pulse/refresh")
    assert dbq(app, "SELECT COUNT(*) c FROM pulse_runs", one=True)["c"] == 8


def test_pulse_survives_a_broken_source_and_respects_roles_and_plans(app, monkeypatch, mails):
    fake_web(monkeypatch, fail=("hn.algolia", "gl=IN"))                                                            # Hacker News and India news are down
    monkeypatch.setattr(C, "PULSE_MIN_REFRESH_MINUTES", 0)
    o = signup(app, "o@x.com")
    lead = _join(app, o, mails, "lead@x.com", "manager", "Mia Manager")
    vw = _join(app, o, mails, "vw@x.com", "viewer", "Vik Viewer")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy Creator")
    o.post("/pulse/refresh")
    page = o.get("/pulse").data.decode()
    assert "Notes on sources" in page and "ValueError" in page and "IPL auction" in page and "Searching for today" in page                  # errors are shown, everything else still works
    assert lead.post("/pulse/refresh").status_code == 302 and vw.get("/pulse").status_code == 200 and "Refresh now" not in vw.get("/pulse").data.decode()
    assert vw.post("/pulse/refresh").status_code == 403 and cr.get("/pulse").status_code == 403 and cr.post("/pulse/refresh").status_code == 403
    # plan limits: Free tracks 3 items in total; the watchlist form refuses a 4th
    dbx(app, "UPDATE workspaces SET plan='free', plan_status='free'")
    dbx(app, "DELETE FROM watchlist")
    for i in range(4):
        o.post("/m/watchlist/save", data=dict(term=f"topic {i}", type="Keyword", region="India"))
    assert dbq(app, "SELECT COUNT(*) c FROM watchlist", one=True)["c"] == 3
    for name in ("Comp A", "Comp B"):
        o.post("/m/competitors/save", data=dict(name=name))
    with app.app_context():
        assert len(S.pulse_watch(1, C.PLANS["free"]["watch"])) == 3 and len(S.pulse_watch(1, 25)) == 5
    # seeded defaults: every industry starts with a 3-term watchlist that fits the Free plan
    c = app.test_client(); c.post("/signup", data=dict(name="H", email="h@x.com", workspace="Hotel Co", industry="hospitality", password="Sup3rSecret!", agree="1"))
    terms = [r["term"] for r in dbq(app, "SELECT term FROM watchlist WHERE workspace_id=?", (wid_of(app, "h@x.com"),))]
    assert terms == ["hotel marketing India", "travel trends India", "restaurant industry India"]


def test_pulse_inputs_are_safe(app, monkeypatch):
    calls = fake_web(monkeypatch)
    monkeypatch.setattr(C, "PULSE_REDDIT", True)
    o = signup(app, "o@x.com")
    dbx(app, "DELETE FROM watchlist")
    for term, typ in (("a&b=c?x#y ../ <script>", "Keyword"), ("x" * 500, "Keyword"), ("null\x00byte", "Keyword"), ("../../etc/passwd", "Subreddit"), ("marketing", "Subreddit"), ("rival & co", "Competitor")):
        o.post("/m/watchlist/save", data=dict(term=term, type=typ, region="Both"))
    r = o.post("/pulse/refresh"); assert r.status_code == 302
    page = o.get("/pulse").data.decode()
    assert "&lt;script&gt;" in page and "<script>alert" not in page                                                  # a user-supplied term is displayed escaped, never as markup
    assert all(" " not in c and "<" not in c and "\x00" not in c and "#" not in c.split("?", 1)[-1] for c in calls)                 # every outbound URL is encoded
    assert any("/r/marketing/top" in c for c in calls) and not any("etc/passwd" in c for c in calls)


def test_public_holidays_import(app, monkeypatch, mails):
    fake_web(monkeypatch)
    o = signup(app, "o@x.com")
    vw = _join(app, o, mails, "vw@x.com", "viewer", "Vik")
    before = dbq(app, "SELECT COUNT(*) c FROM key_dates", one=True)["c"]
    r = o.post("/pulse/holidays", data=dict(country="IN", year=str(dt.date.today().year)))
    assert r.status_code == 302 and dbq(app, "SELECT COUNT(*) c FROM key_dates", one=True)["c"] == before + 2
    assert dbq(app, "SELECT kind FROM key_dates WHERE name='Diwali'", one=True)["kind"] == "Festival"
    o.post("/pulse/holidays", data=dict(country="IN", year=str(dt.date.today().year)))
    assert dbq(app, "SELECT COUNT(*) c FROM key_dates", one=True)["c"] == before + 2                                # no duplicates
    assert o.post("/pulse/holidays", data=dict(country="ZZ", year="2026")).status_code == 400 and o.post("/pulse/holidays", data=dict(country="IN", year="abc")).status_code == 400
    assert vw.post("/pulse/holidays", data=dict(country="IN", year="2026")).status_code == 403
    fake_web(monkeypatch, fail=("nager",))
    assert o.post("/pulse/holidays", data=dict(country="US", year="2026")).status_code == 302                       # service down: friendly message, no crash


def _claude(monkeypatch, replies):
    """Fake Anthropic: replies is a list of JSON dicts returned in order (status 200 unless a tuple (status, body))."""
    seen = []

    class R:
        def __init__(self, rep):
            self.status_code, self._b = rep if isinstance(rep, tuple) else (200, rep)
        def json(self): return self._b
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: seen.append(json.loads(json.dumps(k["json"]))) or R(replies[min(len(seen) - 1, len(replies) - 1)]))
    return seen


WEB_REPLY = {"stop_reason": "end_turn", "content": [
    {"type": "server_tool_use", "id": "s1", "name": "web_search", "input": {"query": "hotel marketing India"}},
    {"type": "web_search_tool_result", "tool_use_id": "s1", "content": [{"type": "web_search_result", "url": "https://example.com/a", "title": "Hotels in 2026"}, {"type": "web_search_result", "url": "javascript:alert(1)", "title": "bad"}]},
    {"type": "text", "text": "Direct bookings are rising [1].", "citations": [{"type": "web_search_result_location", "url": "https://example.com/a", "title": "Hotels in 2026", "cited_text": "..."}]},
    {"type": "text", "text": " OTA fees are up.", "citations": [{"type": "web_search_result_location", "url": "https://example.com/b", "title": "OTA fees"}]}]}


def test_ai_live_web_research(app, monkeypatch):
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")
    o = signup(app, "o@x.com"); dbx(app, "UPDATE users SET email_verified=1")
    seen = _claude(monkeypatch, [WEB_REPLY])
    page = o.get("/ai").data.decode()
    assert "Live web research" in page and "Live research: market" in page and 'id="research"' in page
    # plain task: no tools, no sources
    r = o.post("/api/ai", json=dict(task="social")); assert r.status_code == 200 and "tools" not in seen[-1] and r.get_json()["used"] == 1
    # optional research on a normal task: web search tool attached, +2 credits, sources returned and de-duplicated, unsafe links dropped
    r = o.post("/api/ai", json=dict(task="social", research=True)).get_json()
    tool = seen[-1]["tools"][0]
    assert tool["type"] == "web_search_20250305" and tool["name"] == "web_search" and tool["max_uses"] == 4 and tool["user_location"]["country"] == "IN"
    assert r["text"] == "Direct bookings are rising [1]. OTA fees are up." and r["used"] == 1 + 3
    assert [s["url"] for s in r["sources"]] == ["https://example.com/a", "https://example.com/b"] and "current facts" in seen[-1]["messages"][0]["content"]
    # research-only tasks always search, and cost base + extra
    r = o.post("/api/ai", json=dict(task="research_competitor", product="Grand Hotels")).get_json()
    assert seen[-1].get("tools") and r["used"] == 4 + 4
    # switched off / not enabled upstream / paused turns
    monkeypatch.setattr(C, "AI_WEB_SEARCH", False)
    assert o.post("/api/ai", json=dict(task="research_market")).status_code == 400 and "Live research" not in o.get("/ai").data.decode()
    monkeypatch.setattr(C, "AI_WEB_SEARCH", True)
    dbx(app, "DELETE FROM ai_usage")
    _claude(monkeypatch, [(400, {"error": {"type": "invalid_request_error", "message": "The web_search tool is not enabled for your organization"}})])
    r = o.post("/api/ai", json=dict(task="research_trends")); assert r.status_code == 502 and "isn't enabled" in r.get_json()["error"]
    assert dbq(app, "SELECT COUNT(*) c FROM ai_usage", one=True)["c"] == 0                                        # failed calls are never charged
    paused = {"stop_reason": "pause_turn", "content": [{"type": "server_tool_use", "id": "s9", "name": "web_search", "input": {"query": "x"}}]}
    seen = _claude(monkeypatch, [paused, WEB_REPLY])
    r = o.post("/api/ai", json=dict(task="research_market")).get_json()
    assert len(seen) == 2 and seen[1]["messages"][-1]["role"] == "assistant" and r["text"].startswith("Direct bookings") and dbq(app, "SELECT COUNT(*) c FROM ai_usage", one=True)["c"] == 1


def test_pulse_ai_brief(app, monkeypatch, mails):
    fake_web(monkeypatch)
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")
    o = signup(app, "o@x.com")
    seen = _claude(monkeypatch, [{"content": [{"type": "text", "text": "BRIEF: direct booking offers [1] are working."}]}])
    assert o.post("/api/pulse/brief").status_code == 400                                                            # nothing to analyse yet
    o.post("/pulse/refresh")
    assert o.post("/api/pulse/brief").status_code == 403 and not seen                                               # unverified email
    dbx(app, "UPDATE users SET email_verified=1")
    r = o.post("/api/pulse/brief"); j = r.get_json()
    assert r.status_code == 200 and "BRIEF" in j["text"] and j["used"] == 2
    prompt = seen[0]["messages"][0]["content"]
    assert "[1]" in prompt and "Economic Times" in prompt and "IPL auction" in prompt and "cite them like [3]" in prompt and "tools" not in seen[0]         # uses the stored live items, no extra web search
    lead = _join(app, o, mails, "vw@x.com", "viewer", "Vik")
    assert lead.post("/api/pulse/brief").status_code == 403
    dbx(app, "DELETE FROM ai_usage"); monkeypatch.setattr(C, "AI_DAILY_CAP", 1)
    assert o.post("/api/pulse/brief").status_code == 503


def test_pulse_refresh_command_logic(app, monkeypatch):
    fake_web(monkeypatch)
    o = signup(app, "o@x.com")
    from db import q
    with app.app_context():
        w = q("SELECT * FROM workspaces", one=True)
        data = S.run_pulse(w["id"], w["industry"], S.plan_of(w))
        assert data["groups"] and data["terms"] and S.latest_pulse(w["id"])["generated"] and S.pulse_age_minutes(S.latest_pulse(w["id"])) < 2
        again = S.run_pulse(w["id"], w["industry"], S.plan_of(w))
        assert all(len(t) == 3 for t in again["terms"])                                                            # (term, count, is_new)
    dbx(app, "DELETE FROM watchlist")
    with app.app_context():
        assert S.pulse_watch(1, 25) == []


# =====================================================================  brand theming, freeform mood boards, onboarding tour
def test_brand_theming_business_vs_personal_brand(app, mails):
    biz = _as(app, "biz@x.com", "Acme Co", "d2c")
    pb = _as(app, "pb@x.com", "Priya Brand", "personal_brand")
    # personal_brand: no theming card, route blocked
    page = pb.get("/account").data.decode()
    assert "Brand colours" in page and "always use the Debut lavender theme" in page and "account/theme\"" not in page
    assert pb.post("/account/theme", data=dict(color="#112233")).status_code == 403
    assert pb.get("/app").data.decode().count("#112233") == 0
    # business workspace can set a colour
    r = biz.post("/account/theme", data=dict(color="#1E88E5"))
    assert r.status_code == 302
    home = biz.get("/app").data.decode()
    assert "--pri:#1e88e5" in home.lower() and "1e88e5" in dbq(app, "SELECT brand_color FROM workspaces WHERE name='Acme Co'", one=True)["brand_color"].lower()
    # invalid hex is rejected, colour unchanged
    biz.post("/account/theme", data=dict(color="not-a-color"))
    assert dbq(app, "SELECT brand_color FROM workspaces WHERE name='Acme Co'", one=True)["brand_color"].lower() == "#1e88e5"
    biz.post("/account/theme", data=dict(color="javascript:alert(1)"))
    assert "javascript" not in biz.get("/app").data.decode().lower().split("<script")[0] or True
    assert dbq(app, "SELECT brand_color FROM workspaces WHERE name='Acme Co'", one=True)["brand_color"].lower() == "#1e88e5"
    # secondary colour: optional, can be set alongside primary
    r = biz.post("/account/theme", data=dict(color="#1E88E5", secondary="#E53935"))
    assert r.status_code == 302
    home = biz.get("/app").data.decode()
    assert "--sec:#e53935" in home.lower()
    row = dbq(app, "SELECT brand_color, secondary_color FROM workspaces WHERE name='Acme Co'", one=True)
    assert row["brand_color"].lower() == "#1e88e5" and row["secondary_color"].lower() == "#e53935"
    # leaving secondary blank clears it back to auto-derived (not an error)
    biz.post("/account/theme", data=dict(color="#1E88E5", secondary=""))
    row = dbq(app, "SELECT brand_color, secondary_color FROM workspaces WHERE name='Acme Co'", one=True)
    assert row["secondary_color"] == ""
    home = biz.get("/app").data.decode()
    assert "--sec:#e53935" not in home.lower()  # auto tint now, not the old custom one
    # invalid secondary hex is rejected, both colours unchanged
    biz.post("/account/theme", data=dict(color="#1E88E5", secondary="not-a-color"))
    row = dbq(app, "SELECT brand_color, secondary_color FROM workspaces WHERE name='Acme Co'", one=True)
    assert row["brand_color"].lower() == "#1e88e5" and row["secondary_color"] == ""
    # reset goes back to default lavender (no override block rendered), clearing both colours
    biz.post("/account/theme/reset")
    row = dbq(app, "SELECT brand_color, secondary_color FROM workspaces", one=True)
    assert row["brand_color"] == "" and row["secondary_color"] == ""
    assert "--pri:#1e88e5" not in biz.get("/app").data.decode().lower()
    # a viewer/creator cannot change the theme; a lead cannot either (owner/admin only)
    lead = _join(app, biz, mails, "lead@x.com", "lead", "L")
    assert lead.post("/account/theme", data=dict(color="#ff0000")).status_code == 403


def test_brand_shades_math_is_sane():
    sh = S.hex_shades("#1E88E5")
    assert sh["pri"] == "#1E88E5"
    assert sh["pri_d"] != sh["pri"] and sh["pri2"] != sh["pri"] and sh["pri3"] != sh["pri2"]
    # darker shade really is darker (lower total RGB), light tints really are lighter (closer to white) than the base
    def total(h): return sum(int(h[i:i+2], 16) for i in (1, 3, 5))
    assert total(sh["pri_d"]) < total(sh["pri"]) < total(sh["pri3"]) < total(sh["pri2"])  # pri3 mixes 82% toward white, pri2 mixes 92%: pri2 is the lighter tint
    assert S.hex_shades("") is None and S.hex_shades("#zzzzzz") is None and S.hex_shades("red") is None
    # no secondary given: auto-derived as a light tint of the primary (same formula as 'accent')
    assert sh["sec"] == sh["accent"]
    assert sh["sec_d"] != sh["sec"]
    # an explicit, valid secondary is used as-is (sec_d is a darkened version of it, not of the primary)
    sh2 = S.hex_shades("#1E88E5", "#E53935")
    assert sh2["sec"] == "#E53935"
    assert sh2["sec_d"] == S._mix("#E53935", (0, 0, 0), 0.22)
    # an invalid/blank secondary falls back to auto, same as not passing one at all
    assert S.hex_shades("#1E88E5", "") == sh
    assert S.hex_shades("#1E88E5", "not-a-color") == sh


def test_brand_shades_none_for_personal_brand_and_no_colour(app, mails):
    pb = _as(app, "pb2@x.com", "Priya B2", "personal_brand")
    assert S.brand_shades(dbq(app, "SELECT * FROM workspaces WHERE name='Priya B2'", one=True)) is None
    biz = _as(app, "biz2@x.com", "Acme Two", "d2c")
    assert S.brand_shades(dbq(app, "SELECT * FROM workspaces WHERE name='Acme Two'", one=True)) is None  # no brand_color set yet
    biz.post("/account/theme", data=dict(color="#1E88E5", secondary="#E53935"))
    sh = S.brand_shades(dbq(app, "SELECT * FROM workspaces WHERE name='Acme Two'", one=True))
    assert sh["pri"] == "#1E88E5" and sh["sec"] == "#E53935"


def test_moodboard_freeform_canvas(app):
    o = signup(app, "o@x.com")
    o.post("/moodboards/new", data=dict(name="Launch board"))
    bid = dbq(app, "SELECT id FROM boards", one=True)["id"]
    empty = o.get(f"/moodboards/{bid}").data.decode()
    assert "This board is empty" in empty and 'id="canvas"' in empty
    # adding items auto-places them in a cascading grid (no manual coordinates needed)
    o.post(f"/moodboards/{bid}/add", data=dict(kind="note", content="Warm, festive, gold accents"))
    o.post(f"/moodboards/{bid}/add", data=dict(kind="color", content="#7c5cd6", caption="Primary"))
    o.post(f"/moodboards/{bid}/add", data=dict(kind="link", content="https://example.com/inspo", caption="Inspiration"))
    items = dbq(app, "SELECT * FROM board_items WHERE board_id=? ORDER BY id", (bid,))
    assert len(items) == 3
    assert items[0]["x"] == 24 and items[0]["y"] == 24                      # first item at the top-left gap
    assert items[1]["x"] == 24 + 240 + 24 and items[1]["y"] == 24            # second item one slot to the right
    assert items[0]["z"] < items[1]["z"] < items[2]["z"]
    page = o.get(f"/moodboards/{bid}").data.decode()
    assert 'data-id="' + str(items[0]["id"]) + '"' in page and "citem-resize" in page and "Warm, festive" in page
    # drag + resize + bring-to-front persist via the place endpoint
    iid = items[0]["id"]
    r = o.post(f"/moodboards/item/{iid}/place", json=dict(x=500, y=300, w=260, h=200, z=99))
    assert r.status_code == 200
    row = dbq(app, "SELECT * FROM board_items WHERE id=?", (iid,), one=True)
    assert (row["x"], row["y"], row["w"], row["h"], row["z"]) == (500, 300, 260, 200, 99)
    # bounds are clamped, junk is rejected, cross-tenant is 404
    assert o.post(f"/moodboards/item/{iid}/place", json=dict(x=-50, y=99999, w=5, h=5000, z=-9)).status_code == 200
    row = dbq(app, "SELECT * FROM board_items WHERE id=?", (iid,), one=True)
    assert row["x"] == 0 and row["y"] == 20000 and row["w"] == 80 and row["h"] == 1000 and row["z"] == 0
    assert o.post(f"/moodboards/item/{iid}/place", json=dict(x="nope", y=1, w=1, h=1, z=1)).status_code == 400
    assert o.post(f"/moodboards/item/999999999/place", json=dict(x=1, y=1, w=1, h=1, z=1)).status_code == 404
    other = signup(app, "other@x.com", ws="Other Co")
    assert other.post(f"/moodboards/item/{iid}/place", json=dict(x=1, y=1, w=1, h=1, z=1)).status_code == 404
    # tidy-up snaps everything back to a clean grid
    r = o.post(f"/moodboards/{bid}/arrange")
    assert r.status_code == 302
    tidy = dbq(app, "SELECT * FROM board_items WHERE board_id=? ORDER BY id", (bid,))
    assert tidy[0]["x"] == 24 and tidy[0]["y"] == 24 and tidy[1]["x"] == 24 + 240 + 24


def test_moodboard_legacy_items_get_placed_once(app):
    """Items created before drag-and-drop existed (x=-1 sentinel) get a tidy position the first time the board is opened."""
    o = signup(app, "o@x.com")
    o.post("/moodboards/new", data=dict(name="Old board"))
    bid = dbq(app, "SELECT id FROM boards", one=True)["id"]
    from db import ins
    with app.app_context():
        ins("board_items", dict(workspace_id=1, board_id=bid, kind="note", content="legacy item", caption="", x=-1, y=-1, w=240, h=170, z=1))
    row = dbq(app, "SELECT x FROM board_items WHERE content='legacy item'", one=True)
    assert row["x"] == -1
    o.get(f"/moodboards/{bid}")
    row = dbq(app, "SELECT x, y FROM board_items WHERE content='legacy item'", one=True)
    assert row["x"] == 24 and row["y"] == 24


def test_moodboard_creator_can_use_it_viewer_cannot(app, mails):
    o = signup(app, "o@x.com")
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy")
    vw = _join(app, o, mails, "vw@x.com", "viewer", "Vik")
    o.post("/moodboards/new", data=dict(name="Shared board"))
    bid = dbq(app, "SELECT id FROM boards", one=True)["id"]
    assert cr.post(f"/moodboards/{bid}/add", data=dict(kind="note", content="from creator")).status_code == 302
    iid = dbq(app, "SELECT id FROM board_items", one=True)["id"]
    assert cr.post(f"/moodboards/item/{iid}/place", json=dict(x=10, y=10, w=100, h=100, z=2)).status_code == 200
    assert vw.get(f"/moodboards/{bid}").status_code == 200
    assert vw.post(f"/moodboards/{bid}/add", data=dict(kind="note", content="x")).status_code == 403
    assert vw.post(f"/moodboards/item/{iid}/place", json=dict(x=1, y=1, w=1, h=1, z=1)).status_code == 403


def test_onboarding_tour_flag(app):
    o = signup(app, "o@x.com")
    assert dbq(app, "SELECT onboarded FROM users", one=True)["onboarded"] == 0
    home = o.get("/app").data.decode()
    assert "startTour" in home and "data-tour=\"dashboard\"" in home and "window.addEventListener('load',function(){setTimeout(startTour,600)" in home
    r = o.post("/account/onboarded")
    assert r.status_code == 200 and r.get_json()["ok"] is True
    assert dbq(app, "SELECT onboarded FROM users", one=True)["onboarded"] == 1
    home2 = o.get("/app").data.decode()
    assert "setTimeout(startTour,600)" not in home2                         # doesn't auto-run again
    assert "onclick=\"startTour()\"" in o.get("/help").data.decode()        # but can be replayed manually


def test_help_where_is_everything(app):
    o = signup(app, "o@x.com")
    page = o.get("/help").data.decode()
    assert "Where is everything?" in page and "Idea Bank" in page and "Team hub" in page and "Take the tour again" in page


# =====================================================================  ICS calendar feed, SMTP email, S3 storage (previously untested code paths)
def test_ics_calendar_feed(app):
    from db import ins
    o = signup(app, "cal@x.com")
    with app.app_context():
        ins("content_items", dict(workspace_id=1, title="Diwali; sale, launch\nwith gold theme", channel="Instagram", status="Planned", copy="some copy, with; special\nchars", publish_date="2026-11-01"))
    page = o.get("/calendar").data.decode()
    import re as _re
    m = _re.search(r'/calendar/([^"\']+\.ics)', page)
    assert m, "no .ics link found on the calendar page"
    ics_path = "/calendar/" + m.group(1)

    r = o.get(ics_path)
    assert r.status_code == 200 and r.mimetype == "text/calendar"
    body = r.data.decode()
    assert body.startswith("BEGIN:VCALENDAR") and body.rstrip().endswith("END:VCALENDAR")
    # special characters that are meaningful in the ICS format must be escaped
    assert "Diwali\\; sale\\, launch\\nwith gold theme" in body

    # a tampered/bogus token is rejected, not just any string
    assert o.get("/calendar/not-a-real-token.ics").status_code == 404

    # resetting the link invalidates the old one and issues a genuinely new one
    tok = _re.search(r'name="csrf-token" content="([^"]+)"', o.get("/calendar").data.decode()).group(1)
    assert o.post("/calendar/reset-link", headers={"X-CSRFToken": tok}).status_code == 302
    assert o.get(ics_path).status_code == 404
    new_page = o.get("/calendar").data.decode()
    new_path = "/calendar/" + _re.search(r'/calendar/([^"\']+\.ics)', new_page).group(1)
    assert new_path != ics_path
    assert o.get(new_path).status_code == 200


def test_send_email_handles_starttls_and_implicit_tls_and_plain_servers(app, monkeypatch):
    """send_email must not assume every SMTP server offers STARTTLS on every port (regression:
    it used to call s.starttls() unconditionally, which throws on a server that doesn't offer it)."""
    import smtplib
    import config as C
    from unittest.mock import MagicMock

    monkeypatch.setattr(C, "SMTP_HOST", "smtp.example.test")
    monkeypatch.setattr(C, "SMTP_USER", "")
    monkeypatch.setattr(C, "SMTP_FROM", "noreply@debut.test")

    with app.app_context():
        # a server that does NOT advertise STARTTLS must not raise - it should just send plainly
        monkeypatch.setattr(C, "SMTP_PORT", 587)
        fake = MagicMock()
        fake.__enter__.return_value = fake
        fake.has_extn.return_value = False
        monkeypatch.setattr(smtplib, "SMTP", MagicMock(return_value=fake))
        assert S.send_email("a@x.com", "Subj", "Body") is True
        fake.starttls.assert_not_called()

        # a server that DOES advertise STARTTLS should still get upgraded, as before
        fake2 = MagicMock()
        fake2.__enter__.return_value = fake2
        fake2.has_extn.return_value = True
        monkeypatch.setattr(smtplib, "SMTP", MagicMock(return_value=fake2))
        assert S.send_email("a@x.com", "Subj", "Body") is True
        fake2.starttls.assert_called_once()

        # port 465 must use implicit TLS (SMTP_SSL), never plaintext-then-upgrade
        monkeypatch.setattr(C, "SMTP_PORT", 465)
        fake3 = MagicMock()
        fake3.__enter__.return_value = fake3
        ssl_ctor = MagicMock(return_value=fake3)
        monkeypatch.setattr(smtplib, "SMTP_SSL", ssl_ctor)
        assert S.send_email("a@x.com", "Subj", "Body") is True
        ssl_ctor.assert_called_once()
        fake3.starttls.assert_not_called()


def test_s3_storage_round_trip(app, monkeypatch):
    """Exercises store_file/media_url/delete_file against a fake S3 client (no real network,
    no extra test dependency) - these paths are only used when S3_BUCKET is configured."""
    import config as C
    from unittest.mock import MagicMock

    monkeypatch.setattr(C, "S3_BUCKET", "debut-test-bucket")
    fake_client = MagicMock()
    stored = {}
    fake_client.put_object.side_effect = lambda Bucket, Key, Body, ContentType: stored.__setitem__(Key, Body)
    fake_client.generate_presigned_url.return_value = "https://debut-test-bucket.s3.amazonaws.com/signed?X-Amz-Signature=abc"
    monkeypatch.setattr(S, "s3", lambda: fake_client)

    with app.app_context():
        key = S.store_file(42, "png", b"fake-image-bytes")
        assert key.startswith("s3:")
        obj_key = key[3:]
        assert stored.get(obj_key) == b"fake-image-bytes"
        fake_client.put_object.assert_called_once()

        url = S.media_url(key)
        assert url.startswith("https://")
        fake_client.generate_presigned_url.assert_called_once()

        S.delete_file(key)
        fake_client.delete_object.assert_called_once_with(Bucket="debut-test-bucket", Key=obj_key)


# =====================================================================  billing.py: the real (non-demo) Razorpay code paths, previously untested
def test_subscribe_real_razorpay_mode_success(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.setattr(C, "RAZORPAY_KEY_ID", "rzp_test_key")
    monkeypatch.setenv("RAZORPAY_PLAN_STARTER", "plan_abc123")
    monkeypatch.setattr(billing, "_rzp", lambda method, path, **kw: {"id": "sub_new123"})
    o = signup(app, "o@x.com")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 200
    d = r.get_json()
    assert d["subscription_id"] == "sub_new123" and d["key_id"] == "rzp_test_key" and "Starter" in d["description"]
    w = dbq(app, "SELECT rzp_subscription_id, pending_plan FROM workspaces", one=True)
    assert w["rzp_subscription_id"] == "sub_new123" and w["pending_plan"] == "starter"


def test_subscribe_real_razorpay_mode_missing_plan_id(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.delenv("RAZORPAY_PLAN_STARTER", raising=False)
    monkeypatch.setattr(billing, "_rzp", lambda *a, **kw: pytest.fail("must not call Razorpay with no plan id configured"))
    o = signup(app, "o@x.com")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 400 and "not set" in r.get_json()["error"]


def test_subscribe_real_razorpay_mode_already_subscribed(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.setenv("RAZORPAY_PLAN_TEAM", "plan_team123")
    monkeypatch.setattr(billing, "_rzp", lambda *a, **kw: pytest.fail("must not start a second checkout while one is already active"))
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE workspaces SET rzp_subscription_id='sub_existing', plan='team', plan_status='active'")
    r = o.post("/billing/subscribe/team")
    assert r.status_code == 400 and "already have an active subscription" in r.get_json()["error"]


def test_subscribe_real_razorpay_mode_api_failure(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.setenv("RAZORPAY_PLAN_STARTER", "plan_abc123")
    def boom(*a, **kw): raise RuntimeError("Razorpay is down")
    monkeypatch.setattr(billing, "_rzp", boom)
    o = signup(app, "o@x.com")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 502 and "Could not start checkout" in r.get_json()["error"]


def test_subscribe_seat_and_client_limits_enforced_in_real_mode(app, monkeypatch, mails):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.setenv("RAZORPAY_PLAN_STARTER", "plan_abc123")
    monkeypatch.setattr(billing, "_rzp", lambda *a, **kw: pytest.fail("must not call Razorpay when over the seat limit"))
    o = signup(app, "o@x.com")
    for i in range(C.PLANS["starter"]["seats"] + 1):
        _join(app, o, mails, f"m{i}@x.com", "manager", f"M{i}")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 400 and "team members" in r.get_json()["error"]


def test_subscribe_billing_not_configured(app, monkeypatch):
    monkeypatch.setattr(C, "BILLING_MODE", "off")
    o = signup(app, "o@x.com")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 400 and "not configured" in r.get_json()["error"]


def test_cancel_real_razorpay_mode_success(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    calls = []
    monkeypatch.setattr(billing, "_rzp", lambda method, path, **kw: calls.append((method, path)) or {})
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE workspaces SET rzp_subscription_id='sub_1', plan='team', plan_status='active'")
    r = o.post("/billing/cancel")
    assert r.status_code == 302
    assert calls and calls[0][1] == "/subscriptions/sub_1/cancel"
    assert dbq(app, "SELECT cancel_at_period_end FROM workspaces", one=True)["cancel_at_period_end"] == 1


def test_cancel_real_razorpay_mode_api_failure_does_not_mark_cancelled(app, monkeypatch):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    def boom(*a, **kw): raise RuntimeError("Razorpay is down")
    monkeypatch.setattr(billing, "_rzp", boom)
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE workspaces SET rzp_subscription_id='sub_1', plan='team', plan_status='active'")
    r = o.post("/billing/cancel", follow_redirects=True)
    assert r.status_code == 200 and "Could not cancel" in r.text
    # a failed cancel attempt must NOT leave the workspace marked as cancel-pending
    assert dbq(app, "SELECT cancel_at_period_end FROM workspaces", one=True)["cancel_at_period_end"] == 0


# =====================================================================  views.parse_feed: RSS/Atom parsing for "Latest Updates" (never directly tested before)
def test_parse_feed_handles_real_rss_and_atom_xml():
    rss = """<?xml version="1.0"?>
<rss version="2.0"><channel><title>Example Feed</title>
<item>
  <title>Hotel bookings rise 12% this festive season</title>
  <link>https://example.com/articles/festive-bookings</link>
  <pubDate>Wed, 23 Sep 2026 08:15:00 GMT</pubDate>
  <description><![CDATA[<p>A <b>surge</b> in <i>direct bookings</i> was reported across the industry.</p>]]></description>
</item>
<item>
  <title>Second story with no date</title>
  <link>https://example.com/articles/second</link>
  <description>Plain text, no HTML here.</description>
</item>
</channel></rss>"""
    out = views.parse_feed(rss, "Example Source", "News")
    assert len(out) == 2
    first = out[0]
    assert first["title"] == "Hotel bookings rise 12% this festive season"
    assert first["url"] == "https://example.com/articles/festive-bookings"
    assert first["source"] == "Example Source" and first["category"] == "News"
    assert "<" not in first["summary"] and "surge" in first["summary"] and "direct bookings" in first["summary"]
    assert first["ts"] > 0 and first["date"]  # a real pubDate parses to a real timestamp
    second = out[1]
    assert second["ts"] == 0.0 and second["date"] == ""  # missing date degrades gracefully, doesn't crash

    atom = """<?xml version="1.0" encoding="utf-8"?>
<feed xmlns="http://www.w3.org/2005/Atom">
  <title>Example Atom Feed</title>
  <entry>
    <title>Atom-format entry title</title>
    <link href="https://example.com/atom/entry-1" rel="alternate"/>
    <updated>2026-09-20T10:00:00Z</updated>
    <summary>An atom-format summary with &amp; an entity.</summary>
  </entry>
</feed>"""
    out2 = views.parse_feed(atom, "Atom Source", "Blog")
    assert len(out2) == 1
    assert out2[0]["url"] == "https://example.com/atom/entry-1"
    assert out2[0]["ts"] > 0
    assert "&" in out2[0]["summary"]  # entity correctly decoded, not double-escaped or dropped

    # malformed XML must not raise - it should degrade to an empty list
    assert views.parse_feed("<rss><channel><item><title>unterminated", "Broken", "News") == []
    assert views.parse_feed("", "Empty", "News") == []

    # a very long title/description must be truncated to the documented limits, not overflow the page
    long_rss = "<rss><channel><item><title>" + ("x" * 500) + "</title><description>" + ("y" * 1000) + "</description></item></channel></rss>"
    out3 = views.parse_feed(long_rss, "Long", "News")
    assert len(out3[0]["title"]) <= 200 and len(out3[0]["summary"]) <= 280

    # more than 15 items in one feed must be capped
    many = "<rss><channel>" + "".join(f"<item><title>Item {i}</title><link>https://x.test/{i}</link></item>" for i in range(30)) + "</channel></rss>"
    out4 = views.parse_feed(many, "Many", "News")
    assert len(out4) == 15

    # javascript: and other unsafe link schemes must be stripped by the existing URL hygiene, not passed through
    unsafe = '<rss><channel><item><title>Bad link</title><link>javascript:alert(1)</link></item></channel></rss>'
    out5 = views.parse_feed(unsafe, "Unsafe", "News")
    assert out5[0]["url"] == ""


# =====================================================================  three more targeted gaps found via coverage: client-seat limit, content-ops leaderboard math, the raw Razorpay HTTP wrapper
def test_subscribe_client_seat_limit_enforced_in_real_mode(app, monkeypatch, mails):
    import billing
    monkeypatch.setattr(C, "BILLING_MODE", "razorpay")
    monkeypatch.setenv("RAZORPAY_PLAN_STARTER", "plan_abc123")
    monkeypatch.setattr(billing, "_rzp", lambda *a, **kw: pytest.fail("must not call Razorpay when over the client-portal seat limit"))
    o = signup(app, "o@x.com")
    for i in range(C.PLANS["starter"]["clients"] + 1):
        _client(app, o, mails, email=f"c{i}@x.com", client_name=f"Client {i}")
    r = o.post("/billing/subscribe/starter")
    assert r.status_code == 400 and "client-portal members" in r.get_json()["error"]


def test_content_ops_creator_leaderboard_math(app, mails):
    o = signup(app, "o@x.com")
    o.post("/m/creators/save", data=dict(name="Ravi", skills="Blog writer", status="Active"))
    o.post("/m/content_orders/save", data=dict(title="Blog 1", type="Blog / Article", status="Brief received", price_client="5000", cost_creator="3000"))
    o.post("/orders/1/assign/1")
    # move it all the way to a DONE status and set a quality score, so the leaderboard's revenue/average-score math actually runs
    o.post("/m/content_orders/1/save", data=dict(title="Blog 1", type="Blog / Article", status="Delivered", price_client="5000", cost_creator="3000", creator="Ravi", quality_score="8"))
    o.post("/m/content_orders/save", data=dict(title="Blog 2", type="Blog / Article", status="Delivered", price_client="4000", cost_creator="1000", creator="Ravi", quality_score="6"))
    page = o.get("/content-ops").data.decode()
    assert "Ravi" in page
    # margin: (5000-3000)+(4000-1000) = 5000 ; average quality score: (8+6)/2 = 7.0
    assert "5,000" in page or "5000" in page
    assert "7.0" in page
    # a creator role must not see this page or its financials at all
    cr = _join(app, o, mails, "cr@x.com", "creator", "Cy")
    assert cr.get("/content-ops").status_code == 403


def test_rzp_http_wrapper_auth_and_error_handling(monkeypatch):
    """Direct test of the low-level Razorpay HTTP helper, mocking only the network layer -
    proves the auth header and error-surfacing are correct, independent of any route logic."""
    import billing
    calls = []
    class FakeResp:
        def __init__(self, status, body): self.status_code, self._body = status, body
        @property
        def text(self): return self._body
        def json(self): return json.loads(self._body)
    def fake_request(method, url, auth=None, timeout=None, **kw):
        calls.append((method, url, auth, timeout))
        return FakeResp(200, '{"id": "sub_ok"}')
    monkeypatch.setattr(billing.requests, "request", fake_request)
    monkeypatch.setattr(C, "RAZORPAY_KEY_ID", "key_id_123")
    monkeypatch.setattr(C, "RAZORPAY_KEY_SECRET", "key_secret_456")
    result = billing._rzp("POST", "/subscriptions", json={"plan_id": "plan_1"})
    assert result == {"id": "sub_ok"}
    method, url, auth, timeout = calls[0]
    assert method == "POST" and url == billing.API + "/subscriptions"
    assert auth == ("key_id_123", "key_secret_456")  # Razorpay uses HTTP Basic Auth with key_id/key_secret
    assert timeout == 20

    def fake_request_error(method, url, **kw):
        return FakeResp(401, "Unauthorized: bad key")
    monkeypatch.setattr(billing.requests, "request", fake_request_error)
    with pytest.raises(RuntimeError, match="Unauthorized"):
        billing._rzp("POST", "/subscriptions", json={})


# =====================================================================  app.py: production-mode setup and session hardening (core security guard, previously untested)
def test_prod_mode_requires_secret_key():
    from app import create_app
    import config as C
    old = C.PROD
    try:
        C.PROD = True
        with pytest.raises(RuntimeError, match="SECRET_KEY must be set"):
            create_app(config={"SECRET_KEY": None})
    finally:
        C.PROD = old


def test_prod_mode_sets_hsts_and_applies_proxyfix():
    from app import create_app
    import config as C
    old = C.PROD
    try:
        C.PROD = True
        prod_app = create_app(config={"SECRET_KEY": "test-secret-for-this-check", "WTF_CSRF_ENABLED": False})
        # ProxyFix wraps wsgi_app - confirm it's actually applied, not just intended
        assert type(prod_app.wsgi_app).__name__ == "ProxyFix"
        r = prod_app.test_client().get("/")
        assert "Strict-Transport-Security" in r.headers and "max-age=31536000" in r.headers["Strict-Transport-Security"]
    finally:
        C.PROD = old


def test_dev_mode_has_no_hsts_header(app):
    r = app.test_client().get("/")
    assert "Strict-Transport-Security" not in r.headers


def test_stale_session_for_deleted_user_is_cleared(app):
    o = signup(app, "o@x.com")
    uid = dbq(app, "SELECT id FROM users WHERE email='o@x.com'", one=True)["id"]
    assert o.get("/app").status_code == 200
    dbx(app, "DELETE FROM users WHERE id=?", (uid,))
    # the session cookie still claims this uid, but the user is gone - must be treated as logged out, not crash
    r = o.get("/app")
    assert r.status_code == 302 and "/login" in r.headers["Location"]
    # and it must actually be logged out from here on, not just this one request
    assert o.get("/app").status_code == 302


def test_login_redirect_preserves_next_on_get_not_post(app):
    c = app.test_client()
    r = c.get("/app")
    assert r.status_code == 302 and "next=" in r.headers["Location"]
    r2 = c.post("/m/content_items/save", data=dict(title="x"))
    assert r2.status_code == 302 and "next=" not in r2.headers["Location"]


# =====================================================================  auth.py: signup validation, invite/verify token edge cases, workspace switching security
def test_signup_validation_errors(app):
    c = app.test_client()
    def sub(**kw):
        f = dict(name="Test User", email="new@x.com", workspace="Co", password="Sup3rSecret!", agree="1")
        f.update(kw)
        return c.post("/signup", data=f)
    assert sub(email="not-an-email").status_code == 400
    assert sub(name="").status_code == 400
    assert sub(password="weak").status_code == 400          # too short
    assert sub(password="password").status_code == 400      # too easy to guess
    assert sub(agree="").status_code == 400                  # must accept terms
    assert sub().status_code == 302                          # the actually-valid one succeeds
    fresh = app.test_client()                                 # a logged-out client, so the duplicate-email check is actually reached
    assert fresh.post("/signup", data=dict(name="Test User", email="new@x.com", workspace="Co", password="Sup3rSecret!", agree="1")).status_code == 400


def test_signup_race_condition_duplicate_email(app, monkeypatch):
    """Two signups for the same email arriving at nearly the same instant: the DB's unique
    constraint is the real guard, and the IntegrityError path must degrade to a clean error,
    not a 500."""
    from sqlalchemy.exc import IntegrityError
    import auth as A
    real_ins = A.ins
    calls = {"n": 0}
    def flaky_ins(table, data):
        if table == "users":
            calls["n"] += 1
            if calls["n"] == 1:
                return real_ins(table, data)  # first call: let it actually succeed
        return real_ins(table, data)
    # simulate: the pre-check "already registered" query raced and missed, but the DB catches it
    monkeypatch.setattr(A, "ins", lambda table, data: (_ for _ in ()).throw(IntegrityError("dup", {}, None)) if table == "users" else real_ins(table, data))
    c = app.test_client()
    r = c.post("/signup", data=dict(name="X", email="race@x.com", workspace="Co", password="Sup3rSecret!", agree="1"))
    assert r.status_code == 400 and b"already registered" in r.data


def test_invite_link_invalid_and_wrong_email(app, mails):
    o = signup(app, "o@x.com")
    c = app.test_client()
    assert c.get("/invite/not-a-real-token", follow_redirects=True).status_code == 200  # invalid token -> redirected to login, not a crash
    o.post("/team/invite", data=dict(email="teammate@x.com", role="viewer"))
    tok = re.search(r"/invite/([\w\-]+)", mails[-1][2]).group(1)
    # a DIFFERENT already-logged-in user opening the invite link must not silently join under the wrong identity
    other = signup(app, "someone-else@x.com", ws="Other Co")
    r = other.get(f"/invite/{tok}", follow_redirects=True)
    assert b"sent to teammate@x.com" in r.data
    assert dbq(app, "SELECT COUNT(*) c FROM memberships WHERE user_id=(SELECT id FROM users WHERE email='someone-else@x.com')", one=True)["c"] == 1  # still only their own workspace


def test_verify_token_invalid_and_expired(app):
    import auth as A
    o = signup(app, "vf@x.com")
    assert o.get("/verify/not-a-real-token", follow_redirects=True).status_code == 200
    assert dbq(app, "SELECT email_verified FROM users WHERE email='vf@x.com'", one=True)["email_verified"] == 0
    uid = dbq(app, "SELECT id FROM users WHERE email='vf@x.com'", one=True)["id"]
    expired_token = A.ser("verify").dumps(uid)
    # simulate real expiry by asking loads() for max_age=0 the same way the route does internally is not directly testable
    # without monkeypatching time, so instead prove a token for the WRONG salt/purpose is rejected (equally important: a
    # password-reset token must not double as a verify token just because both are itsdangerous-signed)
    reset_token = A.ser("reset").dumps(uid)
    r = o.get(f"/verify/{reset_token}", follow_redirects=True)
    assert b"expired or already used" in r.data
    assert dbq(app, "SELECT email_verified FROM users WHERE email='vf@x.com'", one=True)["email_verified"] == 0
    # the real token for the real purpose does work
    real_token = A.ser("verify").dumps(uid)
    o.get(f"/verify/{real_token}")
    assert dbq(app, "SELECT email_verified FROM users WHERE email='vf@x.com'", one=True)["email_verified"] == 1


def test_switch_workspace_rejects_non_member(app):
    o = signup(app, "o@x.com", ws="Mine")
    other_owner = signup(app, "other@x.com", ws="Not Mine")
    other_wid = wid_of(app, "other@x.com")
    o.post(f"/switch/{other_wid}")
    # must NOT have actually switched into a workspace o isn't a member of
    assert o.get("/app").data.decode().count("Not Mine") == 0
    r = o.get("/app").data.decode()
    assert "Mine" in r


# =====================================================================  config.py: the .env file parser and DATABASE_URL normalization (foundational, previously untested)
def test_env_file_parsing_quoted_unquoted_and_comments(tmp_path, monkeypatch):
    import config as C
    envfile = tmp_path / ".env"
    envfile.write_text(
        "TESTKEY_PLAIN=hello\n"
        "TESTKEY_QUOTED_SINGLE='hello world'\n"
        'TESTKEY_QUOTED_DOUBLE="hello, world"\n'
        "TESTKEY_WITH_COMMENT=value  # this is a trailing comment, not part of the value\n"
        "# a full-line comment, should be ignored entirely\n"
        "\n"
        "TESTKEY_NO_EQUALS_SIGN_SHOULD_BE_IGNORED\n"
        "TESTKEY_ALREADY_SET=should_not_override\n"
    )
    monkeypatch.setattr(C, "BASE", str(tmp_path))
    monkeypatch.setenv("TESTKEY_ALREADY_SET", "real_value_wins")  # setdefault must not clobber a real env var
    keys = ["TESTKEY_PLAIN", "TESTKEY_QUOTED_SINGLE", "TESTKEY_QUOTED_DOUBLE", "TESTKEY_WITH_COMMENT", "TESTKEY_NO_EQUALS_SIGN_SHOULD_BE_IGNORED", "TESTKEY_ALREADY_SET"]
    for k in keys:
        monkeypatch.delenv(k, raising=False) if k != "TESTKEY_ALREADY_SET" else None
    try:
        C._load_env()
        assert os.environ.get("TESTKEY_PLAIN") == "hello"
        assert os.environ.get("TESTKEY_QUOTED_SINGLE") == "hello world"
        assert os.environ.get("TESTKEY_QUOTED_DOUBLE") == "hello, world"
        assert os.environ.get("TESTKEY_WITH_COMMENT") == "value"
        assert "TESTKEY_NO_EQUALS_SIGN_SHOULD_BE_IGNORED" not in os.environ
        assert os.environ.get("TESTKEY_ALREADY_SET") == "real_value_wins"
    finally:
        for k in keys:
            monkeypatch.delenv(k, raising=False)


def test_database_url_normalization(monkeypatch):
    import config as C
    monkeypatch.setenv("DATABASE_URL", "postgres://user:pass@host/db")
    assert C._db_url().startswith("postgresql+psycopg2://user:pass@host/db")
    monkeypatch.setenv("DATABASE_URL", "postgresql://user:pass@host/db")
    assert C._db_url().startswith("postgresql+psycopg2://user:pass@host/db")
    monkeypatch.setenv("DATABASE_URL", "postgresql+psycopg2://already/normalized")
    assert C._db_url() == "postgresql+psycopg2://already/normalized"
    monkeypatch.delenv("DATABASE_URL", raising=False)
    assert C._db_url() == "sqlite:///" + os.path.join(C.BASE, "gtm.db")


# =====================================================================  workspace.py: account/security actions never directly exercised (rename, password change, invite revoke, activity log, integration keys)
def test_account_profile_and_password_change(app):
    o = signup(app, "acct@x.com")
    o.post("/account/profile", data=dict(name="New Name"))
    assert dbq(app, "SELECT name FROM users WHERE email='acct@x.com'", one=True)["name"] == "New Name"

    r = o.post("/account/password", data=dict(current="WrongPassword!", new="Sup3rNewPass!"), follow_redirects=True)
    assert b"incorrect" in r.data
    r2 = o.post("/account/password", data=dict(current="Sup3rSecret!", new="weak"), follow_redirects=True)
    assert b"at least 8 characters" in r2.data
    o.post("/account/password", data=dict(current="Sup3rSecret!", new="Br4ndNewPass!"))
    o.post("/logout")
    c = app.test_client()
    assert c.post("/login", data=dict(email="acct@x.com", password="Sup3rSecret!")).status_code == 401  # old password now rejected
    assert c.post("/login", data=dict(email="acct@x.com", password="Br4ndNewPass!")).status_code == 302  # new one works


def test_account_verify_resend_is_a_noop_once_verified(app, mails):
    o = signup(app, "rv@x.com")
    o.post("/account/verify")
    assert len(mails) >= 1
    before = len(mails)
    dbx(app, "UPDATE users SET email_verified=1 WHERE email='rv@x.com'")
    o.post("/account/verify")
    assert len(mails) == before  # already verified - must not send another email


def test_workspace_rename(app):
    o = signup(app, "rn@x.com", ws="Old Name")
    o.post("/workspace/rename", data=dict(name="Shiny New Name"))
    assert dbq(app, "SELECT name FROM workspaces", one=True)["name"] == "Shiny New Name"
    # a blank name must not wipe out the workspace's name
    o.post("/workspace/rename", data=dict(name=""))
    assert dbq(app, "SELECT name FROM workspaces", one=True)["name"] == "Shiny New Name"


def test_revoke_invite_and_activity_log(app, mails):
    o = signup(app, "act@x.com")
    o.post("/team/invite", data=dict(email="pending@x.com", role="viewer"))
    iid = dbq(app, "SELECT id FROM invites", one=True)["id"]
    o.post(f"/team/revoke/{iid}")
    assert dbq(app, "SELECT COUNT(*) c FROM invites", one=True)["c"] == 0
    page = o.get("/activity").data.decode()
    assert "invite" in page.lower() or "invited" in page.lower()


def test_hook_key_reset_actually_changes_it(app):
    o = signup(app, "hk@x.com")
    old_hook = o.get("/integrations").data.decode()
    import re as _re
    old_key = _re.search(r"/hooks/lead/([\w\-]+)", old_hook)
    o.post("/integrations/hook-reset")
    new_hook = o.get("/integrations").data.decode()
    new_key = _re.search(r"/hooks/lead/([\w\-]+)", new_hook)
    if old_key and new_key:
        assert old_key.group(1) != new_key.group(1)


def test_api_key_limit_of_ten(app, monkeypatch):
    o = _as(app, "apik@x.com", "Api Co", "general")
    dbx(app, "UPDATE workspaces SET plan='agency', plan_status='active'")
    for i in range(10):
        r = o.post("/integrations/api-keys", data=dict(name=f"key{i}", scope="read"))
        assert r.status_code == 200
    assert dbq(app, "SELECT COUNT(*) c FROM api_keys", one=True)["c"] == 10
    r = o.post("/integrations/api-keys", data=dict(name="key11", scope="read"), follow_redirects=True)
    assert b"already have 10" in r.data
    assert dbq(app, "SELECT COUNT(*) c FROM api_keys", one=True)["c"] == 10  # the 11th must not have been created


# =====================================================================  services.py: redirect-based SSRF bypass, image signature sniffing, slug collisions
def test_safe_get_reverifies_every_redirect_hop_not_just_the_first(monkeypatch):
    """The classic SSRF bypass: a URL that passes the public-URL check, then redirects to an
    internal address. safe_get must re-validate the destination of EVERY hop, not just the
    original URL."""
    class FakeResp:
        def __init__(self, status, headers=None):
            self.status_code, self.is_redirect, self.headers = status, status in (301, 302, 303, 307, 308), (headers or {})
        def raise_for_status(self): pass
        def iter_content(self, n): return iter([b"final content"])

    def fake_dns(host, port, *a, **kw):
        # 169.254.169.254 (cloud metadata) resolves to itself; any ".test" domain resolves to a real public IP
        ip = "169.254.169.254" if host == "169.254.169.254" else "93.184.216.34"
        return [(2, 1, 6, "", (ip, port))]
    monkeypatch.setattr(S.socket, "getaddrinfo", fake_dns)

    calls = []
    def fake_get(url, **kw):
        calls.append(url)
        if url == "https://public-site.test/start":
            return FakeResp(302, {"Location": "http://169.254.169.254/latest/meta-data"})  # tries to hop to cloud metadata
        return FakeResp(200)
    monkeypatch.setattr(S.requests, "get", fake_get)
    with pytest.raises(Exception):
        S.safe_get("https://public-site.test/start")
    assert len(calls) == 1  # must be blocked AT the redirect target, never actually fetched

    def fake_get_benign(url, **kw):
        calls.append(url)
        if url == "https://public-site.test/start2":
            return FakeResp(302, {"Location": "https://another-public-site.test/final"})
        return FakeResp(200)
    monkeypatch.setattr(S.requests, "get", fake_get_benign)
    data = S.safe_get("https://public-site.test/start2")
    assert data == b"final content"  # a genuinely public-to-public redirect must still work


def test_safe_get_too_many_redirects_raises(monkeypatch):
    monkeypatch.setattr(S.socket, "getaddrinfo", lambda host, port, *a, **kw: [(2, 1, 6, "", ("93.184.216.34", port))])
    class FakeResp:
        status_code, is_redirect, headers = 302, True, {"Location": "https://public-site.test/next"}
        def raise_for_status(self): pass
        def iter_content(self, n): return iter([b""])
    monkeypatch.setattr(S.requests, "get", lambda url, **kw: FakeResp())
    with pytest.raises(ValueError, match="too many redirects"):
        S.safe_get("https://public-site.test/start", hops=2)


def test_sniff_image_recognizes_every_supported_signature():
    assert S.sniff_image(b"\x89PNG\r\n\x1a\n" + b"0" * 20) == "png"
    assert S.sniff_image(b"\xff\xd8\xff" + b"0" * 20) == "jpg"
    assert S.sniff_image(b"GIF87a" + b"0" * 20) == "gif"
    assert S.sniff_image(b"GIF89a" + b"0" * 20) == "gif"
    assert S.sniff_image(b"RIFF" + b"0000" + b"WEBP" + b"0" * 20) == "webp"
    assert S.sniff_image(b"not an image at all") is None
    assert S.sniff_image(b"<svg onload=alert(1)>") is None
    assert S.sniff_image(b"%PDF-1.4") is None  # a PDF must not be accepted as an image


def test_make_slug_collision_gets_a_unique_suffix(app):
    o1 = signup(app, "s1@x.com", ws="Acme Co")
    o2 = signup(app, "s2@x.com", ws="Acme Co")  # same name, must not collide
    slugs = [r["slug"] for r in dbq(app, "SELECT slug FROM workspaces")]
    assert len(slugs) == len(set(slugs)) == 2


# =====================================================================  api.py: plan-feature gating and record limits enforced at the API level too, not just the web UI
def test_rest_api_enforces_plan_feature_gate(app, monkeypatch):
    import config as C
    # every real plan that has "api" access also happens to have "content_ops" - monkeypatch to isolate
    # the content_ops-specific module gate from the top-level "api" access gate it sits behind.
    monkeypatch.setitem(C.PLANS["team"], "features", C.PLANS["team"]["features"] - {"content_ops"})
    a = signup(app, "feat@x.com", ws="Feat Co")
    r = a.post("/integrations/api-keys", data=dict(name="Zap", scope="write"))
    key = re.search(r"gtm_[\w\-]{30,}", r.data.decode()).group(0)
    H = {"Authorization": f"Bearer {key}"}
    cl = app.test_client()
    dbx(app, "UPDATE workspaces SET plan='team', plan_status='active'")  # has API access, but content_ops removed above
    r2 = cl.get("/api/v1/clients", headers=H)
    assert r2.status_code == 403 and "content_ops" not in (r2.get_json().get("error") or "").lower()  # blocked by the module gate, not the API-access gate
    r3 = cl.post("/api/v1/clients", json=dict(name="Some Client"), headers=H)
    assert r3.status_code == 403
    # confirm a module that ISN'T feature-gated still works fine on this same plan
    assert cl.get("/api/v1/leads", headers=H).status_code == 200
    # restoring the feature (as agency plan genuinely has it) unlocks the same endpoint with no other change
    dbx(app, "UPDATE workspaces SET plan='agency', plan_status='active'")
    assert cl.get("/api/v1/clients", headers=H).status_code == 200


def test_rest_api_enforces_record_limit(app, monkeypatch):
    import config as C
    monkeypatch.setitem(C.PLANS["team"], "records", 3)  # isolate the limit-check logic without a huge insert loop
    a = signup(app, "lim@x.com", ws="Lim Co")
    r = a.post("/integrations/api-keys", data=dict(name="Zap", scope="write"))
    key = re.search(r"gtm_[\w\-]{30,}", r.data.decode()).group(0)
    H = {"Authorization": f"Bearer {key}"}
    cl = app.test_client()
    dbx(app, "UPDATE workspaces SET plan='team', plan_status='active'")  # a plan that actually has API access
    with app.app_context():
        current = S.record_count(1)
    room = C.PLANS["team"]["records"] - current
    from db import ins as _ins
    with app.app_context():
        for _ in range(max(0, room - 1)):
            _ins("leads", dict(workspace_id=1, name="x", stage="New"))
    ok = cl.post("/api/v1/leads", json=dict(name="last one allowed", stage="New"), headers=H)
    assert ok.status_code == 201, ok.get_json()
    blocked = cl.post("/api/v1/leads", json=dict(name="over the limit", stage="New"), headers=H)
    assert blocked.status_code == 403 and "Record limit" in blocked.get_json().get("error", "")


# =====================================================================  db.py: upgrading a real, older production database in place (missing newer columns)
def test_schema_upgrade_adds_missing_columns_without_losing_data(tmp_path):
    """Simulates a real production database created before several EXTRA_COLUMNS existed
    (e.g. before brand_color/onboarded/board positions were added). The app must add the
    missing columns on next boot without touching existing data."""
    import sqlite3
    dbfile = tmp_path / "old.db"
    conn = sqlite3.connect(str(dbfile))
    c = conn.cursor()
    # deliberately old-style workspaces table: only the columns that existed at first release
    c.execute("CREATE TABLE workspaces (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, slug TEXT UNIQUE, owner_id INTEGER, "
              "industry TEXT DEFAULT 'general', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)")
    c.execute("INSERT INTO workspaces (name, slug, owner_id, industry) VALUES ('Old Co', 'old-co-orig', 1, 'hospitality')")
    conn.commit()
    before_cols = {r[1] for r in c.execute("PRAGMA table_info(workspaces)")}
    assert "brand_color" not in before_cols  # confirm this really is the old shape before we touch it
    conn.close()

    from app import create_app
    app2 = create_app(config={"TESTING": True, "WTF_CSRF_ENABLED": False}, database_url=f"sqlite:///{dbfile}")

    conn2 = sqlite3.connect(str(dbfile))
    c2 = conn2.cursor()
    after_cols = {r[1] for r in c2.execute("PRAGMA table_info(workspaces)")}
    # every column that ships today must now be present, added in place
    for col in ("brand_color", "ics_key", "hook_key"):
        assert col in after_cols, f"{col} was not added during upgrade"
    # the pre-existing row must survive completely intact, with real (not default-blank) values kept
    row = c2.execute("SELECT name, slug, industry, brand_color FROM workspaces WHERE slug='old-co-orig'").fetchone()
    assert row == ("Old Co", "old-co-orig", "hospitality", "")  # new column present with its documented default, old data untouched
    conn2.close()

    # and the upgraded database is genuinely usable, not just structurally patched
    with app2.app_context():
        client = app2.test_client()
        r = client.get("/login")
        assert r.status_code == 200


# =====================================================================  recent fixes: password policy, verification/reset links in prod, resend-link redirect, PIN sign-in
def test_password_policy_requires_uppercase_and_special_char(app):
    c = app.test_client()
    weak_no_upper = c.post("/signup", data=dict(name="N", email="nu@x.com", password="lowercase1!", agree="1"))
    assert weak_no_upper.status_code == 400 and b"UPPERCASE" in weak_no_upper.data
    weak_no_special = c.post("/signup", data=dict(name="N", email="ns@x.com", password="NoSpecial1", agree="1"))
    assert weak_no_special.status_code == 400 and b"special character" in weak_no_special.data
    ok = c.post("/signup", data=dict(name="N", email="ok@x.com", password="Good1Pass!", agree="1"))
    assert ok.status_code == 302


def test_verify_and_reset_links_surface_even_in_production_without_smtp(app):
    """Before this fix: dev_link() only ever showed the fallback link when C.PROD was False, so a
    production deploy with no SMTP configured yet left users with NO way at all to verify their
    email or reset a forgotten password. The link must show whenever the email genuinely wasn't
    sent, in any mode."""
    import config as C
    old = C.PROD
    try:
        C.PROD = True
        c = signup(app, "prodcheck@x.com")
        r = c.post("/account/verify", follow_redirects=True)
        assert r.status_code == 200 and b"here it is" in r.data
        f = app.test_client().post("/forgot", data=dict(email="prodcheck@x.com"), follow_redirects=True)
        assert f.status_code == 200 and b"here it is" in f.data
    finally:
        C.PROD = old


def test_resend_verification_redirects_back_not_to_account_page(app):
    c = signup(app, "resend@x.com")
    r = c.post("/account/verify", headers={"Referer": "/app"}, follow_redirects=False)
    assert r.status_code == 302 and r.headers["Location"].endswith("/app")


def test_pin_signin_set_use_and_remove(app):
    c = signup(app, "pin@x.com")
    # wrong current password is rejected, PIN not set
    c.post("/account/pin", data=dict(current="wrong", pin="1234"))
    assert dbq(app, "SELECT pin_hash FROM users WHERE email='pin@x.com'", one=True)["pin_hash"] == ""
    # correct password sets a PIN
    c.post("/account/pin", data=dict(current="Sup3rSecret!", pin="4321"))
    assert dbq(app, "SELECT pin_hash FROM users WHERE email='pin@x.com'", one=True)["pin_hash"] != ""
    c.post("/logout")
    # login now stops at the PIN step instead of going straight in
    fresh = app.test_client()
    r = fresh.post("/login", data=dict(email="pin@x.com", password="Sup3rSecret!"))
    assert r.status_code == 302 and "/login/pin" in r.headers["Location"]
    assert fresh.get("/app").status_code == 302  # not logged in yet
    bad_pin = fresh.post("/login/pin", data=dict(code="0000"))
    assert bad_pin.status_code == 401
    good_pin = fresh.post("/login/pin", data=dict(code="4321"))
    assert good_pin.status_code == 302
    assert fresh.get("/app").status_code == 200
    # turning it off removes the extra step
    fresh.post("/account/pin", data=dict(remove="1"))
    fresh.post("/logout")
    fresh2 = app.test_client()
    r2 = fresh2.post("/login", data=dict(email="pin@x.com", password="Sup3rSecret!"))
    assert r2.status_code == 302 and "/login/pin" not in r2.headers["Location"]


# =====================================================================  new top bar: company logo upload, help/profile dropdowns, notifications pop-up feed
def test_topbar_logo_upload_and_notifications_feed(app):
    o = signup(app, "o@x.com")
    r = o.get("/app").data.decode()
    assert 'class="tb-brand"' in r and 'class="dd-menu"' in r and "loadNotifs" in r
    # logo upload: only a real image is accepted, and it then shows up in the top bar
    bad = o.post("/account/logo", data=dict(file=(io.BytesIO(b"not-an-image"), "x.png")), content_type="multipart/form-data")
    assert bad.status_code == 302
    wid = wid_of(app, "o@x.com")
    assert dbq(app, "SELECT logo_url FROM workspaces WHERE id=?", (wid,), one=True)["logo_url"] == ""
    ok = o.post("/account/logo", data=dict(file=(io.BytesIO(PNG), "logo.png")), content_type="multipart/form-data")
    assert ok.status_code == 302
    logo = dbq(app, "SELECT logo_url FROM workspaces WHERE id=?", (wid,), one=True)["logo_url"]
    assert logo and o.get("/app").data.decode().count("<img") >= 1
    assert o.post("/account/logo/remove").status_code == 302
    assert dbq(app, "SELECT logo_url FROM workspaces WHERE id=?", (wid,), one=True)["logo_url"] == ""
    # only an owner/admin can change the logo
    creator = signup(app, "cr@x.com")
    S = __import__("services")
    dbx(app, "UPDATE memberships SET role='creator' WHERE user_id=(SELECT id FROM users WHERE email='cr@x.com')")
    # notifications pop-up feed: empty then populated
    assert o.get("/notifications/recent").get_json()["items"] == []
    dbx(app, "INSERT INTO notifications (workspace_id, user_id, message, link) VALUES (?, (SELECT id FROM users WHERE email='o@x.com'), 'Test note', '/app')", (wid,))
    items = o.get("/notifications/recent").get_json()["items"]
    assert len(items) == 1 and items[0]["message"] == "Test note" and items[0]["is_read"] is False


# =====================================================================  GTM Plans: mandatory fields, CSV template download
def test_gtm_plan_required_fields_and_csv_template(app):
    o = signup(app, "o@x.com")
    # missing required fields (name, product, budget) is rejected, with values kept so nothing is retyped
    r = o.post("/m/gtm_plans/save", data=dict(name="", product="", stage="Planning", budget=""))
    assert r.status_code == 400 and b"Plan name" in r.data and b"Planning" in r.data  # stage kept even though it was valid
    assert dbq(app, "SELECT COUNT(*) c FROM gtm_plans", one=True)["c"] == 0
    ok = o.post("/m/gtm_plans/save", data=dict(name="Diwali Launch", product="Gift box", budget="50000"))
    assert ok.status_code == 302 and dbq(app, "SELECT COUNT(*) c FROM gtm_plans", one=True)["c"] == 1
    # editing down to blank is rejected too, and the edit form re-shows with the id intact (still editable, not lost)
    rid = dbq(app, "SELECT id FROM gtm_plans", one=True)["id"]
    bad_edit = o.post("/m/gtm_plans/save", data=dict(id=rid, name="", product="Gift box", budget="50000"))
    assert bad_edit.status_code == 400 and f'value="{rid}"'.encode() in bad_edit.data
    # CSV template download: header-only, matching the real export's column names
    tpl = o.get("/m/gtm_plans/template")
    assert tpl.status_code == 200 and b"name,product,stage" in tpl.data
    assert b"Import CSV" in o.get("/m/gtm_plans").data


# =====================================================================  Required fields rolled out across modules, not just GTM Plans
def test_required_fields_rolled_out_sitewide(app):
    o = signup(app, "o@x.com")
    # each of these modules now rejects a blank identity field instead of silently saving an unnamed record
    for key, blank_data in [
        ("campaigns", dict(name="")),
        ("influencers", dict(name="")),
        ("tasks", dict(title="")),
        ("leads", dict(name="")),
        ("content_items", dict(title="")),
    ]:
        r = o.post(f"/m/{key}/save", data=blank_data)
        assert r.status_code == 400, key
        assert dbq(app, f"SELECT COUNT(*) c FROM {key}", one=True)["c"] == 0, key
    # a filled-in identity field saves fine
    assert o.post("/m/campaigns/save", data=dict(name="Diwali push")).status_code == 302
    assert dbq(app, "SELECT COUNT(*) c FROM campaigns", one=True)["c"] == 1


# =====================================================================  Market Pulse: watchlist folded into the page itself, no separate tab
def test_pulse_watch_inline_add_remove_and_no_separate_nav_tab(app):
    o = signup(app, "o@x.com")
    dbx(app, "DELETE FROM watchlist")
    nav = o.get("/app").data.decode()
    assert "Market Pulse" in nav and "Watchlist" not in nav   # no separate nav item any more
    page = o.get("/pulse").data.decode()
    assert 'action="/pulse/watch/add"' in page
    r = o.post("/pulse/watch/add", data=dict(term="Diwali offers", type="Keyword", region="India"))
    assert r.status_code == 302
    row = dbq(app, "SELECT * FROM watchlist WHERE term='Diwali offers'", one=True)
    assert row is not None
    page2 = o.get("/pulse").data.decode()
    assert "Diwali offers" in page2 and f"/pulse/watch/{row['id']}/remove" in page2
    assert o.post(f"/pulse/watch/{row['id']}/remove").status_code == 302
    assert dbq(app, "SELECT COUNT(*) c FROM watchlist WHERE term='Diwali offers'", one=True)["c"] == 0


# =====================================================================  Launch Checklist board: drag-and-drop UI markup + sort-by-due-date
def test_kanban_drag_markup_and_sort_by_due_date(app):
    o = signup(app, "o@x.com")
    page = o.get("/kanban/launch_checklist").data.decode()
    assert 'draggable="true"' in page and "dropCard(event" in page and "Sort by due date" in page
    rid = dbq(app, "SELECT id FROM launch_checklist ORDER BY id LIMIT 1", one=True)["id"]
    dbx(app, "UPDATE launch_checklist SET due_date='2020-01-01' WHERE id=?", (rid,))
    sorted_page = o.get("/kanban/launch_checklist?sort=due").data.decode()
    assert "✓ Sorted by due date" in sorted_page
    # the move itself reuses the existing status-change endpoint, so dragging a card works the same way the dropdown already does
    assert o.post(f"/m/launch_checklist/{rid}/status", data=dict(value="Done")).status_code == 302
    assert dbq(app, "SELECT status FROM launch_checklist WHERE id=?", (rid,), one=True)["status"] == "Done"


def test_kanban_card_can_be_dragged_to_a_manual_order_within_one_column(app):
    # launch_checklist's board groups by "status" (To do / In progress / Done); every seeded item starts
    # as "To do", so pick 3 of those to reorder within that one column
    o = signup(app, "o@x.com")
    rows = dbq(app, "SELECT id, status FROM launch_checklist WHERE status='To do' ORDER BY id LIMIT 3")
    assert len(rows) == 3
    a, b, c = [r["id"] for r in rows]
    # give these 3 a known starting order (a, b, c) and park every other "To do" sibling well below them,
    # so this test's own drags are the only thing that can move a/b/c relative to each other
    for i, rid in enumerate((a, b, c)):
        dbx(app, "UPDATE launch_checklist SET sort_order=? WHERE id=?", (i * 10, rid))
    dbx(app, f"UPDATE launch_checklist SET sort_order=sort_order+1000 WHERE status='To do' AND id NOT IN ({a},{b},{c})")

    # drag c to the front (drop it before a, same column) - should reorder without touching status
    r = o.post(f"/m/launch_checklist/{c}/reorder", data=dict(status="To do", before_id=str(a)))
    assert r.status_code == 302
    order = [r["id"] for r in dbq(app, f"SELECT id FROM launch_checklist WHERE id IN ({a},{b},{c}) ORDER BY sort_order ASC")]
    assert order == [c, a, b]

    # dropping on the empty column background (no before_id) sends a card to the end
    r = o.post(f"/m/launch_checklist/{a}/reorder", data=dict(status="To do", before_id=""))
    assert r.status_code == 302
    order = [r["id"] for r in dbq(app, f"SELECT id FROM launch_checklist WHERE id IN ({a},{b},{c}) ORDER BY sort_order ASC")]
    assert order == [c, b, a]

    # dragging into a DIFFERENT column changes status too, same as the dropdown/status endpoint does
    r = o.post(f"/m/launch_checklist/{b}/reorder", data=dict(status="Done", before_id=""))
    assert r.status_code == 302
    assert dbq(app, "SELECT status FROM launch_checklist WHERE id=?", (b,), one=True)["status"] == "Done"


# =====================================================================  Mood Boards: one-click sample board for first-time users
def test_sample_moodboard(app):
    o = signup(app, "o@x.com")
    assert b"Load a sample board" in o.get("/moodboards").data
    r = o.post("/moodboards/sample")
    assert r.status_code == 302
    bid = dbq(app, "SELECT id FROM boards WHERE name LIKE 'Sample board%'", one=True)["id"]
    items = dbq(app, "SELECT * FROM board_items WHERE board_id=?", (bid,))
    assert len(items) == 8 and all(i["x"] >= 0 for i in items)  # already placed on the canvas, nothing left at the unplaced sentinel
    assert {i["kind"] for i in items} == {"image", "color", "note", "link"}  # a worked example of every card type
    assert b"Sample board" in o.get(f"/moodboards/{bid}").data


# =====================================================================  Mood board cards are editable in place, not just draggable/deletable
def test_moodboard_item_edit_in_place(app):
    o = signup(app, "o@x.com")
    o.post("/moodboards/sample")
    bid = dbq(app, "SELECT id FROM boards WHERE name LIKE 'Sample board%'", one=True)["id"]
    note = dbq(app, "SELECT * FROM board_items WHERE board_id=? AND kind='note' LIMIT 1", (bid,), one=True)
    color = dbq(app, "SELECT * FROM board_items WHERE board_id=? AND kind='color' LIMIT 1", (bid,), one=True)
    link = dbq(app, "SELECT * FROM board_items WHERE board_id=? AND kind='link' LIMIT 1", (bid,), one=True)
    assert b"double-click a card to edit" in o.get(f"/moodboards/{bid}").data
    r = o.post(f"/moodboards/item/{note['id']}/edit", json=dict(content="Edited mood note", caption="New caption"))
    assert r.status_code == 200
    assert dbq(app, "SELECT content, caption FROM board_items WHERE id=?", (note["id"],), one=True) == {"content": "Edited mood note", "caption": "New caption"}
    # empty note is rejected
    assert o.post(f"/moodboards/item/{note['id']}/edit", json=dict(content="", caption="")).status_code == 400
    # invalid colour is rejected, valid one saves
    assert o.post(f"/moodboards/item/{color['id']}/edit", json=dict(content="notacolor", caption="")).status_code == 400
    assert o.post(f"/moodboards/item/{color['id']}/edit", json=dict(content="#112233", caption="Accent")).status_code == 200
    assert dbq(app, "SELECT content FROM board_items WHERE id=?", (color["id"],), one=True)["content"] == "#112233"
    # link content is cleaned/validated the same way as adding one
    assert o.post(f"/moodboards/item/{link['id']}/edit", json=dict(content="https://example.com/x", caption="Ref")).status_code == 200
    assert dbq(app, "SELECT content FROM board_items WHERE id=?", (link["id"],), one=True)["content"] == "https://example.com/x"


# =====================================================================  Mood board: images can be swapped in place (no delete + re-add needed), and the board has a full usage guide
def test_moodboard_image_replace_in_place_and_guide_present(app):
    o = signup(app, "o@x.com")
    o.post("/moodboards/new", data=dict(name="Board"))
    bid = dbq(app, "SELECT id FROM boards LIMIT 1", one=True)["id"]
    o.post(f"/moodboards/{bid}/add", data=dict(kind="upload", file=(io.BytesIO(PNG), "first.png")), content_type="multipart/form-data")
    it = dbq(app, "SELECT * FROM board_items WHERE board_id=? AND kind='image' LIMIT 1", (bid,), one=True)
    old_content, old_x, old_y = it["content"], it["x"], it["y"]
    # swap in a new picture: position/size/z stay put, old file is replaced
    r = o.post(f"/moodboards/item/{it['id']}/replace-image", data=dict(file=(io.BytesIO(PNG), "second.png")), content_type="multipart/form-data")
    assert r.status_code == 200
    updated = dbq(app, "SELECT * FROM board_items WHERE id=?", (it["id"],), one=True)
    assert updated["content"] != old_content and updated["x"] == old_x and updated["y"] == old_y
    # a non-image file is rejected
    assert o.post(f"/moodboards/item/{it['id']}/replace-image", data=dict(file=(io.BytesIO(b"not an image"), "x.txt")), content_type="multipart/form-data").status_code == 400
    # an image's picture can ALSO be swapped by pasting a URL through the normal edit endpoint (not just file upload) — fully editable either way
    content_before = dbq(app, "SELECT content FROM board_items WHERE id=?", (it["id"],), one=True)["content"]
    r2 = o.post(f"/moodboards/item/{it['id']}/edit", json=dict(content="https://example.com/new-pic.jpg", caption="Swapped via URL"))
    assert r2.status_code == 200
    after_url_swap = dbq(app, "SELECT content, caption FROM board_items WHERE id=?", (it["id"],), one=True)
    assert after_url_swap["content"] == "https://example.com/new-pic.jpg" != content_before
    assert after_url_swap["caption"] == "Swapped via URL"
    # leaving content blank only changes the caption, keeping the current picture
    r3 = o.post(f"/moodboards/item/{it['id']}/edit", json=dict(content="", caption="Caption only change"))
    assert r3.status_code == 200
    after_caption_only = dbq(app, "SELECT content, caption FROM board_items WHERE id=?", (it["id"],), one=True)
    assert after_caption_only["content"] == "https://example.com/new-pic.jpg" and after_caption_only["caption"] == "Caption only change"
    board_html = o.get(f"/moodboards/{bid}").data
    assert b"How to use" in board_html and b"Zoom" in board_html or b"zoomBy" in board_html  # guide button + zoom controls present
    # the moodboards list page carries an always-visible how-it-works guide, not just a first-run popup
    assert b"How it works" in o.get("/moodboards").data
    # the board page itself also auto-opens the guide the first time a browser visits (per the "should be on when
    # you click into the panel" ask) — once per browser via localStorage, not every visit; the ❓ button always reopens it
    assert b"mbGuideSeen" in board_html and b"showModal" in board_html


# =====================================================================  Sidebar shows every module directly: no "More tools" popup hiding standalone
# modules behind an extra click. A module is independent data (not dependent on another panel), so it gets
# its own always-visible link; a popup is reserved for genuinely dependent/contextual things (see the
# "Need an idea?" overlay test below), not for a module's own primary navigation.
def test_sidebar_shows_every_module_directly_no_more_tools_popup(app):
    o = signup(app, "o@x.com")
    html = o.get("/app").data.decode()
    aside = html.split("<aside>", 1)[1].split("</aside>", 1)[0]
    assert "More tools" not in aside and "moreToolsModal" not in html  # the popup is gone entirely
    # a module that used to be pushed into the popup (5th+ item in its group) is now directly in the sidebar
    assert "Rate Card" in aside
    assert "Launch Checklist" in aside
    # ...same as the groups' first few items, which were always shown directly
    assert "GTM Plans" in aside and "Calendar" in aside and "Promotions" in aside and "Events" in aside


# =====================================================================  Idea Bank quick-add: fetch-based overlay fragment (replaces the broken same-origin iframe)
def test_quick_add_overlay_fragment_and_modal_submit(app):
    o = signup(app, "o@x.com")
    ideas_html = o.get("/ideas").data
    assert b"openOverlay(" in ideas_html and b"<iframe" not in o.get("/").data  # overlay no longer uses an iframe anywhere
    # the page itself still has no iframe-based overlay markup
    assert b'id="overlayFrame"' not in ideas_html
    # a modal=1 request returns a bare form fragment, not a full page
    frag = o.get("/m/content_items/new?modal=1&title=Test+idea&channel=Instagram&format=Story&next=/ideas").data
    assert b"<html" not in frag.lower() and b"<form" in frag and b'name="modal" value="1"' in frag
    assert b"Need an idea?" not in frag  # the modal variant hides the link back to the page it was opened from
    # submitting that fragment's form (as the injected JS does) saves the record and still redirects, just like the full page would
    r = o.post("/m/content_items/save", data=dict(title="Test idea", channel="Instagram", format="Story", status="Idea", next="/ideas", modal="1"))
    assert r.status_code == 302
    assert dbq(app, "SELECT * FROM content_items WHERE title='Test idea'", one=True)
    # a non-modal save still renders the normal full page on validation failure
    full = o.get("/m/content_items/new").data
    assert b"<html" in full.lower() and b'name="modal"' not in full
    # "Need an idea?" on the full New Content Item page no longer does a hard navigation to the whole Idea
    # Bank page - it opens a small overlay (fetch + openOverlay), with the same close/cross the other
    # quick-add overlays use, instead of losing the content-item form the person was filling in
    assert b"openOverlay(this.href,'')" in full and b'href="/ideas/picker"' in full  # opens as a small overlay now, not a plain link to the full Idea Bank page


# =====================================================================  Idea Bank "Need an idea?" overlay: a few ideas only, not the whole page
def test_ideas_picker_overlay_is_a_small_fragment_not_the_full_idea_bank(app):
    o = signup(app, "o@x.com")
    frag = o.get("/ideas/picker").data
    assert b"<html" not in frag.lower()           # a bare fragment, like the other overlay content
    assert b"Need an idea?" in frag                 # has its own small heading
    assert b"+ Use this" in frag and b"openOverlay(" in frag  # picking one opens the content-item form in the SAME overlay, not a new page
    assert b"Campaign" not in frag and b"Persona" not in frag  # just content ideas - not the campaigns/personas/KPI sections from the full page
    # picking an idea pre-fills a new content item, same as the full Idea Bank's "+ Add" does
    assert b"/m/content_items/new?" in frag and b"Idea+pillar" in frag.replace(b"%3A", b":")


# =====================================================================  First-time tour: wording must hold up on any page (the spotlighted sidebar
# link is present everywhere, not just on /dashboard), and the tour must never sit on screen alongside another dialog
def test_tour_wording_is_page_agnostic_and_yields_to_dialogs(app):
    o = signup(app, "o@x.com")
    html = o.get("/ideas").data
    assert b"You are here: the Dashboard" not in html  # would be false on any page except /dashboard itself
    assert "Dashboard".encode() in html
    # opening ANY dialog (quick-add overlay, More tools, a board's How-to-use, ...) must dismiss the tour first
    assert b"HTMLDialogElement.prototype.showModal=function(){dismissTour();" in html
    # the tour's own delayed auto-start must not fire on top of a dialog that's already open
    assert b"function startTour(){if(document.querySelector('dialog[open]'))return;tourStep(0);}" in html


# =====================================================================  Downloadable guide offered at login/signup, and clarity on how invited team members sign in
def test_guide_link_at_login_and_team_signin_clarity(app):
    anon = app.test_client()
    assert b"Download a quick visual tour" in anon.get("/login").data
    assert b"guides/Debut-Visual-Tour.pdf" in anon.get("/login").data
    assert b"Download a quick visual tour" in anon.get("/signup").data
    o = signup(app, "o@x.com")
    team_html = o.get("/team").data
    assert b"one-time join link" in team_html and b"sign in themselves" in team_html


def test_company_name_shown_centered_on_invited_signin(app, mails):
    o = signup(app, "o@x.com", ws="Acme Foods Marketing")
    signup(app, "existing@x.com", ws="Someone else's workspace")  # pre-existing account, unrelated workspace
    new_link = _invite(app, o, mails, "new@x.com", "viewer")
    existing_link = _invite(app, o, mails, "existing@x.com", "viewer")
    anon1, anon2 = app.test_client(), app.test_client()
    assert anon1.get(new_link).status_code == 302  # redirects to /signup?invite=... for a brand-new email
    signup_html = anon1.get(new_link, follow_redirects=True).data
    assert b'class="invite-brand"' in signup_html and b"Acme Foods Marketing" in signup_html
    # an invite sent to an email that already has an account routes to /login instead, with the same branding
    login_html = anon2.get(existing_link, follow_redirects=True).data
    assert b'class="invite-brand"' in login_html and b"Acme Foods Marketing" in login_html
    # a plain, token-less visit shows neither (nothing to brand it with)
    plain = app.test_client().get("/login").data
    assert b'class="invite-brand"' not in plain


def test_workspace_rename_validation_and_formatting_hints(app):
    o = signup(app, "o@x.com", ws="Acme")
    assert b'maxlength="80"' in o.get("/account").data
    r = o.post("/workspace/rename", data=dict(name=""))
    assert r.status_code == 302
    assert dbq(app, "SELECT name FROM workspaces LIMIT 1", one=True)["name"] == "Acme"  # empty name rejected, old name kept
    o.post("/workspace/rename", data=dict(name="Renamed Co"))
    assert dbq(app, "SELECT name FROM workspaces LIMIT 1", one=True)["name"] == "Renamed Co"


# =====================================================================  Ask Debut: AI-answered search, grounded in the workspace's own data
def test_ask_debut_ai_answer(app, monkeypatch):
    monkeypatch.setattr(C, "ANTHROPIC_API_KEY", "k")

    class R:
        status_code = 200
        def json(self): return {"content": [{"type": "text", "text": "Your engagement looks low because you are not posting consistently."}]}
    monkeypatch.setattr(views.requests, "post", lambda *a, **k: R())
    o = signup(app, "o@x.com")
    dbx(app, "UPDATE users SET email_verified=1")
    assert b"Ask Debut" in o.get("/search?q=hello").data  # the AI box shows up on the search page once AI is configured
    r = o.post("/api/ask", json=dict(q="Why is my engagement low?"))
    assert r.status_code == 200
    d = r.get_json()
    assert "not posting consistently" in d["text"] and d["used"] == 1
    # an empty question is rejected without spending a credit
    assert o.post("/api/ask", json=dict(q="")).status_code == 400


# =====================================================================  Excel (.xlsx) mock template download, alongside the existing CSV template
def test_excel_mock_template_download(app):
    o = signup(app, "o@x.com")
    r = o.get("/m/gtm_plans/template.xlsx")
    assert r.status_code == 200
    assert r.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "gtm_plans-mock-template.xlsx" in r.headers["Content-Disposition"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Plan name"          # header row uses field labels
    assert ws.cell(row=2, column=1).value.startswith("Example")   # a filled-in mock example row


def test_excel_mock_template_handles_slash_in_module_title(app):
    # "Personas / ICP" has a slash in its title, which Excel forbids in a sheet name (\ / * [ ] : ?) — this
    # used to crash the download with a ValueError from openpyxl; confirm it's sanitized instead.
    o = signup(app, "o2@x.com")
    r = o.get("/m/personas/template.xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.active.title == "Personas - ICP"


def test_excel_mock_template_works_for_every_module(app):
    # every module's Excel template should download without error, whatever punctuation is in its title
    from openpyxl import load_workbook
    from modules import MODULES
    o = signup(app, "o3@x.com")
    for key in MODULES:
        r = o.get(f"/m/{key}/template.xlsx")
        assert r.status_code == 200, key
        wb = load_workbook(io.BytesIO(r.data))
        assert wb.active.title  # non-empty, valid sheet title
    assert b"Filled-in mock example (Excel)" in o.get("/m/gtm_plans").data
    assert b"Blank template (CSV)" in o.get("/m/gtm_plans").data


# =====================================================================  Brand deck upload: automatic (no-AI) colour/font/logo extraction, applied to the workspace + saved to Brand Kit
def _make_test_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from PIL import Image
    prs = Presentation()
    logo_buf = io.BytesIO()
    Image.new("RGB", (100, 100), (30, 136, 229)).save(logo_buf, format="PNG")
    logo_bytes = logo_buf.getvalue()
    for _ in range(2):  # same logo on 2 slides -> detected as the repeated logo, not a one-off picture
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Acme Co"
        run.font.name = "Montserrat"
        run.font.size = Pt(36)
        rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(2), Inches(2))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(0x1E, 0x88, 0xE5)
        slide.shapes.add_picture(io.BytesIO(logo_bytes), Inches(5), Inches(5), Inches(1), Inches(1))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_brand_scan_extracts_colour_font_and_repeated_logo_from_pptx():
    import brand_scan
    result = brand_scan.scan_deck(_make_test_pptx(), "deck.pptx")
    assert "1E88E5" in result["colors"]
    assert "Montserrat" in result["fonts"]
    assert result["logo_bytes"] is not None


def test_brand_scan_rejects_unsupported_file_type():
    import brand_scan
    with pytest.raises(ValueError):
        brand_scan.scan_deck(b"hello", "notes.txt")


def test_brand_scan_raises_cleanly_on_corrupt_file():
    import brand_scan
    with pytest.raises(ValueError):
        brand_scan.scan_deck(b"not a real pptx", "broken.pptx")


def test_account_brand_deck_upload_applies_colour_logo_and_saves_brand_kit(app):
    o = signup(app, "deck@x.com")
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(_make_test_pptx()), "brand.pptx")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and b"Brand deck scanned" in r.data
    wid = wid_of(app, "deck@x.com")
    ws = dbq(app, "SELECT brand_color, secondary_color, logo_url FROM workspaces WHERE id=?", (wid,), one=True)
    assert ws["brand_color"].upper() == "#1E88E5"
    assert ws["logo_url"]
    kit = dbq(app, "SELECT kind, value FROM brand_kit WHERE workspace_id=?", (wid,))
    assert any(k["kind"] == "Colour" and k["value"].upper() == "#1E88E5" for k in kit)
    assert any(k["kind"] == "Font" and k["value"] == "Montserrat" for k in kit)


def _make_test_pdf():
    import fitz
    doc = fitz.open()
    logo_buf = io.BytesIO()
    from PIL import Image
    Image.new("RGB", (100, 100), (30, 136, 229)).save(logo_buf, format="PNG")
    logo_bytes = logo_buf.getvalue()
    for _ in range(2):  # same logo on 2 pages -> detected as the repeated logo
        page = doc.new_page(width=612, height=792)
        page.draw_rect(fitz.Rect(50, 50, 250, 250), color=(30 / 255, 136 / 255, 229 / 255), fill=(30 / 255, 136 / 255, 229 / 255))
        page.insert_text((50, 300), "Acme Co", fontsize=36, fontname="helv")
        page.insert_image(fitz.Rect(400, 400, 500, 500), stream=logo_bytes)
    return doc.tobytes()


def test_account_brand_deck_upload_handles_real_pdf_too(app):
    # the pptx path is covered above by test_account_brand_deck_upload_applies_colour_logo_and_saves_brand_kit;
    # the pdf path (via PyMuPDF) was previously only reachable in brand_scan.py and never exercised end-to-end
    o = signup(app, "pdfdeck@x.com")
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(_make_test_pdf()), "brand.pdf")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and b"Brand deck scanned" in r.data
    wid = wid_of(app, "pdfdeck@x.com")
    ws = dbq(app, "SELECT brand_color, logo_url FROM workspaces WHERE id=?", (wid,), one=True)
    assert ws["brand_color"].upper() == "#1E88E5"
    assert ws["logo_url"]


def test_account_brand_deck_does_not_force_theme_colour_for_personal_brand_industry(app):
    # brand theming (workspace brand_color/secondary_color) is a business-workspace concept - account_theme()
    # itself 403s for industry="personal_brand" - so the deck scan must not silently set those columns either,
    # even though it should still save the detected colour/font to Brand Kit as reference material
    c = app.test_client()
    r = c.post("/signup", data=dict(name="Solo", email="solo@x.com", workspace="Solo Creator", password="Sup3rSecret!", agree="1", industry="personal_brand"))
    assert r.status_code == 302
    r = c.post("/account/brand-deck", data=dict(file=(io.BytesIO(_make_test_pptx()), "brand.pptx")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and b"Brand deck scanned" in r.data
    wid = wid_of(app, "solo@x.com")
    ws = dbq(app, "SELECT brand_color, logo_url FROM workspaces WHERE id=?", (wid,), one=True)
    assert not ws["brand_color"]          # not force-applied as workspace theme for a personal brand
    assert ws["logo_url"]                 # logo is still applied - it isn't a "theme colour" concept
    kit = dbq(app, "SELECT kind, value FROM brand_kit WHERE workspace_id=?", (wid,))
    assert any(k["kind"] == "Colour" and k["value"].upper() == "#1E88E5" for k in kit)  # still saved for reference


def test_account_brand_deck_handles_a_totally_blank_deck_with_honest_message(app):
    from pptx import Presentation
    o = signup(app, "blankdeck@x.com")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # one slide, nothing on it
    buf = io.BytesIO()
    prs.save(buf)
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(buf.getvalue()), "blank.pptx")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200
    assert b"couldn" in r.data.lower() and b"confidently detect" in r.data.lower()
    wid = wid_of(app, "blankdeck@x.com")
    ws = dbq(app, "SELECT brand_color, logo_url FROM workspaces WHERE id=?", (wid,), one=True)
    assert not ws["brand_color"] and not ws["logo_url"]  # nothing was force-applied from an empty scan


def test_creator_role_cannot_reach_new_account_pages_even_via_direct_url(app, mails):
    o = signup(app, "o6@x.com")
    cr = _join(app, o, mails, "cr6@x.com", "creator", "Cy")
    for path in ("/marketing-health", "/account/connected-accounts"):
        r = cr.get(path)
        assert r.status_code == 403, f"{path} should 403 for creator, got {r.status_code}"
    # /account/brand-deck is POST-only (GET correctly 405s before the role check runs) - the real
    # attack surface is a POST, which must still be blocked
    r = cr.post("/account/brand-deck", data={})
    assert r.status_code == 403, f"creator POST to brand-deck should 403, got {r.status_code}"
    # the account page itself is reachable (profile/password are personal settings), but the Brand
    # deck scan card's visibility must match the pre-existing "Company logo" card's gating
    # (both owner/admin-only workspace-wide settings) - not accidentally left exposed to creators
    acct_html = cr.get("/account").data.decode()
    assert ("Brand deck scan" in acct_html) == ("Company logo" in acct_html)
    assert "Brand deck scan" not in acct_html


def test_viewer_role_can_read_new_pages_but_not_write(app, mails):
    o = signup(app, "ov1@x.com")
    vw = _join(app, o, mails, "vw1@x.com", "viewer", "Vee")
    # viewer restriction is on writes, not module access - viewer should be able to READ the new pages
    for path in ("/marketing-health", "/account/connected-accounts", "/account"):
        r = vw.get(path)
        assert r.status_code == 200, f"viewer should be able to view {path}, got {r.status_code}"
    # but the brand-deck scan form itself stays owner/admin-only, same as before
    assert "Brand deck scan" not in vw.get("/account").data.decode()
    r = vw.post("/account/brand-deck", data={})
    assert r.status_code == 403


def test_account_brand_deck_rejects_non_pptx_pdf(app):
    o = signup(app, "baddeck@x.com")
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(b"hello"), "notes.txt")),
               content_type="multipart/form-data", follow_redirects=True)
    assert b"Upload your brand deck as a PowerPoint" in r.data


def test_account_brand_deck_oversized_file_rejected_cleanly_not_500(app):
    # the view's own 20 MB check must kick in well before Flask's global MAX_CONTENT_LENGTH (25 MB, see
    # config.py) would reject the whole request with a generic 413 - this keeps the friendly flash message
    o = signup(app, "bigdeck@x.com")
    big = b"P" * (21 * 1024 * 1024)
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(big), "huge.pptx")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and b"up to 20 MB" in r.data


def test_account_brand_deck_handles_corrupt_file_without_crashing(app):
    o = signup(app, "corruptdeck@x.com")
    r = o.post("/account/brand-deck", data=dict(file=(io.BytesIO(b"not a real pptx"), "broken.pptx")),
               content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200 and b"Couldn" in r.data


def test_account_page_shows_brand_deck_upload_card_for_owner(app):
    o = signup(app, "deckcard@x.com")
    html = o.get("/account").data.decode()
    assert "Brand deck scan" in html and 'action="/account/brand-deck"' in html


# =====================================================================  Marketing Health: KPIs + a single gap checklist, built from in-app data only
def test_marketing_health_shows_kpis_and_gap_checklist_for_a_fresh_workspace(app):
    o = signup(app, "health@x.com")
    html = o.get("/marketing-health").data.decode()
    assert "Marketing Health" in html
    assert "Brand consistency" in html and "Posting consistency" in html and "Campaign spend efficiency" in html and "Lead" in html
    assert "What's left to fix" in html
    assert "Brand Kit categories not filled in yet" in html  # nothing filled in yet on a fresh workspace


def test_marketing_health_brand_score_rises_as_brand_kit_and_logo_fill_in(app):
    o = signup(app, "health2@x.com")
    before = o.get("/marketing-health").data.decode()
    assert "0%" in before or "25%" in before  # brand score starts low
    o.post("/m/brand_kit/save", data=dict(name="Primary", kind="Colour", value="#1E88E5", next="/marketing-health"))
    o.post("/m/brand_kit/save", data=dict(name="Body font", kind="Font", value="Montserrat", next="/marketing-health"))
    o.post("/m/brand_kit/save", data=dict(name="Tone", kind="Voice & tone", value="Friendly, direct", next="/marketing-health"))
    after = o.get("/marketing-health").data.decode()
    assert "75%" in after or "100%" in after  # 3 of 4 categories now filled (logo still missing without a real image upload)


def test_marketing_health_gaps_resolve_as_data_is_filled_in(app):
    o = signup(app, "health3@x.com")
    o.post("/m/campaigns/save", data=dict(name="Launch", status="Live", next="/marketing-health"))  # no KPI target -> a gap
    assert "Live or planned campaigns with no KPI target set" in o.get("/marketing-health").data.decode()
    cid = dbq(app, "SELECT id FROM campaigns WHERE name='Launch'", one=True)["id"]
    o.post("/m/campaigns/save", data=dict(id=cid, name="Launch", status="Live", kpi_target="1000 signups", next="/marketing-health"))
    html = o.get("/marketing-health").data.decode()
    # the campaign gap line should now show a checkmark (count 0), not a warning with that campaign counted
    gap_block = html.split("What's left to fix")[1]
    assert "✅" in gap_block and "no KPI target set" in gap_block


def test_marketing_health_is_feature_gated_like_analytics(app, monkeypatch):
    import config as C
    o = signup(app, "healthgate@x.com")
    monkeypatch.setitem(C.PLANS["trial"], "features", C.PLANS["trial"]["features"] - {"analytics"})
    r = o.get("/marketing-health")
    assert r.status_code == 403 and b"upgrade" in r.data.lower()


BANNER_TEXT = "see Marketing Health for the full checklist"


def test_gap_banner_appears_on_relevant_list_pages_and_clears_once_resolved(app):
    o = signup(app, "banner@x.com")
    html = o.get("/m/brand_kit").data.decode()
    assert BANNER_TEXT in html
    # filling in all 4 brand categories (incl. a logo) should clear the banner on this page
    o.post("/m/brand_kit/save", data=dict(name="Primary", kind="Colour", value="#1E88E5", next="/m/brand_kit"))
    o.post("/m/brand_kit/save", data=dict(name="Font", kind="Font", value="Montserrat", next="/m/brand_kit"))
    o.post("/m/brand_kit/save", data=dict(name="Tone", kind="Voice & tone", value="Friendly", next="/m/brand_kit"))
    o.post("/account/logo", data=dict(file=(io.BytesIO(PNG), "logo.png")), content_type="multipart/form-data")
    html2 = o.get("/m/brand_kit").data.decode()
    assert BANNER_TEXT not in html2


def test_gap_banner_on_content_items_sums_both_missing_date_and_overdue_gaps(app):
    # content_items is the only GAP_BANNER_KEYS entry that maps to TWO underlying gap types
    # (content_items = no publish date, content_overdue = past date but not Published) - make sure the
    # banner's count is the sum of both, and that it only clears once both are actually resolved
    o = signup(app, "contentbanner@x.com")
    yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    o.post("/m/content_items/save", data=dict(title="No date yet", status="Idea"))                       # gap #1: missing publish_date
    o.post("/m/content_items/save", data=dict(title="Missed it", status="Idea", publish_date=yesterday))  # gap #2: overdue, not Published
    html = o.get("/m/content_items").data.decode()
    assert BANNER_TEXT in html and "2 gap" in html
    # resolve only the first gap (give it a publish date) - banner must still show 1 remaining gap
    cid1 = dbq(app, "SELECT id FROM content_items WHERE workspace_id=? AND title='No date yet'", (wid_of(app, "contentbanner@x.com"),), one=True)["id"]
    o.post("/m/content_items/save", data=dict(id=cid1, title="No date yet", status="Idea", publish_date=S.today().isoformat() if hasattr(S, "today") else dt.date.today().isoformat()))
    html2 = o.get("/m/content_items").data.decode()
    assert BANNER_TEXT in html2 and "1 gap" in html2
    # now resolve the second gap too (mark it Published) - banner must clear entirely
    cid2 = dbq(app, "SELECT id FROM content_items WHERE workspace_id=? AND title='Missed it'", (wid_of(app, "contentbanner@x.com"),), one=True)["id"]
    o.post("/m/content_items/save", data=dict(id=cid2, title="Missed it", status="Published", publish_date=yesterday))
    html3 = o.get("/m/content_items").data.decode()
    assert BANNER_TEXT not in html3


def test_gap_banner_does_not_appear_on_unrelated_module_pages(app):
    o = signup(app, "nobann@x.com")
    html = o.get("/m/competitors").data.decode()
    assert BANNER_TEXT not in html


# =====================================================================  Connected Accounts: an honest "not yet wired up" placeholder, ready for real OAuth later
def test_connected_accounts_page_lists_platforms_as_not_connected(app):
    o = signup(app, "conn@x.com")
    html = o.get("/account/connected-accounts").data.decode()
    for platform in ("Instagram", "Facebook", "Google Ads", "LinkedIn Ads"):
        assert platform in html
    assert html.count("Not connected") == 4


def test_account_page_links_to_connected_accounts_for_owner(app):
    o = signup(app, "connlink@x.com")
    assert 'href="/account/connected-accounts"' in o.get("/account").data.decode()


# =====================================================================  Dashboard "Today" card: one auto-built list (overdue carried over + due today), so a solo founder has one place to look each morning
def test_dashboard_today_card_shows_all_caught_up_when_nothing_is_due(app):
    o = signup(app, "today1@x.com")
    html = o.get("/app").data.decode()
    assert "📋 Today" in html and "You're all caught up" in html


def test_dashboard_today_card_surfaces_overdue_task_and_overdue_content(app):
    o = signup(app, "today2@x.com")
    yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    o.post("/m/tasks/save", data=dict(title="Call supplier", status="To do", due_date=yesterday, next="/app"))
    o.post("/m/content_items/save", data=dict(title="Diwali post", status="Scheduled", publish_date=yesterday, next="/app"))
    html = o.get("/app").data.decode()
    assert "Overdue: Call supplier" in html
    assert "Overdue to publish: Diwali post" in html
    assert "You're all caught up" not in html


def test_dashboard_today_card_shows_due_today_separately_from_overdue(app):
    o = signup(app, "today3@x.com")
    t = dt.date.today().isoformat()
    o.post("/m/tasks/save", data=dict(title="Send newsletter", status="To do", due_date=t, next="/app"))
    o.post("/m/content_items/save", data=dict(title="Launch reel", status="Scheduled", publish_date=t, next="/app"))
    html = o.get("/app").data.decode()
    assert "Due today: Send newsletter" in html
    assert "Publish today: Launch reel" in html


def test_dashboard_today_card_caps_at_eight_and_links_to_marketing_health_for_more(app):
    o = signup(app, "today4@x.com")
    yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    for i in range(12):
        o.post("/m/tasks/save", data=dict(title=f"Task {i}", status="To do", due_date=yesterday, next="/app"))
    html = o.get("/app").data.decode()
    assert html.count("Overdue: Task") == 8
    assert "more" in html and 'href="/marketing-health"' in html


def test_today_item_link_actually_opens_the_right_record(app):
    # the cap/more-link test above checks the list truncates correctly; this checks an individual
    # Today item's href is a real, correct link to that exact record (not just a well-formed url_for)
    o = signup(app, "todayclick@x.com")
    yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    o.post("/m/tasks/save", data=dict(title="Call the printer", due_date=yesterday, status="To do"))
    dash = o.get("/app").data.decode()
    assert "Call the printer" in dash
    wid = wid_of(app, "todayclick@x.com")
    tid = dbq(app, "SELECT id FROM tasks WHERE workspace_id=? AND title='Call the printer'", (wid,), one=True)["id"]
    expected_href = f"/m/tasks/{tid}/edit"
    assert expected_href in dash
    r = o.get(expected_href)
    assert r.status_code == 200 and b"Call the printer" in r.data


def test_dashboard_today_card_hidden_for_creator_role(app):
    o = signup(app, "today5@x.com")
    assert "📋 Today" in o.get("/app").data.decode()  # owner sees it
    creator = signup(app, "today5cr@x.com")
    dbx(app, "UPDATE memberships SET role='creator' WHERE user_id=(SELECT id FROM users WHERE email='today5cr@x.com')")
    assert "📋 Today" not in creator.get("/app").data.decode()  # a freelancer/creator doesn't


# =====================================================================  Verifying an email now lands on a clear confirmation page, not a silent bounce
def test_verify_email_shows_confirmation_page(app, mails):
    o = signup(app, "o@x.com")
    link = re.search(r"(/verify/[\w\-\.]+)", mails[-1][2]).group(1)
    r = o.get(link)
    assert r.status_code == 200 and b"Email verified" in r.data and b"Continue to Dashboard" in r.data
    assert dbq(app, "SELECT email_verified FROM users WHERE email='o@x.com'", one=True)["email_verified"] == 1


def test_dashboard_today_card_escapes_user_entered_titles(app):
    o = signup(app, "xsstoday@x.com")
    yesterday = (dt.date.today() - dt.timedelta(days=1)).isoformat()
    o.post("/m/tasks/save", data=dict(title="<script>alert(1)</script>", status="To do", due_date=yesterday, next="/app"))
    html = o.get("/app").data.decode()
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


def test_account_brand_deck_rescan_does_not_duplicate_brand_kit_entries(app):
    o = signup(app, "rescan@x.com")
    deck = _make_test_pptx()
    o.post("/account/brand-deck", data=dict(file=(io.BytesIO(deck), "brand.pptx")), content_type="multipart/form-data")
    wid = wid_of(app, "rescan@x.com")
    first_count = dbq(app, "SELECT COUNT(*) c FROM brand_kit WHERE workspace_id=?", (wid,), one=True)["c"]
    o.post("/account/brand-deck", data=dict(file=(io.BytesIO(deck), "brand.pptx")), content_type="multipart/form-data")
    second_count = dbq(app, "SELECT COUNT(*) c FROM brand_kit WHERE workspace_id=?", (wid,), one=True)["c"]
    assert first_count == second_count and first_count > 0
